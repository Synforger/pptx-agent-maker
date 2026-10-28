"""`compare` on real decks, drawn by LibreOffice: the pages that did not change count zero pixels.

頼みの元は「宣言の 1 頁だけを差し替えた時、ほかの頁が変わっていないこと」を確かめること。
字を含む頁は LibreOffice が書体を置き換える時に揺れることがあるので、ここで使う資料は
**図形だけ**にして、「0 画素」を厳密に言える頁で見る。
"""

from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "preview" / "src"))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.compare import bake  # noqa: E402

#: 試験の絵の細かさ (= 実際は 200。動いた頁が分かれば足りるので粗く焼いて速くする)
TEST_DPI = 40

COLOURS = [RGBColor(0xC0, 0x30, 0x30), RGBColor(0x30, 0x90, 0x40), RGBColor(0x30, 0x40, 0xC0)]


def write_deck(path: Path, *, pages: int = 3, moved: int | None = None) -> Path:
    """A deck of `pages` pages, each a coloured square; the `moved`th page has its square somewhere else."""
    deck = Presentation()
    blank = deck.slide_layouts[6]
    for number in range(1, pages + 1):
        slide = deck.slides.add_slide(blank)
        left = Inches(6 if number == moved else 1)
        square = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1), Inches(2), Inches(2))
        square.fill.solid()
        square.fill.fore_color.rgb = COLOURS[(number - 1) % len(COLOURS)]
    deck.save(str(path))
    return path


@unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftoppm"),
                     "draws decks for real: needs LibreOffice (soffice) and poppler (pdftoppm), which the machine that runs `task ci` before a commit has")
class ComparingRealDecks(unittest.TestCase):

    @classmethod
    def setUpClass(cls) -> None:
        cls.tmp = tempfile.TemporaryDirectory()
        folder = Path(cls.tmp.name)
        cls.before = write_deck(folder / "before.pptx")
        cls.after = write_deck(folder / "after.pptx", moved=2)
        # 常駐のプレビューや利用者の cache に触れない (= プロファイルの置き場を試験の中に移す)
        patches = [mock.patch.dict(os.environ, {"XDG_CACHE_HOME": str(folder / "cache")}),
                   mock.patch.object(bake, "COMPARE_DPI", TEST_DPI)]
        for patch in patches:
            patch.start()
            cls.addClassCleanup(patch.stop)
        cls.addClassCleanup(cls.tmp.cleanup)

    def _run(self, *argv: str) -> tuple[int, str, str]:
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main(list(argv))
        return code, out.getvalue(), err.getvalue()

    def test_the_same_deck_twice_has_no_differing_pixel_on_any_page_and_the_command_ends_with_0(self) -> None:
        code, said, _ = self._run("compare", str(self.before), str(self.before))
        self.assertEqual(code, 0, said)
        self.assertEqual(said.splitlines(), [f"page {number}: 0 px differ (0.0000%)" for number in (1, 2, 3)]
                         + ["changed 0 of 3 pages"])

    def test_only_the_page_that_was_changed_has_differing_pixels_and_the_command_ends_with_1(self) -> None:
        code, said, _ = self._run("compare", str(self.before), str(self.after))
        self.assertEqual(code, 1, said)
        lines = said.splitlines()
        self.assertEqual(lines[0], "page 1: 0 px differ (0.0000%)")
        self.assertFalse(lines[1].startswith("page 2: 0 px"), said)
        self.assertEqual(lines[2], "page 3: 0 px differ (0.0000%)")
        self.assertEqual(lines[-1], "changed 1 of 3 pages")

    def test_a_deck_that_is_not_there_is_not_a_difference(self) -> None:
        code, said, complained = self._run("compare", str(self.before), str(self.before.with_name("nothing.pptx")))
        self.assertEqual(code, 2, "ending with 1 would read as 'a page changed'")
        self.assertEqual(said, "")
        self.assertIn("deck not found", complained)


if __name__ == "__main__":
    unittest.main()
