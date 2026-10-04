"""The timeline: periods left to right, lanes top to bottom, bars and marks on the lanes.

⚠ **この型は、絵で描いて貼った計画の頁を人が直せなかったために在る。**だから見張るのは
3 つ ― 文字が全部 pptx の文字として残ること、位置が期間の番号だけで決まること、
収まらないものを縮めずに止めること。
"""

from __future__ import annotations

import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))
from support.pages import named  # noqa: E402

from pptx_agent_maker import checks  # noqa: E402
from pptx_agent_maker.layout import types  # noqa: E402
from pptx_agent_maker.layout.parts.look import LIGHT, TONES  # noqa: E402
from pptx_agent_maker.layout.parts.elements import Bar, Diamond, Text  # noqa: E402
from pptx_agent_maker.layout.parts.page import PageFullError  # noqa: E402
from pptx_agent_maker.deck.base.look import SLOTS, look_of  # noqa: E402
from pptx_agent_maker.layout.base.tokens import DEFAULT, theme_from  # noqa: E402
from pptx_agent_maker.layout.types.bodies.timeline import ROOMY  # noqa: E402
from pptx_agent_maker.layout.types.core.registry import PageTypeError  # noqa: E402
from pptx_agent_maker.write import add_page, new_deck, save  # noqa: E402

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]


def page(lanes, *, periods=MONTHS, **extra):
    return types.build({"type": "timeline", "title": "A plan", "periods": periods,
                        "lanes": lanes, **extra}, lambda name: Path(name), lambda path: 1.0)


def lane(name="Work", bars=(), marks=()):
    return {"name": name, "bars": list(bars), "marks": list(marks)}


def bar(start, end, text="step", **more):
    return {"from": start, "to": end, "text": text, **more}


def of(built, kind):
    return [element for element in built.build() if element.kind == kind]


def columns(built):
    """The left edge of each period's column, read off the period names."""
    return [named(built, month).rect.left - DEFAULT.spacing.bar_gap for month in MONTHS]


class WhereThingsLand(unittest.TestCase):
    """A position is a period number; the type turns it into a place."""

    def test_a_bar_runs_from_the_start_of_its_first_period_to_the_end_of_its_last(self):
        built = page([lane(bars=[bar(1, 3, "build")])])
        edges, gap = columns(built), DEFAULT.spacing.bar_gap
        drawn = named(built, "build").rect
        self.assertEqual(drawn.left, edges[1] + gap)
        self.assertEqual(drawn.right, edges[3] - gap)

    def test_half_a_period_is_half_a_column(self):
        built = page([lane(bars=[bar(0, 0.5, "a"), bar(0.5, 1, "b")])])
        first, second = named(built, "a").rect, named(built, "b").rect
        self.assertAlmostEqual(first.width, second.width, delta=2)
        self.assertLess(first.right, second.left)

    def test_bars_that_only_touch_share_a_row_and_bars_that_overlap_do_not(self):
        built = page([lane(bars=[bar(0, 2, "first"), bar(2, 4, "second"), bar(3, 5, "third")])])
        first, second, third = (named(built, t).rect for t in ("first", "second", "third"))
        self.assertEqual(first.top, second.top)
        self.assertGreaterEqual(third.top, second.bottom)
        self.assertFalse(second.overlaps(third))

    def test_lanes_stack_top_to_bottom_in_the_order_written(self):
        built = page([lane("One", [bar(0, 6, "a")]), lane("Two", [bar(0, 6, "b")])])
        self.assertLess(named(built, "a").rect.bottom, named(built, "b").rect.top)
        self.assertLess(named(built, "One").rect.top, named(built, "Two").rect.top)

    def test_a_mark_is_a_diamond_centred_on_its_moment_with_its_name_beside_it(self):
        built = page([lane(bars=[bar(4, 6, "later")], marks=[{"at": 2, "text": "release"}])])
        diamond = [e for e in built.build() if isinstance(e, Diamond)][0].rect
        self.assertEqual(diamond.width, diamond.height)
        self.assertAlmostEqual(diamond.left + diamond.width // 2, columns(built)[2], delta=1)
        self.assertGreater(named(built, "release").rect.left, diamond.right)

    def test_marks_at_both_ends_of_the_axis_stay_on_the_page(self):
        built = page([lane(marks=[{"at": 0, "text": "start"}, {"at": 6, "text": "end"}])])
        frame = built.theme.frame()
        for element in built.build():
            self.assertTrue(frame.contains(element.rect), f"{element.kind} left the frame")
        # 右端では名前を置く場所が左にしか無い
        diamonds = sorted((e.rect for e in built.build() if isinstance(e, Diamond)),
                          key=lambda r: r.left)
        self.assertLess(named(built, "end").rect.right, diamonds[-1].left)
        self.assertEqual("right", named(built, "end").align)

    def test_everything_on_a_crowded_page_stays_inside_the_frame(self):
        lanes = [lane(f"Lane {n}", [bar(0, 2, "one"), bar(1, 4, "two"), bar(3.5, 6, "three")],
                      [{"at": n, "text": "check"}]) for n in range(1, 4)]
        built = page(lanes, phases=[{"from": 0, "to": 3, "label": "early"},
                                    {"from": 3, "to": 6, "label": "late"}],
                     milestones=[{"at": 2.5, "text": "review"}, {"at": 6, "text": "close"}])
        frame = built.theme.frame()
        for element in built.build():
            self.assertTrue(frame.contains(element.rect), f"{element.kind} left the frame")


class TheWordsStayWords(unittest.TestCase):
    """Every name is text a person can edit, and a bar carries its own."""

    def test_a_bar_holds_its_name_in_the_same_shape(self):
        drawn = named(page([lane(bars=[bar(0, 3, "build the core")])]), "build the core")
        self.assertIsInstance(drawn, Bar)

    def test_no_name_is_set_below_the_floor(self):
        built = page([lane("A lane", [bar(0, 1, "x"), bar(2, 2.2, "a long name for a short bar")],
                           [{"at": 5, "text": "mark"}])], milestones=[{"at": 3, "text": "date"}])
        sizes = {e.size for e in built.build() if isinstance(e, (Bar, Text))}
        self.assertGreaterEqual(min(sizes), DEFAULT.type.minimum)

    def test_a_name_too_long_for_one_line_takes_two_inside_the_bar(self):
        short = named(page([lane(bars=[bar(0, 1.5, "short")])]), "short").rect
        text = "hand over the face input to the team"
        long = named(page([lane(bars=[bar(0, 1.5, text)])]), text)
        self.assertIsInstance(long, Bar)
        self.assertGreater(long.rect.height, short.height)
        fits = "hand over the input"
        self.assertEqual(short.height, named(page([lane(bars=[bar(0, 1.5, fits)])]), fits).rect.height)

    def test_a_single_word_is_never_folded_in_the_middle(self):
        """1 語の名前を 2 行に折ると、語の途中で切れる (= 焼いて初めて出た)。"""
        built = page([lane(bars=[bar(2, 2.3, "auditing")])], periods=MONTHS)
        self.assertIsInstance(named(built, "auditing"), Text)
        self.assertEqual(1, len([e for e in built.build() if isinstance(e, Bar) and not e.text]))

    def test_a_name_too_long_for_two_lines_goes_beside_the_bar_uncut(self):
        text = "a name far too long for a bar this short"
        built = page([lane(bars=[bar(1, 1.5, text)])])
        label = named(built, text)
        self.assertIsInstance(label, Text)
        empty = [e for e in built.build() if isinstance(e, Bar) and not e.text]
        self.assertEqual(1, len(empty))
        self.assertGreaterEqual(label.rect.left, empty[0].rect.right)
        self.assertEqual(label.rect.top, empty[0].rect.top)

    def test_a_name_beside_a_bar_keeps_the_next_bar_off_it(self):
        text = "a name far too long for a bar this short"
        built = page([lane(bars=[bar(1, 1.5, text), bar(2, 4, "next")])])
        self.assertGreaterEqual(named(built, "next").rect.top, named(built, text).rect.bottom)

    def test_a_tentative_bar_is_a_dashed_outline_on_paper(self):
        built = page([lane(bars=[bar(0, 2, "sure"), bar(2, 4, "maybe", tentative=True)])])
        sure, maybe = named(built, "sure"), named(built, "maybe")
        self.assertTrue(maybe.dashed and maybe.outline)
        self.assertEqual(DEFAULT.palette.paper, maybe.fill)
        self.assertFalse(sure.dashed)
        self.assertNotEqual(DEFAULT.palette.paper, sure.fill)

    def test_neighbouring_lanes_differ_in_tone_and_use_only_palette_colours(self):
        built = page([lane("One", [bar(0, 6, "a")]), lane("Two", [bar(0, 6, "b")])])
        self.assertNotEqual(named(built, "a").fill, named(built, "b").fill)
        palette = set(vars(DEFAULT.palette).values())
        # 薄い地の枠は palette がその地から出す色 (= 頁が自分で選んだ色ではない)
        palette |= {DEFAULT.palette.edge(getattr(DEFAULT.palette, tone)) for tone in LIGHT}
        for element in built.build():
            for colour in (getattr(element, "colour", ""), getattr(element, "fill", ""),
                           getattr(element, "outline", "")):
                if colour:
                    self.assertIn(colour, palette, f"{element.kind} uses a colour of its own")


class TheToneOfALane(unittest.TestCase):
    """A lane may say which of the palette's roles its bars take; unsaid, two grounds alternate."""

    def lanes(self, *tones):
        return [{"name": f"Lane {n}", "bars": [bar(0, 6, f"bar {n}")], **({"tone": tone} if tone else {})}
                for n, tone in enumerate(tones, start=1)]

    def test_unsaid_the_lanes_alternate_between_two_light_grounds(self):
        built = page(self.lanes(None, None, None, None))
        palette = DEFAULT.palette
        self.assertEqual([palette.box, palette.band, palette.box, palette.band],
                         [named(built, f"bar {n}").fill for n in range(1, 5)])

    def test_a_lane_takes_the_tone_it_names_and_leaves_the_others_alone(self):
        built = page(self.lanes(None, "tint", None, "accent"))
        palette = DEFAULT.palette
        self.assertEqual([palette.box, palette.tint, palette.box, palette.accent],
                         [named(built, f"bar {n}").fill for n in range(1, 5)])

    def test_every_bar_of_the_lane_takes_it(self):
        built = page([{"name": "One", "tone": "tint",
                       "bars": [bar(0, 2, "a"), bar(2, 4, "b"), bar(1, 3, "c")]}])
        self.assertEqual({DEFAULT.palette.tint}, {named(built, t).fill for t in "abc"})

    def test_a_name_on_a_dark_ground_is_set_in_the_papers_colour(self):
        palette = DEFAULT.palette
        for tone, ink in (("box", palette.ink), ("band", palette.ink), ("tint", palette.ink),
                          ("accent", palette.paper), ("good", palette.paper), ("bad", palette.paper)):
            with self.subTest(tone):
                drawn = named(page(self.lanes(tone)), "bar 1")
                self.assertEqual(getattr(palette, tone), drawn.fill)
                self.assertEqual(ink, drawn.colour)

    def test_three_lanes_can_be_told_apart_by_three_light_grounds(self):
        built = page(self.lanes("box", "band", "tint"))
        fills = [named(built, f"bar {n}").fill for n in range(1, 4)]
        self.assertEqual(3, len(set(fills)))
        self.assertEqual({DEFAULT.palette.ink}, {named(built, f"bar {n}").colour for n in range(1, 4)})

    def test_a_tentative_bar_stays_a_dashed_outline_whatever_the_lanes_tone(self):
        built = page([{"name": "One", "tone": "accent",
                       "bars": [bar(0, 2, "sure"), bar(2, 4, "maybe", tentative=True)]}])
        maybe = named(built, "maybe")
        self.assertTrue(maybe.dashed)
        self.assertEqual(DEFAULT.palette.paper, maybe.fill)
        self.assertEqual(DEFAULT.palette.muted, maybe.colour)

    def test_a_name_beside_a_dark_bar_stays_dark_on_the_paper(self):
        text = "a name far too long for a bar this short"
        built = page([{"name": "One", "tone": "accent", "bars": [bar(1, 1.3, text)]}])
        self.assertIsInstance(named(built, text), Text)
        self.assertEqual(DEFAULT.palette.ink, named(built, text).colour)

    def test_a_tone_that_is_not_a_role_is_refused_and_says_which_lane(self):
        with self.assertRaises(PageTypeError) as raised:
            page(self.lanes(None, "green"))
        said = str(raised.exception)
        self.assertIn("lane 2", said)
        for tone in TONES:
            self.assertIn(tone, said)

    def test_the_projects_own_tint_reaches_the_bars(self):
        theme = theme_from({"palette": {"tint": "abcdef"}})
        built = types.build({"type": "timeline", "title": "A plan", "periods": MONTHS,
                             "lanes": self.lanes("tint")}, lambda name: Path(name),
                            lambda path: 1.0, theme)
        self.assertEqual("ABCDEF", named(built, "bar 1").fill)

    def test_the_tint_is_never_read_from_a_templates_theme(self):
        """テーマの空いた枠には濃い色が入っていることが多い。薄い地をそこから読むと名前が読めなくなる。"""
        self.assertNotIn("tint", [name for _slot, name in SLOTS])
        template = REPO / "src" / "pptx_agent_maker" / "templates" / "project" / "specimen.pptx"
        self.assertNotIn("tint", look_of(template).get("palette", {}))


class TheEdgeOfABar(unittest.TestCase):
    """A bar on a light ground carries a line of its own colour, deeper.

    ⚠ **薄い地は列の地とも紙とも近い。**枠が無かった間は、棒の端が背景に溶けて、どこからどこまでの
    棒なのか読めなかった (= 実物のデッキで差し戻された)。
    """

    def lanes(self, *tones):
        return [{"name": f"Lane {n}", "tone": tone, "bars": [bar(0, 6, f"bar {n}")]}
                for n, tone in enumerate(tones, start=1)]

    def test_a_bar_on_a_light_ground_has_a_solid_edge_of_its_own_ground_made_deeper(self):
        palette = DEFAULT.palette
        built = page(self.lanes(*LIGHT))
        for n, tone in enumerate(LIGHT, start=1):
            with self.subTest(tone):
                drawn = named(built, f"bar {n}")
                self.assertEqual(palette.edge(getattr(palette, tone)), drawn.outline)
                self.assertFalse(drawn.dashed)

    def test_the_edge_differs_from_the_ground_the_paper_and_the_wash(self):
        palette = DEFAULT.palette
        for tone in LIGHT:
            with self.subTest(tone):
                edge = palette.edge(getattr(palette, tone))
                self.assertNotIn(edge, {getattr(palette, tone), palette.paper, palette.wash})

    def test_the_three_light_grounds_have_three_different_edges(self):
        built = page(self.lanes(*LIGHT))
        self.assertEqual(3, len({named(built, f"bar {n}").outline for n in range(1, 4)}))

    def test_a_bar_on_a_dark_ground_has_none(self):
        for tone in ("accent", "good", "bad"):
            with self.subTest(tone):
                self.assertEqual("", named(page(self.lanes(tone)), "bar 1").outline)

    def test_unsaid_the_alternating_grounds_each_take_their_own_edge(self):
        palette = DEFAULT.palette
        built = page([lane("One", [bar(0, 6, "a")]), lane("Two", [bar(0, 6, "b")])])
        self.assertEqual([palette.edge(palette.box), palette.edge(palette.band)],
                         [named(built, text).outline for text in "ab"])

    def test_a_tentative_bar_keeps_its_own_dashed_line(self):
        built = page([{"name": "One", "tone": "tint",
                       "bars": [bar(0, 2, "sure"), bar(2, 4, "maybe", tentative=True)]}])
        maybe = named(built, "maybe")
        self.assertTrue(maybe.dashed)
        self.assertEqual(DEFAULT.palette.muted, maybe.outline)

    def test_the_projects_own_ground_brings_its_edge_along(self):
        theme = theme_from({"palette": {"tint": "abcdef"}})
        built = types.build({"type": "timeline", "title": "A plan", "periods": MONTHS,
                             "lanes": self.lanes("tint")}, lambda name: Path(name),
                            lambda path: 1.0, theme)
        drawn = named(built, "bar 1")
        self.assertEqual(theme.palette.edge("ABCDEF"), drawn.outline)
        self.assertNotEqual(DEFAULT.palette.edge(DEFAULT.palette.tint), drawn.outline)

    def test_the_edge_takes_no_room(self):
        """枠は棒の端の上に引く線で、棒の位置も大きさも変えない。"""
        light, dark = page(self.lanes("box")), page(self.lanes("accent"))
        self.assertEqual(named(dark, "bar 1").rect, named(light, "bar 1").rect)

    def test_the_period_cells_and_the_phase_bands_have_none(self):
        built = page(self.lanes("box"), phases=[{"from": 0, "to": 6, "label": "this half"}])
        for kind in ("period", "span"):
            self.assertEqual({""}, {piece.outline for piece in of(built, kind)}, kind)


class ABarAcrossLanes(unittest.TestCase):
    """A bar may cover the lanes under its own: one stretch shared by lanes that part later.

    ⚠ **またがる棒は、覆うレーンの高さを全部取る。**同じ時期に同じレーンへ置かれた物は、段を
    分けて逃がせない (= 重ねて描くことになる) ので、縮めも重ねもせずに拒む。
    """

    LONG = "a name far too long for a bar this short"

    def shared(self, **more):
        """Two lanes that share months 0-3 and part after: (the page, the shared bar)."""
        built = page([lane("Inside", [bar(0, 3, "shared", spans=2, **more), bar(3, 6, "ours")]),
                      lane("Outside", [bar(3, 6, "theirs")])])
        return built, named(built, "shared")

    def test_it_runs_from_the_top_of_its_own_lane_to_the_bottom_of_the_last_lane_it_spans(self):
        built, shared = self.shared()
        self.assertEqual(named(built, "ours").rect.top, shared.rect.top)
        self.assertEqual(named(built, "theirs").rect.bottom, shared.rect.bottom)

    def test_its_ends_fall_where_any_bars_would(self):
        built = page([lane("One", [bar(0, 3, "shared", spans=2)]), lane("Two"),
                      lane("Three", [bar(0, 3, "alone")])])
        shared, alone = named(built, "shared").rect, named(built, "alone").rect
        self.assertEqual((alone.left, alone.right), (shared.left, shared.right))

    def test_it_is_one_shape_with_its_name_inside(self):
        built, shared = self.shared()
        self.assertIsInstance(shared, Bar)
        self.assertEqual(1, len([e for e in built.build() if getattr(e, "text", None) == "shared"]))

    def test_it_covers_every_row_of_a_lane_that_has_several(self):
        built = page([lane("One", [bar(0, 2, "shared", spans=2), bar(2, 5, "a"), bar(3, 6, "b")]),
                      lane("Two", [bar(2, 5, "c"), bar(3, 6, "d")])])
        shared = named(built, "shared").rect
        self.assertEqual(named(built, "a").rect.top, shared.top)
        self.assertGreater(named(built, "d").rect.top, named(built, "c").rect.top)
        self.assertEqual(named(built, "d").rect.bottom, shared.bottom)

    def test_it_can_cover_more_than_two(self):
        built = page([lane("One", [bar(0, 3, "shared", spans=3), bar(3, 6, "a")]),
                      lane("Two", [bar(3, 6, "b")]), lane("Three", [bar(3, 6, "c")]),
                      lane("Four", [bar(0, 6, "d")])])
        shared = named(built, "shared").rect
        self.assertEqual(named(built, "c").rect.bottom, shared.bottom)
        self.assertLess(shared.bottom, named(built, "d").rect.top)

    def test_one_lane_is_what_a_bar_spans_when_nothing_is_said(self):
        said = page([lane(bars=[bar(0, 3, "a", spans=1)])])
        unsaid = page([lane(bars=[bar(0, 3, "a")])])
        self.assertEqual(unsaid.build(), said.build())

    def test_it_takes_the_tone_of_the_lane_it_is_written_in(self):
        palette = DEFAULT.palette
        built = page([{"name": "One", "tone": "tint", "bars": [bar(0, 6, "shared", spans=2)]},
                      {"name": "Two", "tone": "accent"}])
        shared = named(built, "shared")
        self.assertEqual((palette.tint, palette.ink, palette.edge(palette.tint)),
                         (shared.fill, shared.colour, shared.outline))

    def test_a_tentative_one_is_still_a_dashed_line_on_paper(self):
        _built, shared = self.shared(tentative=True)
        self.assertTrue(shared.dashed)
        self.assertEqual(DEFAULT.palette.paper, shared.fill)

    def test_it_is_laid_over_the_line_between_the_lanes_it_spans(self):
        """区切りの線を棒が隠すから 1 本に見える。線が後なら、棒を横切って 2 本に見える。"""
        built, shared = self.shared()
        elements = built.build()
        between = [e for e in elements if e.kind == "rule" and e.rect.width > e.rect.height
                   and shared.rect.top < e.rect.top < shared.rect.bottom]
        self.assertTrue(between)
        self.assertGreater(elements.index(shared), max(elements.index(line) for line in between))

    def test_a_lane_that_is_only_spanned_needs_nothing_of_its_own(self):
        built = page([lane("One", [bar(0, 6, "shared", spans=2)]), lane("Two")])
        two = named(built, "Two").rect
        shared = named(built, "shared").rect
        self.assertGreater(shared.bottom, two.top)
        self.assertGreater(shared.height, 2 * DEFAULT.line_height(named(built, "shared").size))

    def test_a_lane_with_nothing_in_it_and_nothing_over_it_is_still_refused(self):
        with self.assertRaises(PageTypeError) as raised:
            page([lane("One", [bar(0, 6, "shared", spans=2)]), lane("Two"), lane("Three")])
        self.assertIn("lane 3 has neither bars nor marks", str(raised.exception))

    def test_bars_that_only_touch_its_ends_are_taken(self):
        built = page([lane("One", [bar(0, 2, "before"), bar(2, 4, "shared", spans=2), bar(4, 6, "after")]),
                      lane("Two", [bar(0, 2, "under before"), bar(4, 6, "under after")])])
        shared = named(built, "shared").rect
        for text in ("before", "under before"):
            self.assertLess(named(built, text).rect.right, shared.left)
        for text in ("after", "under after"):
            self.assertGreater(named(built, text).rect.left, shared.right)

    def refused(self, lanes, *said):
        with self.assertRaises(PageTypeError) as raised:
            page(lanes)
        for word in said:
            self.assertIn(word, str(raised.exception))

    def test_a_bar_under_it_in_a_lane_it_spans_is_refused_and_both_are_named(self):
        self.refused([lane("One", [bar(0, 3, "shared", spans=2)]), lane("Two", [bar(2, 4, "under")])],
                     "lane 2, bar 1", "lane 1, bar 1", "spans lane 2")

    def test_a_bar_beside_it_in_its_own_lane_over_the_same_periods_is_refused(self):
        self.refused([lane("One", [bar(0, 3, "shared", spans=2), bar(2, 4, "mine")]), lane("Two")],
                     "lane 1, bar 2", "lane 1, bar 1", "spans lane 1")

    def test_a_mark_on_it_is_refused_and_so_is_one_on_either_end(self):
        for at in (0, 1.5, 3):
            with self.subTest(at):
                self.refused([lane("One", [bar(0, 3, "shared", spans=2)]),
                              lane("Two", marks=[{"at": at, "text": "gate"}])],
                             "lane 2, mark 1", "lane 1, bar 1")

    def test_two_of_them_cannot_span_one_lane_over_the_same_periods(self):
        self.refused([lane("One", [bar(0, 3, "upper", spans=2)]),
                      lane("Two", [bar(2, 5, "lower", spans=2)]), lane("Three")],
                     "lane 1, bar 1", "lane 2, bar 1", "both span lane 2")

    def test_two_of_them_one_after_the_other_are_taken(self):
        built = page([lane("One", [bar(0, 3, "upper", spans=2)]),
                      lane("Two", [bar(3, 6, "lower", spans=2)]), lane("Three")])
        self.assertLess(named(built, "upper").rect.top, named(built, "lower").rect.top)
        self.assertLess(named(built, "upper").rect.bottom, named(built, "lower").rect.bottom)

    def test_it_cannot_span_past_the_last_lane(self):
        self.refused([lane("One", [bar(0, 6, "a")]), lane("Two", [bar(0, 3, "shared", spans=2)])],
                     "lane 2, bar 1", "spans 2 lanes", "only 1 is left")

    def test_how_many_lanes_is_a_whole_number_of_one_or_more(self):
        for value in (0, -1, 1.5, "2", True, None):
            with self.subTest(value):
                self.refused([lane("One", [bar(0, 3, "shared", spans=value)]), lane("Two")],
                             "lane 1, bar 1", "`spans`", "whole number of lanes")

    def test_a_name_beside_a_neighbour_goes_to_the_side_it_does_not_stand_on(self):
        """右が先。ただし右へ出すとまたがる棒の上に乗るなら、左へ出す。"""
        text = "a name that goes beside it"
        lanes = [lane("One", [bar(2.6, 2.9, text), bar(3.5, 4.5, "shared", spans=2)]),
                 lane("Two", [bar(0, 3, "under")])]
        name = named(page(lanes), text)
        self.assertIsInstance(name, Text)
        self.assertLess(name.rect.right, named(page(lanes), "shared").rect.left)
        self.assertEqual("right", name.align)
        # 同じ棒は、またがる棒が無ければ右へ出る (= 左へ出たのは、またがる棒を避けたから)
        lanes[0]["bars"][1] = bar(3.5, 4.5, "shared")
        self.assertEqual("left", named(page(lanes), text).align)

    def test_a_name_with_it_on_both_sides_stays_inside_its_bar(self):
        built = page([lane("One", [bar(0, 2.8, "before", spans=2), bar(2.9, 3.1, self.LONG),
                                   bar(3.2, 6, "after", spans=2)]), lane("Two")])
        self.assertIsInstance(named(built, self.LONG), Bar)

    def test_a_marks_name_with_it_on_both_sides_stops_the_page_and_says_why(self):
        with self.assertRaises(PageFullError) as raised:
            page([lane("One", [bar(0, 2.8, "before", spans=2), bar(3.2, 6, "after", spans=2)],
                       [{"at": 3, "text": self.LONG}]), lane("Two")])
        self.assertIn("lane 1, mark 1", str(raised.exception))
        self.assertIn("spanning", str(raised.exception))

    def test_a_name_too_long_for_the_lanes_makes_them_taller_instead_of_being_cut(self):
        text = " ".join(["a long name"] * 12)
        built = page([lane("One", [bar(0, 1, text, spans=2), bar(1, 6, "ours")]),
                      lane("Two", [bar(1, 6, "theirs")])])
        long = named(built, text)
        inner = long.rect.width - 2 * DEFAULT.spacing.bar_pad_x
        needed = DEFAULT.wrapped_height(text, inner, long.size) + 2 * DEFAULT.spacing.bar_pad_y
        self.assertGreaterEqual(long.rect.height, needed)
        self.assertGreater(named(built, "ours").rect.height, named(built, "theirs").rect.height)
        self.assertLessEqual(long.rect.bottom, built.theme.frame().bottom)

    def test_everything_stays_inside_the_frame_and_passes_the_checks_once_baked(self):
        built = page([lane("One", [bar(0, 3, "shared", spans=3), bar(3, 5, "a"), bar(4, 6, "b")]),
                      lane("Two", [bar(3, 6, "c")], [{"at": 4, "text": "gate"}]),
                      lane("Three", [bar(3, 4, "d"), bar(4, 6, "last", spans=2, tentative=True)]),
                      lane("Four", [bar(0, 4, "e")])],
                     phases=[{"from": 0, "to": 6, "label": "this half"}],
                     milestones=[{"at": 5.5, "text": "review"}])
        frame = built.theme.frame()
        for element in built.build():
            self.assertTrue(frame.contains(element.rect), element.kind)
        with tempfile.TemporaryDirectory() as tmp:
            deck = new_deck()
            add_page(deck, built.build())
            baked = save(deck, Path(tmp) / "across.pptx")
            self.assertEqual([], checks.run_all(baked))
            with zipfile.ZipFile(baked) as archive:
                xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        self.assertEqual(1, xml.count(">shared<"))
        self.assertEqual({"rect"}, set(re.findall(r'<a:prstGeom prst="(\w+)"', xml)))


class DatesThatCutAcross(unittest.TestCase):
    """A milestone is a line through every lane, with its name flying from the top of it."""

    LANES = [lane("One", [bar(0, 6, "a")]), lane("Two", [bar(0, 6, "b")])]

    def test_the_line_runs_from_its_name_to_the_bottom_of_the_last_lane(self):
        built = page(self.LANES, milestones=[{"at": 3, "text": "review"}])
        line, name = of(built, "line")[0].rect, named(built, "review").rect
        self.assertEqual(line.top, name.top)
        self.assertGreater(line.bottom, named(built, "b").rect.bottom)
        self.assertAlmostEqual(line.left + line.width // 2, columns(built)[3], delta=1)
        self.assertEqual(name.left - line.right, DEFAULT.spacing.bar_pad_x - DEFAULT.spacing.hairline)

    def test_the_line_is_laid_under_the_bars(self):
        """上に引くと、棒の名前を貫く (= 焼いて初めて出た)。"""
        kinds = [e.kind for e in page(self.LANES, milestones=[{"at": 3, "text": "x"}]).build()]
        self.assertLess(kinds.index("line"), kinds.index("bar"))

    #: 名前が線に貫かれうる並び ― 近い日付が続く / 右端で名前が左へ出て手前の線に掛かる /
    #: 長い名前の下を、後から来た線がくぐる
    CROWDED = {
        "close together": [{"at": 3, "text": "a long review name"}, {"at": 3.3, "text": "budget lock"},
                           {"at": 3.6, "text": "another one close by"}, {"at": 5.9, "text": "launch"}],
        "a name flying left at the edge": [{"at": 5, "text": "early"},
                                           {"at": 5.95, "text": "a long name near the edge"}],
        "a late line under a long name": [
            {"at": 1, "text": "a long name that runs on and on a while"},
            {"at": 2, "text": "another long name that runs far across the page"},
            {"at": 3.5, "text": "short"}],
    }

    def test_no_line_cuts_through_another_milestones_name(self):
        for case, stones in self.CROWDED.items():
            with self.subTest(case):
                built = page(self.LANES, milestones=stones)
                names = [named(built, stone["text"]).rect for stone in stones]
                for line in of(built, "line"):
                    for name in names:
                        self.assertFalse(line.rect.overlaps(name), "a line runs through a name")
                for index, one in enumerate(names):
                    for other in names[index + 1:]:
                        self.assertFalse(one.overlaps(other), "two names share a place")

    def test_the_crowded_cases_are_the_cases_they_say_they_are(self):
        """見本の名前が短くなると、規則の枝を通らないまま緑になる (= 幅の数え方を変えた時に起きた)。"""
        def row(built, text):
            tops = sorted({e.rect.top for e in built.build() if e.kind == "label" and e.bold})
            return tops.index(named(built, text).rect.top)

        # 1 つめの名前は、右へ出すと後ろの線に貫かれるので左へ出る。残りは段を分ける
        built = page(self.LANES, milestones=self.CROWDED["close together"])
        self.assertEqual([0, 0, 1], [row(built, s["text"]) for s in self.CROWDED["close together"][:3]])
        self.assertEqual("right", named(built, "a long review name").align)
        built = page(self.LANES, milestones=self.CROWDED["a name flying left at the edge"])
        self.assertEqual((1, 0), (row(built, "early"), row(built, "a long name near the edge")))
        built = page(self.LANES, milestones=self.CROWDED["a late line under a long name"])
        self.assertEqual(2, row(built, "short"), "the short name had room on the first row but for the line")

    def test_a_name_with_no_room_on_the_right_flies_left_from_its_line(self):
        built = page(self.LANES, milestones=self.CROWDED["a name flying left at the edge"])
        name = named(built, "a long name near the edge")
        line = max(of(built, "line"), key=lambda e: e.rect.left).rect
        self.assertEqual("right", name.align)
        self.assertLess(name.rect.right, line.left)
        self.assertEqual(name.rect.top, line.top)

    def test_a_name_beside_a_bar_moves_to_the_side_no_line_crosses(self):
        text = "a name far too long for a bar this short"
        free = page([lane(bars=[bar(3, 3.2, text)])])
        crossed = page([lane(bars=[bar(3, 3.2, text)])], milestones=[{"at": 4, "text": "x"}])
        self.assertEqual("left", named(free, text).align)
        self.assertEqual("right", named(crossed, text).align)

    def test_a_tentative_bar_hides_the_line_behind_it(self):
        built = page([lane(bars=[bar(2, 4, "maybe", tentative=True)])],
                     milestones=[{"at": 3, "text": "x"}])
        self.assertTrue(named(built, "maybe").fill, "an unfilled bar lets the line through its name")


class WhatItRefuses(unittest.TestCase):
    """A declaration nobody can draw stops here, and says what to change."""

    def refused(self, lanes, message, **extra):
        with self.assertRaises(PageTypeError) as raised:
            page(lanes, **extra)
        self.assertIn(message, str(raised.exception))

    def test_a_position_outside_the_periods(self):
        self.refused([lane(bars=[bar(0, 7)])], "outside the 6 periods")
        self.refused([lane(marks=[{"at": -1, "text": "x"}])], "outside the 6 periods")

    def test_a_position_that_is_not_a_number(self):
        self.refused([lane(bars=[bar("Jan", 2)])], "period number")
        self.refused([lane(bars=[bar(True, 2)])], "period number")
        self.refused([lane(bars=[{"to": 2, "text": "x"}])], "period number")

    def test_a_bar_that_ends_before_it_starts(self):
        self.refused([lane(bars=[bar(3, 3)])], "end after it starts")
        self.refused([lane(bars=[bar(4, 2)])], "end after it starts")

    def test_a_key_nobody_reads_at_every_level(self):
        self.refused([{"name": "x", "bars": [bar(0, 1)], "colour": "red"}], "does not take colour")
        self.refused([lane(bars=[bar(0, 1, left=3)])], "does not take left")
        self.refused([lane(marks=[{"at": 1, "text": "x", "x": 3}])], "does not take x")
        self.refused([lane(bars=[bar(0, 1)])], "does not take width",
                     phases=[{"from": 0, "to": 1, "label": "p", "width": 3}])
        self.refused([lane(bars=[bar(0, 1)])], "does not take top",
                     milestones=[{"at": 1, "text": "m", "top": 3}])

    def test_milestones_written_under_a_lane_say_where_they_belong(self):
        """TOML では `[[pages.lanes]]` の下に書いたキーはそのレーンの物になる。"""
        self.refused([{"name": "x", "bars": [bar(0, 1)], "milestones": [{"at": 1, "text": "m"}]}],
                     "above the first `[[pages.lanes]]`")

    def test_no_periods_no_lanes_or_a_lane_with_nothing_on_it(self):
        self.refused([lane(bars=[bar(0, 1)])], "`periods` is empty", periods=[])
        self.refused([], "`lanes` is empty")
        self.refused([lane()], "neither bars nor marks")

    def test_phases_that_overlap(self):
        self.refused([lane(bars=[bar(0, 1)])], "starts before the one before it ends",
                     phases=[{"from": 0, "to": 3, "label": "a"}, {"from": 2, "to": 5, "label": "b"}])


class WhatDoesNotFitIsNotShrunk(unittest.TestCase):

    def test_more_lanes_than_the_page_holds(self):
        with self.assertRaises(PageFullError) as raised:
            page([lane(f"Lane {n}", [bar(0, 6, "work")]) for n in range(30)])
        self.assertIn("will not shrink", str(raised.exception))

    def test_a_period_name_wider_than_its_column(self):
        with self.assertRaises(PageFullError):
            page([lane(bars=[bar(0, 1)])], periods=["a period with a very long name"] * 12)

    def test_a_few_bars_do_not_swell_to_fill_the_page(self):
        """余った高さを配り切ると、棒 1 本の線表が頁いっぱいの帯になる。"""
        built = page([lane(bars=[bar(0, 6, "only")])])
        natural = DEFAULT.line_height(DEFAULT.type.plan[0]) + 2 * DEFAULT.spacing.bar_pad_y
        self.assertLessEqual(named(built, "only").rect.height, round(natural * (1 + ROOMY)))
        self.assertGreater(named(built, "only").rect.height, natural)
        bottom = max(e.rect.bottom for e in built.build())
        self.assertLess(bottom, built.theme.frame().bottom - natural, "one bar filled the page")

    def test_spare_height_goes_to_the_bars_so_the_page_is_used_to_its_foot(self):
        """頁の下に空きを残さない (= 実物のデッキで、下半分が空いた工程表が差し戻された)。"""
        lanes = [lane(f"Lane {n}", [bar(0, 6, "work")]) for n in range(6)]
        built = page(lanes)
        natural = DEFAULT.line_height(DEFAULT.type.plan[0]) + 2 * DEFAULT.spacing.bar_pad_y
        heights = {e.rect.height for e in built.build() if e.kind == "bar"}
        self.assertEqual(1, len(heights))
        self.assertGreater(heights.pop(), natural)
        bottom = max(e.rect.bottom for e in built.build())
        self.assertLessEqual(built.theme.frame().bottom - bottom, len(lanes))

    def test_a_full_page_keeps_its_bars_at_their_natural_height(self):
        lanes = [lane(f"Lane {n}", [bar(0, 6, "work")]) for n in range(12)]
        built = page(lanes)
        size = named(built, "work").size
        natural = DEFAULT.line_height(size) + 2 * DEFAULT.spacing.bar_pad_y
        heights = {e.rect.height for e in built.build() if e.kind == "bar"}
        self.assertLess(max(heights), round(natural * 1.2))


class TheSizeOfTheWords(unittest.TestCase):
    """The words are set at the largest size of `plan` the page holds, and never below the smallest."""

    def sizes(self, built):
        return {e.size for e in built.build() if e.kind in ("bar", "span", "period", "label")}

    def test_a_page_with_room_takes_the_larger_size(self):
        built = page([lane(bars=[bar(0, 6, "work")], marks=[{"at": 1, "text": "mark"}])],
                     phases=[{"from": 0, "to": 6, "label": "phase"}],
                     milestones=[{"at": 3, "text": "date"}])
        self.assertEqual({DEFAULT.type.plan[0]}, self.sizes(built))

    def test_a_crowded_page_takes_the_smaller_one_for_every_word_alike(self):
        lanes = [lane(f"Lane {n}", [bar(0, 3, "work"), bar(4, 4.1, "a name beside its bar")])
                 for n in range(12)]
        built = page(lanes)
        self.assertEqual({DEFAULT.type.plan[-1]}, self.sizes(built))
        one_fewer = page(lanes[:10])
        self.assertEqual({DEFAULT.type.plan[0]}, self.sizes(one_fewer), "ten lanes have room for the larger size")

    def test_the_smallest_size_is_the_body_size_not_the_footnote_size(self):
        self.assertEqual(DEFAULT.type.body, min(DEFAULT.type.plan))
        self.assertGreater(min(DEFAULT.type.plan), DEFAULT.type.caption)
        self.assertEqual(sorted(DEFAULT.type.plan, reverse=True), list(DEFAULT.type.plan))

    def test_a_page_that_does_not_fit_at_the_smallest_size_stops_and_says_the_size(self):
        with self.assertRaises(PageFullError) as raised:
            page([lane(f"Lane {n}", [bar(0, 6, "work")]) for n in range(30)])
        self.assertIn(f"{min(DEFAULT.type.plan):g}pt", str(raised.exception))

    def test_a_lanes_name_is_set_at_its_own_size_in_bold(self):
        built = page([lane("Operations", [bar(0, 6, "work")])])
        name = named(built, "Operations")
        self.assertEqual(("lane", DEFAULT.type.stage, True), (name.kind, name.size, name.bold))

    def test_a_long_lane_name_folds_and_the_lane_makes_room_for_it(self):
        long_name = "the team that looks after everything nobody else has the time to look after"
        # 頁を詰める (= 余った高さが配られると、名前のぶんを取らなくても枠が足りてしまう)
        others = [lane(f"Lane {n}", [bar(0, 6, f"work {n}")]) for n in range(7)]
        built = page([lane(long_name, [bar(0, 6, "work")]), lane("Next", [bar(0, 6, "more")]), *others])
        self.assertGreaterEqual(named(built, "Next").rect.top, named(built, long_name).rect.bottom)
        room = named(built, long_name).rect
        self.assertLessEqual(room.width, built.theme.frame().width // 4)
        needed = DEFAULT.wrapped_height(long_name, room.width, DEFAULT.type.stage, bold=True)
        self.assertGreater(needed, DEFAULT.line_height(DEFAULT.type.stage), "the name does not fold")
        self.assertGreaterEqual(room.height, needed, "the lane is too short for its own name")


class WhereAPeriodIs(unittest.TestCase):
    """A reader has to see at a glance which period a bar falls in."""

    def built(self, **extra):
        return page([lane("One", [bar(0, 6, "a")]), lane("Two", [bar(1, 2, "b")])], **extra)

    def test_every_period_has_a_heading_cell_like_a_tables_heading_row(self):
        cells = of(self.built(), "period")
        self.assertEqual(MONTHS, [cell.text for cell in cells])
        for cell in cells:
            self.assertIsInstance(cell, Bar)
            self.assertTrue(cell.bold)
            self.assertEqual((DEFAULT.palette.accent, DEFAULT.palette.paper), (cell.fill, cell.colour))
        for one, other in zip(cells, cells[1:]):
            self.assertEqual(one.rect.top, other.rect.top)
            self.assertGreater(other.rect.left, one.rect.right)

    def test_a_phase_sits_above_the_period_cells_with_its_ends_on_theirs(self):
        built = self.built(phases=[{"from": 1, "to": 4, "label": "phase"}])
        cells, phase = of(built, "period"), named(built, "phase").rect
        self.assertEqual(phase.left, cells[1].rect.left)
        self.assertEqual(phase.right, cells[3].rect.right)
        self.assertLess(phase.bottom, cells[0].rect.top)

    def test_every_other_column_is_washed_from_the_headings_down_to_the_last_lane(self):
        built = self.built()
        stripes, cells = of(built, "stripe"), of(built, "period")
        self.assertEqual(len(MONTHS) // 2, len(stripes))
        self.assertEqual({DEFAULT.palette.wash}, {stripe.colour for stripe in stripes})
        edges = columns(built)
        self.assertEqual([edges[1], edges[3], edges[5]], [stripe.rect.left for stripe in stripes])
        for stripe in stripes:
            self.assertEqual(cells[0].rect.bottom, stripe.rect.top)
            self.assertEqual(named(built, "b").rect.bottom + DEFAULT.spacing.row_gap, stripe.rect.bottom)

    def test_the_wash_and_the_lines_are_laid_under_the_bars(self):
        kinds = [e.kind for e in self.built(milestones=[{"at": 3, "text": "x"}]).build()]
        first_bar = kinds.index("bar")
        for under in ("stripe", "edge", "rule", "line"):
            self.assertLess(max(i for i, kind in enumerate(kinds) if kind == under), first_bar, under)

    def test_a_line_stands_where_each_period_ends_darker_than_the_lines_between_lanes(self):
        built = self.built()
        edges = of(built, "edge")
        self.assertEqual(len(MONTHS) + 1, len(edges))
        self.assertEqual({DEFAULT.palette.muted}, {edge.colour for edge in edges})
        self.assertNotEqual(DEFAULT.palette.muted, DEFAULT.palette.rule)
        self.assertEqual({DEFAULT.palette.rule}, {e.colour for e in built.build()
                                                  if e.kind == "rule" and e.rect.width > e.rect.height})
        for edge, column in zip(edges, columns(built)):
            self.assertAlmostEqual(edge.rect.left + edge.rect.width // 2, column, delta=1)

    def test_the_projects_own_wash_reaches_the_columns(self):
        theme = theme_from({"palette": {"wash": "eeeeee"}})
        built = types.build({"type": "timeline", "title": "A plan", "periods": MONTHS,
                             "lanes": [lane(bars=[bar(0, 6)])]}, lambda name: Path(name),
                            lambda path: 1.0, theme)
        self.assertEqual({"EEEEEE"}, {stripe.colour for stripe in of(built, "stripe")})

    def test_the_wash_is_never_read_from_a_templates_theme(self):
        self.assertNotIn("wash", [name for _slot, name in SLOTS])


class OnceBaked(unittest.TestCase):
    """What the writer puts in the file: shapes PowerPoint has no reason to repair."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        built = page([lane("Core", [bar(0, 2, "build"), bar(2, 2.3, "a name that goes beside it"),
                                    bar(3, 5, "maybe", tentative=True)],
                           [{"at": 2, "text": "release"}])],
                     phases=[{"from": 0, "to": 6, "label": "this half"}],
                     milestones=[{"at": 5.5, "text": "review"}])
        deck = new_deck()
        add_page(deck, built.build())
        cls.deck = save(deck, Path(cls.tmp.name) / "timeline.pptx")
        with zipfile.ZipFile(cls.deck) as archive:
            cls.xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        cls.shapes = re.findall(r"<p:sp>.*?</p:sp>", cls.xml, re.S)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def shape(self, text):
        return next(s for s in self.shapes if f">{text}<" in s)

    def test_the_only_preset_shape_is_the_rectangle(self):
        self.assertEqual({"rect"}, set(re.findall(r'<a:prstGeom prst="(\w+)"', self.xml)))

    def test_the_diamond_is_a_square_turned_45_degrees(self):
        turned = [s for s in self.shapes if 'rot="2700000"' in s]
        self.assertEqual(1, len(turned))
        width, height = re.search(r'<a:ext cx="(\d+)" cy="(\d+)"', turned[0]).groups()
        self.assertEqual(width, height)

    def test_a_tentative_bar_has_a_dashed_line_and_a_light_bar_a_solid_one(self):
        self.assertIn('<a:prstDash val="dash"/>', self.shape("maybe"))
        edge = DEFAULT.palette.edge(DEFAULT.palette.box)
        self.assertNotIn("prstDash", self.shape("build"))
        self.assertRegex(self.shape("build"),
                         rf'<a:ln w="{DEFAULT.spacing.hairline}">\s*<a:solidFill>\s*'
                         rf'<a:srgbClr val="{edge}"/>')

    def test_a_shape_with_no_edge_says_so(self):
        """線を書かない図形はテーマの線を引く。持たないなら「無い」と書く。"""
        for text in ("Jan", "this half"):
            self.assertRegex(self.shape(text), r"<a:ln[^>]*>\s*<a:noFill/>", text)

    def test_no_shape_borrows_the_themes_shape_style(self):
        """参照が残ると、LibreOffice で焼いた絵にだけ影が付く。"""
        self.assertNotIn("<p:style>", self.xml)

    def test_a_bar_and_its_name_are_one_shape(self):
        self.assertIn("<a:solidFill>", self.shape("build").split("<p:txBody>")[0])

    def test_a_name_beside_something_is_not_wrapped(self):
        for text in ("release", "review", "a name that goes beside it"):
            self.assertIn('wrap="none"', self.shape(text), text)
        self.assertIn('wrap="square"', self.shape("build"))

    def test_a_period_cell_is_a_filled_shape_with_its_name_in_bold(self):
        cell = self.shape("Jan")
        self.assertIn(f'<a:srgbClr val="{DEFAULT.palette.accent}"/>', cell.split("<p:txBody>")[0])
        self.assertRegex(cell, r'<a:rPr[^>]*b="1"')

    def test_no_word_is_written_below_the_body_size(self):
        sizes = {int(size) for size in re.findall(r'sz="(\d+)"', self.xml)}
        self.assertGreaterEqual(min(sizes), int(DEFAULT.type.body * 100))

    def test_no_run_is_empty(self):
        self.assertNotRegex(self.xml, r"<a:t\s*/>|<a:t></a:t>")

    def test_the_checks_find_nothing(self):
        self.assertEqual([], checks.run_all(self.deck))

    def test_python_pptx_opens_it_again(self):
        from pptx import Presentation
        self.assertEqual(1, len(Presentation(str(self.deck)).slides))


class TheWidthOfOneLine(unittest.TestCase):
    """Whether a name fits on one line is measured, not assumed."""

    def test_a_full_width_character_is_one_em_and_a_narrow_one_less(self):
        em = DEFAULT.width("全", 10)
        self.assertEqual(em * 3, DEFAULT.width("全角字", 10))
        self.assertLess(DEFAULT.width("abc", 10), em * 3)
        self.assertGreater(DEFAULT.width("abc", 10), em * 3 // 2,
                           "narrower than what the overlap check assumes leaves no slack")

    def test_bold_runs_wider(self):
        self.assertGreater(DEFAULT.width("name", 10, bold=True), DEFAULT.width("name", 10))

    def test_only_the_longest_line_counts(self):
        self.assertEqual(DEFAULT.width("abcdef", 10), DEFAULT.width("ab\nabcdef\nabc", 10))


if __name__ == "__main__":
    unittest.main()
