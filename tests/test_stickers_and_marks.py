"""A sticker on a page, and a status mark in a table's cell.

⚠ **頁の性格 (= 暫定・イメージ) と、表の中の状態 (= どこまで済んだか、良いか悪いか) は、言葉で
書くと頁ごとに書き方が割れる。**札は題の帯の右端に 1 つ、印は決まった 8 つの字 ― どちらも頁が
自分で形を決める口を持たない。

札は題の隣に立つので、題の字が札の下へ回り込まないこと (= 題はそのぶん狭い幅で折れ、行もその幅で
数える) を見張る。印は字のまま表に入るので、案件の書体に字形が無くても同じ字が出ること (= 印の
字形を持つ書体を名指しする) を見張る。
"""

from __future__ import annotations

import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx import Presentation  # noqa: E402

from pptx_agent_maker import DEFAULT, Page, PageFullError, PageTypeError, checks  # noqa: E402
from pptx_agent_maker.checks import declared  # noqa: E402
from pptx_agent_maker.layout import types  # noqa: E402
from pptx_agent_maker.layout.base.tokens import MARK_FACE, theme_from  # noqa: E402
from pptx_agent_maker.layout.parts.elements import Bar, Mark, Table  # noqa: E402
from pptx_agent_maker.layout.types.core.read import HARVEY, MARKS, read_cell  # noqa: E402
from pptx_agent_maker.project.files.manifest import Manifest  # noqa: E402
from pptx_agent_maker.write import add_page, new_deck, save  # noqa: E402
from support.pages import build  # noqa: E402
from test_types import MINIMAL  # noqa: E402

S, T, P = DEFAULT.spacing, DEFAULT.type, DEFAULT.palette
ONE = "題は結論の文で書く"
#: 24pt で 1 行に入るのは全角 36 字。札を持つ頁は、そのぶん少ない
ALMOST_FULL = "字" * 34
FOLDED = "題は結論の文で書くので長くなりやすい。" * 3


def placed(page: Page, kind: str):
    return next(e for e in page.elements if e.kind == kind)


class ASticker(unittest.TestCase):
    def test_a_page_without_one_is_laid_out_as_it_always_was(self) -> None:
        for kicker in ("", "01 | 背景"):
            with self.subTest(kicker=kicker):
                self.assertEqual(Page(ONE, kicker=kicker).elements, Page(ONE, kicker=kicker, sticker="").elements)
                self.assertNotIn("sticker", {e.kind for e in Page(ONE, kicker=kicker).elements})

    def test_it_stands_at_the_right_end_of_the_titles_band_as_wide_as_its_words(self) -> None:
        for kicker in ("", "01 | 背景"):
            with self.subTest(kicker=kicker):
                page = Page(ONE, kicker=kicker, sticker="暫定")
                tag = placed(page, "sticker")
                self.assertIsInstance(tag, Bar)
                self.assertEqual("暫定", tag.text)
                self.assertEqual(DEFAULT.frame().right, tag.rect.right)
                self.assertEqual(DEFAULT.width("暫定", T.body) + 2 * S.pad, tag.rect.width)
                self.assertEqual(DEFAULT.line_height(T.body) + 2 * S.bar_pad_y, tag.rect.height)
                self.assertTrue(DEFAULT.frame().contains(tag.rect))
                self.assertLessEqual(tag.rect.bottom, placed(page, "rule").rect.top)

    def test_it_is_grey_words_in_a_thin_grey_edge_on_no_ground(self) -> None:
        tag = placed(Page(ONE, sticker="Draft"), "sticker")
        self.assertEqual((P.muted, P.muted, "", T.body), (tag.colour, tag.outline, tag.fill, tag.size))
        self.assertFalse(tag.heavy or tag.dashed or tag.bold)

    def test_the_titles_box_ends_a_gap_before_it(self) -> None:
        for kicker in ("", "01 | 背景"):
            for title in (ONE, FOLDED):
                with self.subTest(kicker=kicker, title=title[:8]):
                    page = Page(title, kicker=kicker, sticker="イメージ")
                    words, tag = placed(page, "title").rect, placed(page, "sticker").rect
                    self.assertEqual(tag.left - S.gap_m, words.right)
                    self.assertFalse(words.overlaps(tag))
                    self.assertEqual(Page(title, kicker=kicker).body.left, words.left)

    def test_it_sits_at_the_middle_of_the_titles_box_however_many_lines_the_title_folds_to(self) -> None:
        """題は枠の真ん中に寄る。札も同じ真ん中に置けば、題が何行に折れても 2 つは揃う。"""
        for kicker in ("", "01 | 背景"):
            for title in (ONE, FOLDED):
                with self.subTest(kicker=kicker, title=title[:8]):
                    page = Page(title, kicker=kicker, sticker="暫定")
                    words, tag = placed(page, "title").rect, placed(page, "sticker").rect
                    self.assertAlmostEqual(words.top + words.height / 2, tag.top + tag.height / 2, delta=1)

    def test_a_title_folds_sooner_beside_a_sticker_and_the_band_grows_for_it(self) -> None:
        self.assertEqual(1, DEFAULT.title_lines(ALMOST_FULL))
        self.assertEqual(2, DEFAULT.title_lines(ALMOST_FULL, "イメージ"))
        line = DEFAULT.line_height(T.title)
        self.assertEqual(Page(ALMOST_FULL).body.top + line, Page(ALMOST_FULL, sticker="イメージ").body.top)

    def test_the_room_a_sticker_takes_is_its_width_and_the_gap_beside_it(self) -> None:
        em = T.title * 12700
        room = DEFAULT.frame().width - 2 * S.text_inset - DEFAULT.sticker_width("暫定") - S.gap_m
        fits = int(room / em)
        self.assertEqual(1, DEFAULT.title_lines("字" * fits, "暫定"))
        self.assertEqual(2, DEFAULT.title_lines("字" * (fits + 1), "暫定"))

    def test_a_title_is_counted_at_the_width_of_the_box_it_is_placed_in(self) -> None:
        """⚠ 行を数える幅と、題を置く枠の幅が食い違うと、帯の高さが題の行と合わない。札の幅だけでなく、
        札との間の空きも、題の枠から引かれている。"""
        base, line = Page(ONE, sticker="暫定").body.top, DEFAULT.line_height(T.title)
        counted = set()
        for full in range(30, 37):
            for half in range(4):
                title = "字" * full + " " + "i" * half
                page = Page(title, sticker="暫定")
                room = placed(page, "title").rect.width - 2 * S.text_inset
                lines = DEFAULT.wraps(title, room, T.title, bold=True)
                counted.add(lines)
                self.assertEqual(base + (lines - 1) * line, page.body.top, title)
        self.assertEqual({1, 2}, counted, "the titles tried do not cross the edge this test is about")

    def test_a_sticker_that_leaves_the_title_no_room_stops_the_page(self) -> None:
        with self.assertRaises(PageFullError) as stopped:
            Page("Unbreakable" * 3, sticker="この頁はまだ暫定の内容で、数字は来週の集計で差し替える予定です。" * 2)
        self.assertIn("sticker", str(stopped.exception))

    def test_every_type_takes_one(self) -> None:
        for name, data in MINIMAL.items():
            with self.subTest(name):
                page = build({"type": name, "title": "だい", "sticker": " 暫定 ", **data})
                self.assertEqual("暫定", placed(page, "sticker").text)

    def test_a_page_carries_one_sticker_written_as_its_words(self) -> None:
        for sticker in (["暫定", "イメージ"], 3, True, {"text": "暫定"}):
            with self.subTest(sticker=sticker), self.assertRaises(PageTypeError) as refused:
                build({"type": "cards", "title": "だい", "cards": [["a", "b"]], "sticker": sticker})
            self.assertIn("one sticker", str(refused.exception))

    def test_the_check_of_a_long_title_counts_beside_the_sticker(self) -> None:
        """⚠ 検査が札を知らずに数えると、頁は 3 行ぶんの帯を取ったのに検査は 2 行と言う。"""
        title = "字" * 70
        self.assertEqual((2, 3), (DEFAULT.title_lines(title), DEFAULT.title_lines(title, "イメージ")))
        with tempfile.TemporaryDirectory() as tmp:
            head = 'specimen = "specimen.pptx"\nout = "deck.pptx"\n'
            page = f'[[pages]]\nkind = "declare"\ntype = "cards"\ntitle = "{title}"\ncards = [["a", "b"]]\n'
            plain, tagged = Path(tmp) / "plain.toml", Path(tmp) / "tagged.toml"
            plain.write_text(head + page, encoding="utf-8")
            tagged.write_text(head + page + 'sticker = " イメージ "\n', encoding="utf-8")
            self.assertEqual([], declared.long_title.run(Manifest.load(plain), DEFAULT, {}))
            found = declared.long_title.run(Manifest.load(tagged), DEFAULT, {})
            self.assertEqual([("long_title", 1)], [(f.check, f.page) for f in found])

    def test_once_baked_the_checks_find_nothing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            deck = new_deck()
            for title in (ONE, FOLDED):
                page = build({"type": "cards", "title": title, "kicker": "01 | 背景", "sticker": "イメージ",
                              "cards": [["a", "b"]]})
                add_page(deck, page.build())
            saved = save(deck, Path(tmp) / "tagged.pptx")
            self.assertEqual([], checks.run_all(saved))
            self.assertEqual(2, len(Presentation(str(saved)).slides))


class AMarkInACell(unittest.TestCase):
    TABLE = [["項目", "進み", "状態"], ["設計", {"harvey": 4}, {"mark": "ok"}],
             ["試作", {"harvey": 2}, {"mark": "partial"}], ["評価", {"harvey": 0}, {"mark": "ng"}]]

    def test_a_harvey_ball_is_filled_by_quarters(self) -> None:
        self.assertEqual(["○", "◔", "◑", "◕", "●"],
                         [read_cell({"harvey": level}, "a cell") for level in range(5)])
        self.assertEqual(5, len(set(HARVEY)))

    def test_good_partly_and_bad_each_have_a_mark_of_their_own(self) -> None:
        self.assertEqual({"ok": "✓", "partial": "△", "ng": "×"},
                         {name: read_cell({"mark": name}, "a cell") for name in ("ok", "partial", "ng")})
        self.assertEqual(MARKS.keys(), {"ok", "partial", "ng"})
        self.assertFalse(set(MARKS.values()) & set(HARVEY), "a mark and a Harvey ball share a character")

    def test_a_mark_is_told_from_the_same_character_written_as_words(self) -> None:
        self.assertIsInstance(read_cell({"harvey": 0}, "a cell"), Mark)
        self.assertNotIsInstance(read_cell("○", "a cell"), Mark)
        self.assertEqual("12", read_cell(12, "a cell"))

    def test_anything_else_written_as_a_table_of_keys_is_refused(self) -> None:
        for cell in ({"harvey": 5}, {"harvey": -1}, {"harvey": 2.0}, {"harvey": True}, {"harvey": "2"},
                     {"mark": "good"}, {"mark": 1}, {"harvey": 1, "mark": "ok"}, {"text": "a"}, {}):
            with self.subTest(cell=cell), self.assertRaises(PageTypeError) as refused:
                read_cell(cell, "table: row 2, column 3")
            self.assertIn("table: row 2, column 3", str(refused.exception))

    def test_every_table_takes_marks_and_knows_where_they_are(self) -> None:
        wanted = frozenset({(1, 1), (1, 2), (2, 1), (2, 2), (3, 1), (3, 2)})
        pages = {
            "under a body": build({"type": "figure", "title": "だい", "figure": "dot.png", "table": self.TABLE}),
            "a board": build({"type": "board", "title": "だい", "table": self.TABLE}),
            "a cell of a compose": build({"type": "compose", "title": "だい",
                                          "rows": [{"cells": [{"table": self.TABLE}]}]}),
        }
        for where, page in pages.items():
            with self.subTest(where):
                table = next(e for e in page.build() if isinstance(e, Table))
                self.assertEqual(wanted, table.marks)
                self.assertEqual(("設計", "●", "✓"), table.rows[1])

    def test_a_mark_written_wrong_says_which_cell(self) -> None:
        bad = [["項目", "進み"], ["設計", {"harvey": 9}]]
        with self.assertRaises(PageTypeError) as refused:
            build({"type": "board", "title": "だい", "table": bad})
        self.assertIn("table: row 2, column 2", str(refused.exception))
        with self.assertRaises(PageTypeError) as refused:
            build({"type": "compose", "title": "だい", "rows": [{"cells": [{"card": ["a", "b"]}, {"table": bad}]}]})
        self.assertIn("compose: row 1, cell 2: table: row 2, column 2", str(refused.exception))

    def test_a_row_of_marks_is_as_tall_as_a_row_of_words(self) -> None:
        page = build({"type": "board", "title": "だい", "table": self.TABLE})
        table = next(e for e in page.build() if isinstance(e, Table))
        self.assertEqual(len(self.TABLE) * DEFAULT.table_row_height(), table.rect.height)

    def _cells(self, theme=DEFAULT) -> list[str]:
        """The XML of every cell of the table, once the page is written out."""
        page = build({"type": "board", "title": "だい", "table": self.TABLE}, theme)
        with tempfile.TemporaryDirectory() as tmp:
            deck = new_deck(theme)
            add_page(deck, page.build(), theme)
            saved = save(deck, Path(tmp) / "marks.pptx")
            self.assertEqual([], checks.run_all(saved))
            self.assertEqual(1, len(Presentation(str(saved)).slides))
            with zipfile.ZipFile(saved) as archive:
                xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        return re.findall(r"<a:tc>.*?</a:tc>", xml, re.S)

    def test_a_mark_is_written_as_a_character_in_the_face_that_carries_it(self) -> None:
        """⚠ **印の字形を持つ書体は限られる。**案件の書体が Arial なら、4 分の 1 の円も ✓ も無い。"""
        for theme in (DEFAULT, theme_from({"font": "Arial"})):
            with self.subTest(theme.type.family):
                cells = self._cells(theme)
                self.assertEqual(12, len(cells))
                faces = [re.findall(r'<a:(latin|ea|cs|sym) typeface="([^"]*)"', cell) for cell in cells]
                for index, found in enumerate(faces):
                    if index >= 3 and index % 3:
                        self.assertEqual([("latin", MARK_FACE), ("ea", MARK_FACE), ("sym", MARK_FACE)], found)
                    else:
                        self.assertEqual([("latin", theme.type.family)], found)
                self.assertIn("<a:t>●</a:t>", cells[4])
                self.assertIn("<a:t>✓</a:t>", cells[5])

    def test_the_faces_of_a_mark_come_in_the_order_the_file_format_keeps(self) -> None:
        """⚠ 順を違えた file を PowerPoint は修復しようとする (= 色、欧文、東アジア、記号の順)。"""
        mark = self._cells()[4]
        properties = re.search(r"<a:rPr[^>]*>(.*?)</a:rPr>", mark, re.S).group(1)
        self.assertEqual(["solidFill", "srgbClr", "latin", "ea", "sym"], re.findall(r"<a:(\w+)[ >]", properties))

    def test_a_mark_is_set_like_the_cells_beside_it(self) -> None:
        words, mark = self._cells()[3], self._cells()[4]
        for attribute in (r'algn="(\w+)"', r'sz="(\d+)"', r'<a:srgbClr val="(\w+)"'):
            self.assertEqual(re.findall(attribute, words), re.findall(attribute, mark), attribute)

    def test_the_listing_of_types_says_what_a_cell_may_hold(self) -> None:
        said = types.describe()
        self.assertIn("harvey = 0 to 4", said)
        self.assertIn("mark = ok | partial | ng", said)
        self.assertIn("sticker", said)


if __name__ == "__main__":
    unittest.main()
