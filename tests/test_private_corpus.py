"""The push check against private decks (= .tooling/local-ci/private-corpus-check.py).

⚠ **語の一覧は、実物から写した値を知らない。**試しに使った実物の頁の値を test や例に
写すと、名前の一覧には無いので素通りする。この検査は私的なデッキの文言そのものと突き合わせる。
ここで使うデッキは合成したもの (= 私的なデッキは repo に入らない)。
"""

from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parent.parent
SCRIPT = REPO / ".tooling" / "local-ci" / "private-corpus-check.py"


class CheckingAgainstPrivateDecks(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        base = Path(self.tmp.name)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for text in ("社外に出さない測定の説明文", "値は 0.4567 だった"):
            box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(400000))
            box.text_frame.paragraphs[0].add_run().text = text
        (base / "decks").mkdir()
        deck.save(str(base / "decks" / "private.pptx"))
        self.config = base / "config"
        self.config.mkdir()
        (self.config / "sources.txt").write_text(str(base / "decks") + "\n", encoding="utf-8")
        (self.config / "words.txt").write_text("内部の呼び名\n", encoding="utf-8")
        self.env = {**os.environ, "PRIVATE_CORPUS_DIR": str(self.config),
                    "XDG_CACHE_HOME": str(base / "cache")}
        self.base = base

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _check(self, text: str) -> subprocess.CompletedProcess:
        target = self.base / "outgoing.txt"
        target.write_text(text, encoding="utf-8")
        return subprocess.run([sys.executable, str(SCRIPT), "--text", str(target)],
                              capture_output=True, text=True, env=self.env)

    def test_a_run_from_a_private_deck_is_stopped(self) -> None:
        done = self._check('title = "社外に出さない測定の説明文"\n')
        self.assertEqual(done.returncode, 1, done.stdout)

    def test_a_value_from_a_private_deck_is_stopped(self) -> None:
        """実物の頁から写した測定値 (= 語の一覧には決して載らない物)。"""
        done = self._check('table = [["幅", "0.4567"]]\n')
        self.assertEqual(done.returncode, 1, done.stdout)

    def test_a_listed_word_is_stopped(self) -> None:
        self.assertEqual(self._check("# 内部の呼び名 の頁\n").returncode, 1)

    def test_made_up_values_pass(self) -> None:
        done = self._check('table = [["幅", "1.5"]]\ntitle = "見本の題"\n')
        self.assertEqual(done.returncode, 0, done.stdout)

    def test_an_allowed_phrase_passes(self) -> None:
        (self.config / "allow.txt").write_text("社外に出さない測定の説明文\n", encoding="utf-8")
        self.assertEqual(self._check('title = "社外に出さない測定の説明文"\n').returncode, 0)

    def test_a_machine_with_no_private_decks_says_it_did_not_check(self) -> None:
        (self.config / "sources.txt").unlink()
        done = self._check('title = "社外に出さない測定の説明文"\n')
        self.assertEqual(done.returncode, 0)
        self.assertIn("NOT CHECKED", done.stdout)


if __name__ == "__main__":
    unittest.main()
