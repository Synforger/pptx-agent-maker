"""Each check must fire on a deck that is actually broken, and stay quiet on a good one.

⚠ **検査は、わざと壊したもので効くことを見てから使う。**「赤 0」が意味を持つのは
そのあと。ここで作る壊れた頁は、宣言層では**組めない**ものばかり ― だから
python-pptx で直に置いている (= 複製と輸入で入ってくる頁と同じ立場)。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx_agent_maker import DEFAULT, Page  # noqa: E402
from pptx_agent_maker.checks import CHECKS, run_all  # noqa: E402
from pptx_agent_maker.checks.rules import (  # noqa: E402
    empty_cells, internal_names, off_page, overlap, type_floor, unreplaced, untyped_parts)
from pptx_agent_maker.write import add_page, new_deck, save  # noqa: E402
from test_deck import make_dot  # noqa: E402

CONFIG = {"stale_words": ["W3 の値"]}


def raw_deck(destination: Path, place) -> Path:
    """A deck built without the declarative layer — the way an imported page arrives."""
    deck = Presentation()
    deck.slide_width = Emu(DEFAULT.slide.width)
    deck.slide_height = Emu(DEFAULT.slide.height)
    place(deck.slides.add_slide(deck.slide_layouts[6]))
    destination.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(destination))
    return destination


def text_at(slide, text: str, *, left=Emu(1000000), size=Pt(14)):
    box = slide.shapes.add_textbox(left, Emu(1000000), Emu(3000000), Emu(500000))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = size
    return box


class ChecksFireTest(unittest.TestCase):
    """bad/ 側 — 壊れた頁で必ず鳴る。"""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_a_shape_past_the_edge_is_caught(self) -> None:
        deck = raw_deck(self.dir / "off.pptx",
                        lambda s: text_at(s, "はみ出した帯", left=Emu(DEFAULT.slide.width - 200000)))
        self.assertTrue(off_page.run(deck, {}), "a shape hanging off the page went unreported")

    def test_type_below_the_floor_is_caught(self) -> None:
        deck = raw_deck(self.dir / "small.pptx", lambda s: text_at(s, "読めない注記", size=Pt(7)))
        self.assertTrue(type_floor.run(deck, {}))

    def test_a_blank_cell_is_caught(self) -> None:
        def place(slide):
            table = slide.shapes.add_table(2, 2, Emu(500000), Emu(500000),
                                           Emu(4000000), Emu(1000000)).table
            table.cell(0, 0).text = "手法"
            table.cell(0, 1).text = "値"
            table.cell(1, 0).text = "A"  # (1, 1) left blank on purpose
        deck = raw_deck(self.dir / "blank.pptx", place)
        self.assertTrue(empty_cells.run(deck, {}))

    def test_an_internal_name_is_caught(self) -> None:
        deck = raw_deck(self.dir / "internal.pptx",
                        lambda s: text_at(s, "数値は src/scoring/main.py が出す"))
        self.assertTrue(internal_names.run(deck, {}))

    def test_words_the_specimen_left_behind_are_caught(self) -> None:
        deck = raw_deck(self.dir / "stale.pptx", lambda s: text_at(s, "W3 の値をそのまま掲載"))
        self.assertTrue(unreplaced.run(deck, CONFIG))

    def test_words_printed_on_top_of_each_other_are_caught(self) -> None:
        """⚠ **枠ではなく描かれた文字で見る** ― 枠で見ていた頃は、画面では離れている
        物が上がり、実デッキでは上がったものが全部空振りだった。
        """
        def two_lines_in_one_place(slide):
            text_at(slide, "うえに乗る長い文字の行")
            box = slide.shapes.add_textbox(Emu(1000000), Emu(1050000),
                                           Emu(3000000), Emu(500000))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = "したに敷かれる長い文字の行"
            run.font.size = Pt(14)

        deck = raw_deck(self.dir / "stacked.pptx", two_lines_in_one_place)
        self.assertTrue(overlap.run(deck), "重なっている文字が上がらない")

    def test_a_part_with_no_content_type_is_caught(self) -> None:
        """⚠ 頁には何も出ない壊れ方 ― 絵の file だけが写って、種類の登録が無い。"""
        import zipfile
        good = raw_deck(self.dir / "typed.pptx", lambda s: text_at(s, "ふつう"))
        broken = self.dir / "untyped.pptx"
        # 実際に起きた形 = 絵の file だけが入り、png の登録が無い (= 文字だけのデッキには無い)
        with zipfile.ZipFile(good) as source, zipfile.ZipFile(broken, "w") as target:
            self.assertNotRegex(source.read("[Content_Types].xml").decode(), 'Extension="png"')
            for item in source.infolist():
                target.writestr(item, source.read(item))
            target.write(REPO / "tests" / "data" / "dot.png", "ppt/media/image9.png")
        self.assertFalse(untyped_parts.run(good), "a sound package was reported")
        self.assertTrue(untyped_parts.run(broken), "a part with no type went through")

    def test_every_check_has_a_fixture_that_fires_it(self) -> None:
        """A check nobody proved is a check nobody can trust."""
        proven = {off_page.NAME, overlap.NAME, type_floor.NAME, empty_cells.NAME,
                  internal_names.NAME, unreplaced.NAME, untyped_parts.NAME}
        self.assertEqual({check.NAME for check in CHECKS}, proven,
                         "a check exists with no fixture proving it fires")


class ChecksStayQuietTest(unittest.TestCase):
    """clean/ 側 — 宣言層で組んだ頁では 1 件も鳴らない。"""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        make_dot()
        page = Page("読める頁", condition="条件", conclusion="結論", footer="出所")
        left, right = page.body.columns([2, 1], gap=DEFAULT.spacing.gap_m)
        page.figure(left, REPO / "tests" / "data" / "dot.png", 1.0, caption="図")
        rows = [["列", "値"], ["A", "1"]]
        table, note = right.split_top(DEFAULT.table_height(len(rows)), gap=DEFAULT.spacing.gap_s)
        page.table(table, rows)
        page.note(note, "読み方")
        deck = new_deck()
        add_page(deck, page.build())
        self.deck = save(deck, self.dir / "clean.pptx")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_a_declared_page_reports_nothing(self) -> None:
        findings = run_all(self.deck, CONFIG)
        self.assertEqual(findings, [], "\n".join(f.render() for f in findings))


if __name__ == "__main__":
    unittest.main()
