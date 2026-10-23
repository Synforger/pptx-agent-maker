"""The frame each declared picture is given, read without building (= `frames`).

⚠ **絵は縦横比を保って枠の中に置かれる。**枠が分からないまま描いたグラフは、枠より縦長なら高さで、横長なら
幅で縮められ、描いた字が頁の上で小さくなる。枠を読んでその大きさで描けば、絵は枠いっぱいに置かれる ― 枠は
頁を組まないと決まらないので、build と同じ道で組んで読む。焼かないので、`out` の deck には触らない。
"""

from __future__ import annotations

import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.deck.frames import frames  # noqa: E402
from pptx_agent_maker.project import Workspace, create  # noqa: E402
from pptx_agent_maker.project.files.manifest import Manifest  # noqa: E402

CORPUS = REPO / "tests" / "corpus" / "assets"
HEAD = 'specimen = "specimen.pptx"\nout = "deck.pptx"\n'
COVER = '[[pages]]\nkind = "copy"\npage = 1\n'


def figure_page(picture: str, **more: str) -> str:
    keys = "".join(f'{key} = "{value}"\n' for key, value in more.items())
    return f'[[pages]]\nkind = "declare"\ntype = "figure"\ntitle = "A picture"\nfigure = "{picture}"\n{keys}'


class TheFrameOfEachPicture(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "project"
        create(self.root)
        for name in ("wide.png", "tall.png", "square.png"):
            shutil.copy(CORPUS / name, self.root / "assets" / name)

    def frames_of(self, *pages: str) -> list[dict]:
        (self.root / "deck.toml").write_text(HEAD + "\n".join(pages), encoding="utf-8")
        workspace = Workspace.load(self.root)
        return frames(workspace, Manifest.load(workspace.manifest("deck")))

    def test_each_picture_says_its_page_its_name_its_frame_and_where_it_is_placed(self) -> None:
        found = self.frames_of(COVER, figure_page("wide.png"), figure_page("tall.png"))
        self.assertEqual([(2, "wide.png"), (3, "tall.png")], [(f["page"], f["picture"]) for f in found])
        for entry in found:
            frame, placed = entry["frame"], entry["placed"]
            self.assertLessEqual(placed["width_emu"], frame["width_emu"])
            self.assertLessEqual(placed["height_emu"], frame["height_emu"])
            self.assertAlmostEqual(frame["width_emu"] / 914400, frame["width_in"], places=3)
            self.assertAlmostEqual(frame["height_emu"] / 914400, frame["height_in"], places=3)

    def test_a_picture_fills_the_side_its_shape_runs_out_on(self) -> None:
        """横に長い絵は幅いっぱいで低く、縦に長い絵は高さいっぱいで細い (= 余りが隙間になる)。"""
        wide, tall = self.frames_of(figure_page("wide.png"), figure_page("tall.png"))
        self.assertEqual(wide["frame"]["width_emu"], wide["placed"]["width_emu"])
        self.assertLess(wide["placed"]["height_emu"], wide["frame"]["height_emu"])
        self.assertEqual(tall["frame"]["height_emu"], tall["placed"]["height_emu"])
        self.assertLess(tall["placed"]["width_emu"], tall["frame"]["width_emu"])

    def test_the_frame_leaves_out_the_pictures_name(self) -> None:
        """名 (= `caption`) は絵の下に付き、その行ぶんは枠に入らない。"""
        plain, named = self.frames_of(figure_page("tall.png"), figure_page("tall.png", caption="What it shows"))
        self.assertEqual(plain["frame"]["width_emu"], named["frame"]["width_emu"])
        self.assertLess(named["frame"]["height_emu"], plain["frame"]["height_emu"])

    def test_the_same_picture_on_two_pages_gives_a_frame_for_each(self) -> None:
        found = self.frames_of(figure_page("square.png"), figure_page("square.png", caption="Named"))
        self.assertEqual([1, 2], [entry["page"] for entry in found])
        self.assertNotEqual(found[0]["frame"], found[1]["frame"])

    def test_where_it_says_a_picture_is_placed_is_where_the_build_puts_it(self) -> None:
        from pptx import Presentation

        found = self.frames_of(figure_page("wide.png"), figure_page("tall.png", caption="Named"))
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(0, main(["build", str(self.root), "deck", "--skip-checks"]))
        built = Presentation(self.root / "deck.pptx")
        for entry in found:
            pictures = [shape for shape in built.slides[entry["page"] - 1].shapes if shape.shape_type == 13]
            self.assertEqual([(entry["placed"]["width_emu"], entry["placed"]["height_emu"])],
                             [(shape.width, shape.height) for shape in pictures])

    def test_nothing_is_built(self) -> None:
        self.frames_of(figure_page("wide.png"))
        self.assertFalse((self.root / "deck.pptx").exists())

    def test_the_command_prints_the_list_as_json(self) -> None:
        (self.root / "deck.toml").write_text(HEAD + COVER + "\n" + figure_page("wide.png"), encoding="utf-8")
        said = io.StringIO()
        with contextlib.redirect_stdout(said):
            code = main(["frames", str(self.root), "deck"])
        self.assertEqual(0, code)
        listed = json.loads(said.getvalue())
        self.assertEqual([2], [entry["page"] for entry in listed])
        self.assertEqual({"width_emu", "height_emu", "width_in", "height_in"}, set(listed[0]["frame"]))


if __name__ == "__main__":
    unittest.main()
