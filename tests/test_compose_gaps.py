"""compose: the gap across a row and the gap between rows, cells of no words, rules over a picture, one `weight`.

見張るのは 4 つ ― 段のマスとマスの間 (`gap_x`) と段と段の間 (`gap_y`) が別に取れ、書かなければ `gap_m` に従うこと、
空の文章のマスが高さを取らず何も描かないこと、図のマスが `guides` で絵の高さの割合の所に絵の幅いっぱいの線を引くこと、
`weight` の検査が段・マス・道のりの段で 1 つの関数と 1 つの文面になること。
"""

from __future__ import annotations

import math
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))
from support.pages import box_of, build  # noqa: E402

from pptx_agent_maker import DEFAULT, cm  # noqa: E402
from pptx_agent_maker.layout.base.geometry import Rect  # noqa: E402
from pptx_agent_maker.layout.base.tokens import SPACED, TIGHTENED, ThemeError, Under, theme_from  # noqa: E402
from pptx_agent_maker.layout.parts.elements import Figure, Line, Text  # noqa: E402
from pptx_agent_maker.layout.parts.page import Page, PageFullError  # noqa: E402
from pptx_agent_maker.layout.types.core.read import read_weight  # noqa: E402
from pptx_agent_maker.layout.types.core.registry import PageTypeError  # noqa: E402
from pptx_agent_maker.write import add_page, new_deck, save  # noqa: E402

SQUARE, WIDE = "dot.png", "wide.png"      # 8 x 8, 16 x 4


def using(**spacing):
    """A deck whose page uses a spacing of its own, like the deck to be read."""
    return theme_from({"uses": {"gapped": {"like": "read", "spacing": spacing}}}).using("gapped")


def compose(rows, theme=DEFAULT, **extra):
    return build({"type": "compose", "title": "A page", "rows": rows, **extra}, theme)


def row(*cells, **more):
    return {"cells": list(cells), **more}


def card(heading, body="", **more):
    return {"card": {"heading": heading, "body": body, **more}}


class Refusing(unittest.TestCase):
    def refused(self, rows, expected, theme=DEFAULT, **extra) -> None:
        """Build a page that has to stop, and check what the stop says."""
        with self.assertRaises(PageTypeError) as stopped:
            compose(rows, theme, **extra).build()
        self.assertIn(expected, str(stopped.exception))


GRID = [row(card("A", "a"), card("B", "b")), row(card("C", "c"), card("D", "d"))]


def boxes(built, *headings):
    return [box_of(built, heading).rect for heading in headings]


class TheGapAcrossAndTheGapDown(unittest.TestCase):
    """A row's cells stand `gap_x` apart and the rows `gap_y` apart; unsaid, both are `gap_m`."""

    def test_cells_stand_apart_by_gap_x_and_rows_by_gap_y(self) -> None:
        theme = using(gap_x=2.0, gap_y=0.1)
        a, b, c, _d = boxes(compose(GRID, theme), "A", "B", "C", "D")
        self.assertEqual(cm(2.0), b.left - a.right)
        self.assertEqual(cm(0.1), c.top - a.bottom)
        self.assertNotEqual(theme.spacing.gap_m, theme.spacing.gap_x)

    def test_the_two_gaps_are_independent_of_each_other(self) -> None:
        a, b, c, _d = boxes(compose(GRID, using(gap_x=0.2)), "A", "B", "C", "D")
        self.assertEqual((cm(0.2), DEFAULT.spacing.gap_m), (b.left - a.right, c.top - a.bottom))
        a, b, c, _d = boxes(compose(GRID, using(gap_y=1.2)), "A", "B", "C", "D")
        self.assertEqual((DEFAULT.spacing.gap_m, cm(1.2)), (b.left - a.right, c.top - a.bottom))

    def test_a_use_that_says_only_gap_m_sets_both(self) -> None:
        """今まで `gap_m` で隙間を決めていた案件は、書き換えずに同じ頁になる。"""
        theme = using(gap_m=1.5)
        self.assertEqual((cm(1.5),) * 3, (theme.spacing.gap_m, theme.spacing.gap_x, theme.spacing.gap_y))
        a, b, c, _d = boxes(compose(GRID, theme), "A", "B", "C", "D")
        self.assertEqual((cm(1.5), cm(1.5)), (b.left - a.right, c.top - a.bottom))

    def test_gap_m_with_one_of_them_written_apart(self) -> None:
        theme = using(gap_m=1.5, gap_y=0.2)
        self.assertEqual((cm(1.5), cm(1.5), cm(0.2)), (theme.spacing.gap_m, theme.spacing.gap_x, theme.spacing.gap_y))

    def test_unsaid_they_are_the_gap_of_the_deck(self) -> None:
        for name in ("present", "read", "sheet"):
            with self.subTest(name):
                s = DEFAULT.using(name).spacing
                self.assertEqual((s.gap_m, s.gap_m), (s.gap_x, s.gap_y))

    def test_a_use_may_write_them_and_the_scale_tightens_them(self) -> None:
        self.assertLessEqual({"gap_x", "gap_y"}, set(SPACED))
        self.assertLessEqual({"gap_x", "gap_y"}, set(TIGHTENED))
        half = theme_from({"uses": {"gapped": {"like": "read", "spacing": {"scale": 0.5}}}}).using("gapped").spacing
        self.assertEqual((round(DEFAULT.spacing.gap_x * 0.5), round(DEFAULT.spacing.gap_y * 0.5)),
                         (half.gap_x, half.gap_y))
        written = using(scale=0.5, gap_x=2.0).spacing
        self.assertEqual((cm(2.0), round(DEFAULT.spacing.gap_y * 0.5)), (written.gap_x, written.gap_y))

    def test_a_negative_gap_is_refused(self) -> None:
        with self.assertRaises(ThemeError) as stopped:
            using(gap_x=-0.1)
        self.assertIn("gap_x is -0.1cm", str(stopped.exception))

    def test_rows_nested_in_a_cell_stand_apart_the_same_way(self) -> None:
        built = compose([row(card("Beside"), {"rows": [row(card("Upper"), card("Right")), row(card("Lower"))]})],
                        using(gap_x=2.0, gap_y=0.1))
        beside, upper, right, lower = boxes(built, "Beside", "Upper", "Right", "Lower")
        self.assertEqual(cm(0.1), lower.top - upper.bottom)
        self.assertEqual(cm(2.0), right.left - upper.right)
        # 入れ子の高さも `gap_y` で数える (= 隣のカードは入れ子の上から下まで)
        self.assertEqual((upper.top, lower.bottom), (beside.top, beside.bottom))

    def test_a_panel_divides_its_rows_the_same_way_and_is_as_tall_as_they_are(self) -> None:
        theme = using(gap_x=2.0, gap_y=0.1)
        rows = [row(card("One"), card("Two")), row(card("Three"))]
        built = compose([row({"panel": {"heading": "Held", "rows": rows}}, card("Beside"))], theme)
        one, two, three, beside = boxes(built, "One", "Two", "Three", "Beside")
        panel = next(e for e in built.build() if e.kind == "panel")
        self.assertEqual(cm(2.0), two.left - one.right)
        self.assertEqual(cm(0.1), three.top - one.bottom)
        pad = theme.spacing.pad
        self.assertEqual(pad, panel.rect.bottom - three.bottom, "the panel ends one pad under its last row")
        self.assertEqual(panel.rect.height, beside.height, "and the card beside it stands as tall")


class ACellOfNoWordsTakesNoHeight(unittest.TestCase):
    """`text = ""` is room a row may leave and an arrow may cross, not a line of words."""

    @staticmethod
    def top_of(rows, heading):
        return box_of(compose(rows), heading).rect.top

    def test_a_row_of_only_empty_cells_is_no_taller_than_its_gaps(self) -> None:
        gap = DEFAULT.spacing.gap_y
        a, d = boxes(compose([row(card("A")), row({"text": ""}, {"text": "  "}), row(card("D"))]), "A", "D")
        self.assertEqual(2 * gap, d.top - a.bottom)

    def test_an_empty_cell_beside_others_leaves_the_row_as_tall_as_they_are(self) -> None:
        with_it = boxes(compose([row({"text": ""}, card("B", "body"))]), "B")[0]
        without = boxes(compose([row(card("B", "body"))]), "B")[0]
        self.assertEqual(without.height, with_it.height)

    def test_an_empty_cell_beside_a_short_cell_does_not_make_the_row_taller_than_the_short_cell(self) -> None:
        built = compose([row({"text": ""}, {"text": "Short"}), row(card("Under"))])
        short = next(e for e in built.build() if getattr(e, "text", "") == "Short")
        under = box_of(built, "Under").rect
        self.assertEqual(under.top - DEFAULT.spacing.gap_y, short.rect.bottom)

    def test_it_draws_nothing(self) -> None:
        built = compose([row({"text": ""}, card("B"))])
        self.assertEqual([], [e for e in built.build() if isinstance(e, Text) and not e.text.strip()])

    def test_a_box_of_code_with_no_words_is_still_a_box(self) -> None:
        """コードの枠は、中が空でも枠が見える物なので高さを取る。"""
        built = compose([row({"text": "", "role": "mono"}), row(card("D"))])
        code = next(e for e in built.build() if e.kind == "code_box")
        self.assertGreater(code.rect.height, 0)

    def test_an_arrow_still_crosses_the_empty_row_between_two_cells(self) -> None:
        built = compose([row({**card("A"), "id": "a"}), row({"text": ""}), row({**card("D"), "id": "d"})],
                        links=[{"from": "a", "to": "d"}])
        lines = [e for e in built.build() if isinstance(e, Line)]
        a, d = boxes(built, "A", "D")
        arrow = next(e for e in lines if e.arrow)
        self.assertEqual((a.bottom, d.top), (arrow.start[1], arrow.end[1]))


class RulesOverAPicture(Refusing):
    """`guides` rules the picture at shares of its own height, each rule as wide as the picture."""

    @staticmethod
    def rules(built, picture):
        return sorted((e for e in built.build() if isinstance(e, Line) and not e.arrow
                       and picture.left <= e.start[0] <= picture.right and picture.top <= e.start[1] <= picture.bottom),
                      key=lambda line: line.start[1])

    @staticmethod
    def picture(built):
        return next(e for e in built.build() if isinstance(e, Figure) and e.kind == "figure").rect

    def test_each_rule_runs_the_width_of_the_picture_at_its_share_of_the_height(self) -> None:
        built = compose([row({"figure": SQUARE, "guides": [0.1, 0.5, 0.75]})])
        picture = self.picture(built)
        rules = self.rules(built, picture)
        self.assertEqual(3, len(rules))
        for rule, share in zip(rules, (0.1, 0.5, 0.75)):
            with self.subTest(share=share):
                self.assertEqual(rule.start[1], rule.end[1], "a horizontal rule")
                self.assertEqual((picture.left, picture.right), (rule.start[0], rule.end[0]))
                self.assertEqual(picture.top + round(share * picture.height), rule.start[1])
                self.assertFalse(rule.arrow)

    def test_the_rules_are_measured_on_the_picture_not_on_the_cell(self) -> None:
        """正方形の絵は、横長のマスの中で幅が狭い (= 線もそこまで)。名が下に付いても、測るのは絵だけ。"""
        built = compose([row({"figure": SQUARE, "guides": [0.99], "caption": "A name under the picture"}),
                         row(card("Under"))])
        picture = self.picture(built)
        (rule,) = self.rules(built, picture)
        caption = next(e for e in built.build() if getattr(e, "text", "") == "A name under the picture")
        self.assertEqual(picture.top + round(0.99 * picture.height), rule.start[1])
        self.assertLess(rule.start[1], picture.bottom)
        self.assertLess(picture.bottom, caption.rect.top)
        self.assertEqual((picture.left, picture.right), (rule.start[0], rule.end[0]))

    def test_the_colour_is_the_accent_unless_a_tone_is_asked_for(self) -> None:
        plain = compose([row({"figure": SQUARE, "guides": [0.5]})])
        (rule,) = self.rules(plain, self.picture(plain))
        self.assertEqual(Page("x", DEFAULT).tone_ink("accent"), rule.colour)
        asked = compose([row({"figure": SQUARE, "guides": [0.5], "guide_tone": "bad"})])
        (rule,) = self.rules(asked, self.picture(asked))
        self.assertEqual(Page("x", DEFAULT).tone_ink("bad"), rule.colour)
        self.assertNotEqual(Page("x", DEFAULT).tone_ink("accent"), rule.colour)

    def test_a_tone_the_project_named_is_read_like_any_other(self) -> None:
        theme = theme_from({"grounds": {"Late": "8B1E3F"}})
        built = compose([row({"figure": SQUARE, "guides": [0.5], "guide_tone": "Late"})], theme)
        (rule,) = self.rules(built, self.picture(built))
        self.assertEqual(Page("x", theme).tone_ink("Late"), rule.colour)

    def test_a_picture_without_guides_has_no_rule(self) -> None:
        built = compose([row({"figure": SQUARE})])
        self.assertEqual([], self.rules(built, self.picture(built)))

    def test_a_rule_is_as_thin_as_the_other_lines_of_the_toolkit(self) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        built = compose([row({"figure": SQUARE, "guides": [0.5]})])
        with tempfile.TemporaryDirectory() as folder:
            deck, out = new_deck(), Path(folder) / "ruled.pptx"
            add_page(deck, built.build())
            save(deck, out)
            slide = Presentation(str(out)).slides[0]
        widths = {shape.line.width for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE}
        self.assertEqual({DEFAULT.spacing.hairline * 3 // 2}, widths)

    def test_what_is_not_a_share_is_refused_with_the_reason(self) -> None:
        for value in (0, 1, -0.1, 1.5, True, "0.5", None, math.nan, math.inf):
            with self.subTest(value=value):
                self.refused([row({"figure": SQUARE, "guides": [0.5, value]})], "above 0 and below 1")
        for value in ([], 0.5, "0.5", {"at": 0.5}):
            with self.subTest(value=value):
                self.refused([row({"figure": SQUARE, "guides": value})], "`guides` is a list of numbers between 0 and 1")

    def test_a_tone_that_is_no_tone_is_refused(self) -> None:
        self.refused([row({"figure": SQUARE, "guides": [0.5], "guide_tone": "teal"})], "guide_tone")

    def test_the_tone_without_rules_is_refused(self) -> None:
        self.refused([row({"figure": SQUARE, "guide_tone": "bad"})], "`guide_tone` is the colour of the rules")

    def test_only_a_figure_is_ruled(self) -> None:
        self.refused([row({**card("A"), "guides": [0.5]})], "`guides` belongs to a figure")
        self.refused([row({**card("A"), "guide_tone": "bad"})], "`guide_tone` belongs to a figure")

    def test_marks_and_rules_are_two_ways_to_read_the_same_height(self) -> None:
        self.refused([row({"figure": SQUARE, "guides": [0.5], "marks": [{"at": [0.5, 0.5], "text": "here"}]})],
                "`marks` or `guides`, not both")

    def test_a_ruled_figure_is_read_inside_a_panel_as_well(self) -> None:
        self.refused([row({"panel": {"rows": [row({"figure": SQUARE, "guides": [2]})]}})], "above 0 and below 1")


class AFigureStandsWhereItsCellPutsIt(Refusing):
    """A picture lower than its cell stands at the top unless `valign` says the middle or the bottom."""

    @staticmethod
    def stood(**more):
        built = compose([row({"figure": WIDE, **more})])
        return next(e for e in built.build() if isinstance(e, Figure) and e.kind == "figure").rect, \
            Page("A page", DEFAULT).body

    def test_unsaid_it_stands_at_the_top_as_it_always_did(self) -> None:
        picture, body = self.stood()
        self.assertEqual(body.top, picture.top)
        self.assertLess(picture.height, body.height, "the cell is taller than the picture, or this proves nothing")
        self.assertEqual(picture, self.stood(valign="top")[0])

    def test_in_the_middle_and_at_the_bottom(self) -> None:
        top, body = self.stood()
        middle, _ = self.stood(valign="middle")
        bottom, _ = self.stood(valign="bottom")
        spare = body.height - top.height
        self.assertAlmostEqual(body.top + spare // 2, middle.top, delta=1)
        self.assertAlmostEqual(body.bottom, bottom.bottom, delta=1)
        self.assertEqual((top.left, top.width, top.height), (middle.left, middle.width, middle.height))
        self.assertEqual((top.left, top.width, top.height), (bottom.left, bottom.width, bottom.height))

    def test_the_name_goes_with_the_picture(self) -> None:
        built = compose([row({"figure": WIDE, "caption": "A name", "valign": "bottom"})])
        name = next(e for e in built.build() if getattr(e, "text", "") == "A name")
        self.assertAlmostEqual(Page("A page", DEFAULT).body.bottom, name.rect.bottom, delta=DEFAULT.spacing.text_inset)

    def test_the_rules_follow_the_picture(self) -> None:
        built = compose([row({"figure": WIDE, "guides": [0.5], "valign": "bottom"})])
        picture = next(e for e in built.build() if isinstance(e, Figure) and e.kind == "figure").rect
        (rule,) = [e for e in built.build() if isinstance(e, Line)
                   if picture.top <= e.start[1] <= picture.bottom and e.start[0] == picture.left]
        self.assertEqual(picture.top + round(0.5 * picture.height), rule.start[1])

    def test_a_picture_that_fills_the_height_of_its_cell_has_nowhere_to_move(self) -> None:
        stood = []
        for valign in ("top", "middle", "bottom"):
            built = compose([row({"figure": SQUARE, "valign": valign})])
            stood.append(next(e for e in built.build() if isinstance(e, Figure) and e.kind == "figure").rect)
        self.assertEqual(1, len(set(stood)))

    def test_what_is_not_a_place_is_refused(self) -> None:
        for value in ("centre", "", 1, None, ["top"]):
            with self.subTest(value=value):
                self.refused([row({"figure": WIDE, "valign": value})], "`valign` is 'top' or 'middle' or 'bottom'")

    def test_only_a_figure_has_a_place_in_its_cell(self) -> None:
        self.refused([row({**card("A"), "valign": "middle"})], "`valign` belongs to a figure")

    def test_a_figure_with_marks_stands_where_its_names_put_it(self) -> None:
        self.refused([row({"figure": WIDE, "valign": "middle", "marks": [{"at": [0.5, 0.5], "text": "here"}]})],
                     "`valign` is for a figure without them")


class AFigureKeepsRoomBesideIt(Refusing):
    """`inset` narrows a picture by the same room on its left and its right; the cell and its neighbours stay put."""

    @staticmethod
    def stood(*beside, **more):
        built = compose([row({"figure": WIDE, **more}, *beside)]).build()
        picture = next(e for e in built if isinstance(e, Figure) and e.kind == "figure").rect
        return picture, [e.rect for e in built if isinstance(e, Text) and e.text == "beside"]

    def test_unsaid_or_zero_the_picture_is_as_wide_as_its_cell(self) -> None:
        self.assertEqual(self.stood()[0], self.stood(inset=0)[0])

    def test_the_picture_is_narrowed_on_both_sides(self) -> None:
        wide, _ = self.stood()
        kept, _ = self.stood(inset=0.28)
        self.assertAlmostEqual(wide.left + cm(0.28), kept.left, delta=1)
        self.assertAlmostEqual(wide.right - cm(0.28), kept.right, delta=1)
        self.assertLess(kept.height, wide.height, "a narrower picture keeps its shape, so it is lower")
        low, _ = self.stood(inset=0.28, valign="bottom")
        self.assertAlmostEqual(Page("A page", DEFAULT).body.bottom, low.bottom, delta=1)

    def test_the_cell_beside_it_does_not_move(self) -> None:
        _, plain = self.stood({"text": "beside"})
        _, kept = self.stood({"text": "beside"}, inset=0.5)
        self.assertEqual(plain, kept)

    def test_what_inset_cannot_be(self) -> None:
        for inset in (-0.1, "0.3", True):
            with self.subTest(inset=inset):
                self.refused([row({"figure": WIDE, "inset": inset})], "`inset` is the room left and right")
        self.refused([row({"text": "words", "inset": 0.3})], "`inset` belongs to a figure")
        self.refused([row({"figure": WIDE, "inset": 0.3, "marks": [{"at": [0.5, 0.5], "text": "here"}]})],
                     "`inset` is for a figure without them")

    def test_an_inset_that_leaves_no_picture_stops(self) -> None:
        with self.assertRaises(PageFullError) as stopped:
            compose([row({"figure": WIDE, "inset": 20})]).build()
        self.assertIn("no room for the picture", str(stopped.exception))


class TheGapBetweenABandAndTheBody(unittest.TestCase):
    """The title, the bands and the footer stand `band_gap` from the body; unsaid it is `gap_m`, as it always was."""

    FRAME = {"condition": "A condition", "conclusion": "A conclusion", "footer": "A source"}

    def body(self, theme):
        return Page("A page", theme, **self.FRAME).body

    def test_it_moves_the_body_at_the_top_and_at_the_bottom_and_nothing_else(self) -> None:
        plain = self.body(DEFAULT)
        apart = self.body(using(band_gap=1.6))
        # 題の下と条件の帯の下、結論の帯の上 (足元の帯の上も) で 1 つずつ
        extra = cm(1.6) - DEFAULT.spacing.gap_m
        self.assertEqual(plain.top + 2 * extra, apart.top)
        self.assertEqual(plain.bottom - 2 * extra, apart.bottom)

    def test_the_gap_between_boxes_does_not_move_it(self) -> None:
        self.assertEqual(self.body(DEFAULT), self.body(using(gap_m=0.1, gap_x=0.1, gap_y=0.1, band_gap=0.6)))

    def test_a_use_that_says_only_gap_m_moves_both(self) -> None:
        theme = using(gap_m=0.1)
        self.assertEqual(theme.spacing.gap_m, theme.spacing.band_gap)
        self.assertEqual(self.body(DEFAULT).top + 2 * (cm(0.1) - DEFAULT.spacing.gap_m), self.body(theme).top)

    def test_it_follows_the_scale_and_a_use_may_write_it(self) -> None:
        self.assertIn("band_gap", SPACED)
        self.assertIn("band_gap", TIGHTENED)
        half = theme_from({"uses": {"gapped": {"like": "read", "spacing": {"scale": 0.5}}}}).using("gapped").spacing
        self.assertEqual(round(DEFAULT.spacing.band_gap * 0.5), half.band_gap)
        self.assertEqual(cm(1.0), using(scale=0.5, band_gap=1.0).spacing.band_gap)

    def test_the_title_the_layout_places_is_set_the_same_way(self) -> None:
        """題の枠をテンプレートが持つ頁 (= `_take_at`) も、題の下は `band_gap`。"""
        seat = Rect(cm(1.2), cm(0.9), cm(20), cm(1.5))

        def top_of(theme):
            return Page("A page", replace(theme, under=Under(title=True, title_frame=seat))).body.top

        self.assertEqual(seat.bottom + DEFAULT.spacing.gap_m, top_of(DEFAULT))
        self.assertEqual(seat.bottom + cm(1.6), top_of(using(band_gap=1.6)))


class OneWeight(unittest.TestCase):
    """The share of a row, a cell or a stage of a road is read in one place, and says one thing when it is wrong."""

    def test_unsaid_it_is_one_and_a_number_is_a_float(self) -> None:
        self.assertEqual(1.0, read_weight({}, "x"))
        self.assertEqual(2.0, read_weight({"weight": 2}, "x"))
        self.assertEqual(0.5, read_weight({"weight": 0.5}, "x"))
        self.assertIsInstance(read_weight({"weight": 3}, "x"), float)

    def test_what_is_not_a_number_above_zero_is_refused(self) -> None:
        for value in (0, -1, True, False, "2", None, [1], math.nan, math.inf, -math.inf):
            with self.subTest(value=value):
                with self.assertRaises(PageTypeError) as stopped:
                    read_weight({"weight": value}, "where")
                self.assertEqual(f"where: `weight` is a number above 0 (= a share), not {value!r}", str(stopped.exception))

    def test_a_row_a_cell_and_a_stage_stop_with_the_same_words(self) -> None:
        said = []
        for build_it in (
                lambda: compose([row(card("A"), weight=0)]),
                lambda: compose([row({**card("A"), "weight": 0})]),
                lambda: build({"type": "roadmap", "title": "T", "stages": [
                    {"name": "One", "weight": 0}, {"name": "Two"}]}),
        ):
            with self.assertRaises(PageTypeError) as stopped:
                build_it().build()
            said.append(str(stopped.exception).split(": `weight`", 1)[1])
        self.assertEqual(1, len(set(said)), said)


if __name__ == "__main__":
    unittest.main()
