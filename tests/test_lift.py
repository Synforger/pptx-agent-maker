"""Lifting into the template every later project starts from.

⚠ **上げた物は、次の案件すべてに配られる。**ここは「指した語のほかは、案件の文言・絵・表の値・
代替テキストがテンプレートに残らない」「止まった時はテンプレートに何も書かない」「次の案件で
埋め忘れると検査が止める」を見る。
"""

from __future__ import annotations

import io
import shutil
import sys
import tempfile
import tomllib
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.checks import run_all  # noqa: E402
from pptx_agent_maker.checks.rules import untyped_parts  # noqa: E402
from pptx_agent_maker.project import create  # noqa: E402
from pptx_agent_maker.project.files.recipes import load_parts  # noqa: E402

DATA = Path(__file__).resolve().parent / "data"

W1 = '''specimen = "specimen.pptx"
out = "w1.pptx"

[[pages]]
kind = "declare"
type = "figure"
title = "手法 A の結果"
figure = "dot.png"
table = [["項目", "手法 A"], ["幅", "1.5"], ["高さ", "1.5"]]
'''

RECIPES = '''[recipes.result]
type = "figure"
title = "{method} の結果"
footer = "手法 A の採点表から"
'''


class Lifting(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        base = Path(self.tmp.name)
        # 会社テンプレートの代わり = ツール既定のテンプレートと設定を置いた folder
        seed = create(base / "seed")
        self.template = base / "template"
        self.template.mkdir()
        for name in ("specimen.pptx", "workspace.toml"):
            shutil.copy(seed / name, self.template / name)
        self.project = create(base / "project", specimen=self.template)
        shutil.copy(DATA / "dot.png", self.project / "assets" / "dot.png")
        (self.project / "w1.toml").write_text(W1, encoding="utf-8")
        (self.project / "recipes.toml").write_text(RECIPES, encoding="utf-8")
        self.assertEqual(self._run("build", str(self.project), "w1", "--skip-checks")[0], 0)
        self.before = {name: (self.template / name).read_bytes()
                       for name in ("specimen.pptx", "workspace.toml")}

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _run(self, *argv: str) -> tuple[int, str]:
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main(list(argv))
        return code, out.getvalue() + err.getvalue()

    def _lift(self, *extra: str) -> tuple[int, str]:
        return self._run("lift", str(self.project), str(self.template), *extra)

    def _untouched(self) -> None:
        for name, data in self.before.items():
            self.assertEqual((self.template / name).read_bytes(), data, f"{name} was written")
        self.assertFalse((self.template / "recipes.toml").exists())
        self.assertFalse((self.template / "_archive").exists(), "a refused lift left a copy behind")

    # -- a page -------------------------------------------------------------

    def _page_words(self, number: int = 2) -> tuple[list[str], list[str]]:
        page = Presentation(str(self.template / "specimen.pptx")).slides[number - 1]
        words = [s.text_frame.text for s in page.shapes if s.has_text_frame and s.text_frame.text]
        cells = [c.text for s in page.shapes if s.has_table for r in s.table.rows for c in r.cells]
        return words, cells

    def test_with_nothing_kept_every_word_stands_in(self) -> None:
        """⚠ 既定は抜く側 ― 指し忘れた語は、前の先方の中身ではなく仮の語になる。"""
        code, said = self._lift("--page", "w1.pptx:1")
        self.assertEqual(code, 0, said)
        words, cells = self._page_words()
        for text in words + cells:
            self.assertTrue(text.startswith("<文言 "), f"{text!r} went into the template as it was")

    def test_the_same_words_stand_in_under_the_same_number(self) -> None:
        self._lift("--page", "w1.pptx:1", "--keep", "項目")
        _words, cells = self._page_words()
        self.assertEqual(cells[0], "項目")
        # 2 行に同じ値 1.5 → 同じ仮の語 (= 次の案件は 1 回の差し替えで両方埋まる)
        self.assertEqual(cells[3], cells[5])
        self.assertEqual(len(set(cells[1:])), 4)

    def test_a_lifted_page_carries_nothing_of_the_project(self) -> None:
        code, said = self._lift("--page", "w1.pptx:1", "--replace", "手法 A=<手法名>",
                                "--keep", "項目")
        self.assertEqual(code, 0, said)
        specimen = self.template / "specimen.pptx"
        deck = Presentation(str(specimen))
        self.assertEqual(len(deck.slides), 2, "the page should be added after the template's own")
        page = deck.slides[1]
        words = [s.text_frame.text for s in page.shapes if s.has_text_frame]
        cells = [c.text for s in page.shapes if s.has_table for r in s.table.rows for c in r.cells]
        self.assertIn("<手法名> の結果", words)
        self.assertNotIn("手法 A", " ".join(words + cells))
        self.assertEqual(cells[:2], ["項目", "<手法名>"])
        self.assertTrue(all(c.startswith("<文言 ") for c in cells[2:]), "values stayed in the table")
        self.assertNotIn("1.5", " ".join(cells))
        pictures = [s for s in page.shapes if s.shape_type == 13]
        self.assertTrue(pictures)
        for picture in pictures:
            self.assertNotEqual(picture.image.blob, (DATA / "dot.png").read_bytes())
        with zipfile.ZipFile(specimen) as archive:
            xml = "".join(archive.read(n).decode("utf-8") for n in archive.namelist()
                          if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        self.assertNotIn("dot.png", xml, "the alt text still names the project's file")
        self.assertEqual(untyped_parts.run(specimen), [])

    def test_the_templates_own_pages_stay_as_they_were(self) -> None:
        texts = lambda path: [s.text_frame.text for s in Presentation(str(path)).slides[0].shapes  # noqa: E731
                              if s.has_text_frame]
        original = self.tmp.name + "/original.pptx"
        Path(original).write_bytes(self.before["specimen.pptx"])
        self._lift("--page", "w1.pptx:1", "--replace", "手法 A=<手法名>")
        self.assertEqual(texts(self.template / "specimen.pptx"), texts(original))

    def test_the_stand_in_words_are_checked_in_the_next_project(self) -> None:
        self._lift("--page", "w1.pptx:1", "--replace", "手法 A=<手法名>")
        settings = tomllib.loads((self.template / "workspace.toml").read_text(encoding="utf-8"))
        self.assertIn("<手法名>", settings["checks"]["stale_words"])
        self.assertIn("<文言", settings["checks"]["stale_words"])

        following = create(Path(self.tmp.name) / "next", specimen=self.template)
        (following / "w1.toml").write_text(
            'specimen = "specimen.pptx"\nout = "w1.pptx"\n\n[[pages]]\nkind = "copy"\npage = 2\n',
            encoding="utf-8")
        code, said = self._run("build", str(following), "w1")
        self.assertEqual(code, 1)
        self.assertIn("<手法名>", said, "a stand-in left in place went through")

    #: 人が書き直した `[checks] stale_words` (= 1 行、1 語 1 行、最後の語の行で閉じる、空)。
    #: 1 行の方は、引用符の中に `]` が在る (= エスケープを読み違えると、そこを配列の終わりと取る)
    ARRAYS = {
        "one line": 'stale_words = ["案件名", "第 \\"N]\\" 回"]  # the specimen\'s own\n',
        "one per line": ('stale_words = [\n    "案件名",\n'
                         '    "第 N 回",  # the specimen\'s own [kept]\n]\n'),
        "closed on the last line": 'stale_words = ["案件名",\n  "第 N 回"]\n',
        "empty": "stale_words = [\n]\n",
    }

    def test_the_stand_in_words_join_the_array_however_it_is_written(self) -> None:
        """⚠ 複数行の配列は 1 行目だけが書き直され、2 行目から先が宙に浮いて設定が読めなくなった。"""
        settings = self.template / "workspace.toml"
        base = settings.read_text(encoding="utf-8").split("[checks]")[0]
        for shape, array in self.ARRAYS.items():
            with self.subTest(shape):
                written = base + "[checks]\n" + array + "near_alignment = false\n"
                settings.write_text(written, encoding="utf-8")
                before = tomllib.loads(written)["checks"]["stale_words"]
                code, said = self._lift("--page", "w1.pptx:1", "--replace", "手法 A=<手法名>")
                self.assertEqual(code, 0, said)
                text = settings.read_text(encoding="utf-8")
                checks = tomllib.loads(text)["checks"]
                self.assertEqual(checks["stale_words"], before + ["<手法名>", "<文言"])
                self.assertIs(checks["near_alignment"], False, "the rest of [checks] was lost")
                if "#" in array:
                    self.assertIn("# the specimen's own", text, "the person's comment was dropped")
                self.assertEqual(text.count("\n") > written.count("\n"), "\n" in array.rstrip("\n"),
                                 "an array on one line was spread over many, or one on many lines was folded")

    def test_stand_in_words_that_would_not_land_in_checks_write_nothing(self) -> None:
        """A `stale_words` of another table comes first: the words would go there, not to the check."""
        settings = self.template / "workspace.toml"
        settings.write_text('[elsewhere]\nstale_words = ["x"]\n\n[checks]\nstale_words = ["案件名"]\n',
                            encoding="utf-8")
        self.before["workspace.toml"] = settings.read_bytes()
        code, said = self._lift("--page", "w1.pptx:1")
        self.assertNotEqual(code, 0)
        self.assertIn("stale_words", said)
        self._untouched()

    def test_a_field_powerpoint_fills_is_left_alone(self) -> None:
        """頁番号の欄は PowerPoint が埋める ― 仮の語にすると、次の案件で人が直せない
        「埋め忘れ」を検査が言い続ける。"""
        from lxml import etree
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(2000000), Emu(500000))
        box.text_frame.paragraphs[0].add_run().text = "頁の文言"
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        paragraph = box.text_frame.paragraphs[0]._p
        paragraph.append(etree.fromstring(
            f'<a:fld xmlns:a="{a}" id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" '
            f'type="slidenum"><a:t>&#8249;#&#8250;</a:t></a:fld>'))
        deck.save(str(self.project / "numbered.pptx"))
        code, said = self._lift("--page", "numbered.pptx:1")
        self.assertEqual(code, 0, said)
        with zipfile.ZipFile(self.template / "specimen.pptx") as archive:
            xml = "".join(archive.read(n).decode("utf-8") for n in archive.namelist()
                          if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        self.assertIn('type="slidenum"><a:t>\u2039#\u203a</a:t>', xml)
        self.assertNotIn("頁の文言", xml)

    def test_the_template_as_it_was_is_kept_to_undo_a_lift(self) -> None:
        self._lift("--page", "w1.pptx:1")
        kept = list((self.template / "_archive").iterdir())
        self.assertEqual(len(kept), 1)
        for name, data in self.before.items():
            self.assertEqual((kept[0] / name).read_bytes(), data, f"{name} was not kept as it was")

    def test_only_the_last_few_copies_are_kept(self) -> None:
        """⚠ **控えは上げるたびに 1 つ増え、消す所が無かった。**テンプレートの folder は案件に配る物で、
        そこに見本の写しが際限なく貯まった。残すのは新しい方から数えた世代だけ ― 取り消したいのは
        直前の上げで、何十回も前の物ではない。"""
        from pptx_agent_maker.project.commands.lift import KEPT

        for _round in range(KEPT + 2):
            code, said = self._lift("--page", "w1.pptx:1")
            self.assertEqual(code, 0, said)
        kept = sorted(path.name for path in (self.template / "_archive").iterdir())
        self.assertEqual(KEPT, len(kept))
        # 残っているのは最後の世代 (= 直前の上げを取り消せる): いちばん新しい控えは、最後の上げの前の見本
        newest = self.template / "_archive" / kept[-1]
        self.assertEqual(len(Presentation(str(newest / "specimen.pptx")).slides) + 1,
                         len(Presentation(str(self.template / "specimen.pptx")).slides))

    def test_what_a_person_put_in_the_archive_is_left_alone(self) -> None:
        mine = self.template / "_archive" / "kept-by-hand"
        mine.mkdir(parents=True)
        (mine / "note.txt").write_text("why this one is kept", encoding="utf-8")
        from pptx_agent_maker.project.commands.lift import KEPT

        for _round in range(KEPT + 1):
            self._lift("--page", "w1.pptx:1")
        self.assertTrue((mine / "note.txt").is_file())
        self.assertEqual(KEPT + 1, len(list((self.template / "_archive").iterdir())))

    # -- a recipe -----------------------------------------------------------

    def test_a_recipe_goes_up_with_its_words_replaced_and_reaches_the_next_project(self) -> None:
        code, said = self._lift("--recipe", "result", "--replace", "手法 A=<手法名>",
                                "--keep", "{method} の結果")
        self.assertEqual(code, 0, said)
        recipe = tomllib.loads((self.template / "recipes.toml").read_text(encoding="utf-8"))
        self.assertEqual(recipe["recipes"]["result"]["footer"], "<手法名> の採点表から")
        self.assertEqual(recipe["recipes"]["result"]["title"], "{method} の結果", "the hole was lost")
        self.assertEqual(recipe["recipes"]["result"]["type"], "figure", "the type name is not words")
        following = create(Path(self.tmp.name) / "next", specimen=self.template)
        self.assertTrue((following / "recipes.toml").is_file(), "init did not hand the recipes on")

    def test_a_part_goes_up_with_its_words_replaced_and_its_choices_and_holes_kept(self) -> None:
        (self.project / "recipes.toml").write_text(
            RECIPES + '\n[parts.figure_of_hand]\nrows = [ { cells = [ { text = "{figure}", role = "display" } ] },'
            ' { cells = [ { card = { heading = "手法 A の数", body = "{meaning}", tone = "accent" } } ] } ]\n',
            encoding="utf-8")
        code, said = self._lift("--part", "figure_of_hand", "--replace", "手法 A=<手法名>")
        self.assertEqual(code, 0, said)
        self.assertIn("lifted part 'figure_of_hand'", said)
        part = tomllib.loads((self.template / "recipes.toml").read_text(encoding="utf-8"))["parts"]["figure_of_hand"]
        (number,), (card,) = part["rows"][0]["cells"], part["rows"][1]["cells"]
        self.assertEqual(number, {"text": "{figure}", "role": "display"}, "a hole and a role are not words")
        self.assertEqual(card["card"]["tone"], "accent")
        self.assertEqual(card["card"]["body"], "{meaning}")
        self.assertIn("<手法名>", card["card"]["heading"])
        following = create(Path(self.tmp.name) / "next", specimen=self.template)
        self.assertIn("figure_of_hand", load_parts(following))

    def test_a_recipe_name_the_template_has_is_refused(self) -> None:
        self._lift("--recipe", "result", "--replace", "手法 A=<手法名>")
        code, said = self._lift("--recipe", "result", "--replace", "手法 A=<手法名>")
        self.assertEqual(code, 1)
        self.assertIn("already has a recipe", said)

    # -- refusals that write nothing ----------------------------------------

    def test_a_replace_that_never_matches_writes_nothing(self) -> None:
        """⚠ 綴り違いの置き換えは、抜いたつもりの語をそのまま残す。"""
        code, said = self._lift("--page", "w1.pptx:1", "--replace", "手法Ａ=<手法名>")
        self.assertEqual(code, 1)
        self.assertIn("never matched", said)
        self._untouched()

    def test_a_page_carrying_a_chart_is_refused(self) -> None:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("s", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0),
                               Emu(3000000), Emu(2000000), data)
        deck.save(str(self.project / "charted.pptx"))
        code, said = self._lift("--page", "charted.pptx:1")
        self.assertEqual(code, 1)
        self.assertIn("chart", said)
        self._untouched()

    def test_a_folder_without_a_template_is_refused(self) -> None:
        empty = Path(self.tmp.name) / "empty"
        empty.mkdir()
        code, said = self._run("lift", str(self.project), str(empty), "--recipe", "result")
        self.assertEqual(code, 1)
        self.assertIn("no specimen.pptx", said)


if __name__ == "__main__":
    unittest.main()
