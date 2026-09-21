"""A person's edit must survive, be visible, and be foldable back in.

⚠ 実際に起きたこと ― 手で直した版を「余計な複製」と判断して消し、**復元できなくした**。
それと、番人の差分が先頭 30 字で切れていたせいで、頁後半の削除が「一致」と出た。
ここで守るのは、消えないことと、差分が全文で取れること。
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx_agent_maker.deck.build import HandEditedError, build  # noqa: E402
from pptx_agent_maker.project import Workspace, create  # noqa: E402
from pptx_agent_maker.project.manifest import Manifest  # noqa: E402
from pptx_agent_maker.review import changes, keep_safe, last_machine_build  # noqa: E402
from test_deck import a_specimen, make_dot  # noqa: E402

MANIFEST = '''
specimen = "base/specimen.pptx"
out = "built.pptx"

[[pages]]
kind = "copy"
page = 1

[[pages]]
kind = "declare"
module = "example_page"
'''


def edit(deck: Path, old: str, new: str) -> None:
    """Stand in for someone editing the deck in PowerPoint."""
    presentation = Presentation(str(deck))
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
    presentation.save(str(deck))


class ReviewTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "project"
        create(self.root)
        make_dot()
        a_specimen(self.root / "base" / "specimen.pptx", ["一枚目"])
        shutil.copy(REPO / "tests" / "data" / "dot.png", self.root / "assets" / "example.png")
        (self.root / "manifests" / "deck.toml").write_text(MANIFEST, encoding="utf-8")
        self.workspace = Workspace.load(self.root)
        self.deck = build(self.workspace, Manifest.load(self.workspace.manifest("deck")))

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_building_twice_is_fine(self) -> None:
        """Same declaration, same deck: nothing to protect."""
        build(self.workspace, Manifest.load(self.workspace.manifest("deck")))

    def test_a_hand_edit_stops_the_next_build(self) -> None:
        edit(self.deck, "一枚目", "人が直した")
        with self.assertRaises(HandEditedError):
            build(self.workspace, Manifest.load(self.workspace.manifest("deck")))

    def test_the_hand_edit_is_shelved_before_anything_else(self) -> None:
        edit(self.deck, "一枚目", "人が直した")
        with self.assertRaises(HandEditedError):
            build(self.workspace, Manifest.load(self.workspace.manifest("deck")))
        shelved = sorted((self.workspace.output / "_edits").glob("*.pptx"))
        self.assertTrue(shelved, "the edit was not kept anywhere")

    def test_the_change_is_visible_as_words(self) -> None:
        reference = last_machine_build(self.deck)
        self.assertIsNotNone(reference, "the machine's own build was not kept")
        edit(self.deck, "一枚目", "人が直した")
        found = changes(reference, self.deck)
        # edit() replaces the word everywhere it appears: the title, the band, the caption.
        self.assertEqual([(c.before, c.after) for c in found],
                         [("一枚目", "人が直した"),
                          ("一枚目 の条件", "人が直した の条件"),
                          ("一枚目 の図", "人が直した の図")])

    def test_a_change_late_on_a_page_is_still_seen(self) -> None:
        """The old diff truncated at 30 characters and called this 'identical'."""
        reference = last_machine_build(self.deck)
        edit(self.deck, "出所", "出所: 差し替えた脚注")
        found = changes(reference, self.deck)
        self.assertTrue(any("差し替えた脚注" in c.after for c in found))

    def test_text_is_read_without_swallowing_the_markup(self) -> None:
        """A self-closing <a:t/> once made the reader eat the XML after it."""
        reference = last_machine_build(self.deck)
        edit(self.deck, "一枚目", "短い差し替え")
        for change in changes(reference, self.deck):
            self.assertNotIn("<a:", change.before + change.after,
                             "markup leaked into what should be words")

    def test_shelving_an_untouched_deck_does_nothing(self) -> None:
        self.assertIsNone(keep_safe(self.deck))


if __name__ == "__main__":
    unittest.main()
