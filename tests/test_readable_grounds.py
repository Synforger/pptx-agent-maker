"""Words on a ground the project coloured stay readable, and a template's own pictures are its own.

⚠ **差し色は案件が変える。**表の見出し行・差し色の箱・線表の期間の帯は字を白に決め打っていたので、
薄い差し色 (= 桃色) の案件では白い字が地に溶け、検査 `contrast` が宣言で組んだ頁で当たった。
⚠ **雛形の絵は雛形の側で直す。**複製した表紙のロゴが粗いと、案件には差し替える元が無い。
"""

from __future__ import annotations

import contextlib
import io
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.layout.base.tokens import contrast  # noqa: E402
from pptx_agent_maker.project import create  # noqa: E402

PALE = "F070A4"


class ReadableGrounds(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "project"
        create(self.root)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _build(self, pages: str, settings: str = "") -> tuple[int, str]:
        if settings:
            path = self.root / "workspace.toml"
            path.write_text(path.read_text(encoding="utf-8") + "\n" + settings, encoding="utf-8")
        (self.root / "w1.toml").write_text(
            'specimen = "specimen.pptx"\nout = "w1.pptx"\n\n' + pages, encoding="utf-8")
        said = io.StringIO()
        with contextlib.redirect_stdout(said), contextlib.redirect_stderr(said):
            return main(["build", str(self.root), "w1"]), said.getvalue()

    def _colours_of(self, words: str) -> set[str]:
        """The colours written on the runs that say `words`, over the whole deck."""
        with zipfile.ZipFile(self.root / "w1.pptx") as deck:
            xml = "".join(deck.read(n).decode() for n in deck.namelist() if n.startswith("ppt/slides/slide"))
        return set(re.findall(r'<a:r><a:rPr[^>]*>(?:(?!</a:rPr>).)*?<a:srgbClr val="(\w{6})"(?:(?!</a:r>).)*?'
                              rf'<a:t>{re.escape(words)}</a:t>', xml, re.S))

    def test_words_on_a_pale_accent_take_the_colour_that_reads(self) -> None:
        pages = '''
[[pages]]
kind = "declare"
type = "board"
title = "見出し行の地は差し色"
table = [["日付", "媒体"], ["10/16", "紙"]]
footer = "出所"

[[pages]]
kind = "declare"
type = "cards"
title = "差し色の箱"
cards = [{ heading = "強い札", body = "本文", tone = "accent" }]

[[pages]]
kind = "declare"
type = "timeline"
title = "期間の帯は差し色"
periods = ["上期", "下期"]
lanes = [{ name = "担当", bars = [{ text = "作業", from = 1, to = 2 }] }]
'''
        code, said = self._build(pages, f'[theme.palette]\naccent = "{PALE}"\n')
        self.assertEqual(0, code, said)
        for words in ("日付", "強い札", "上期"):
            (colour,) = self._colours_of(words)
            self.assertGreaterEqual(contrast(colour, PALE), 4.5, f"{words} is #{colour} on #{PALE}")

    def test_the_default_dark_accent_still_carries_white_words(self) -> None:
        code, said = self._build('[[pages]]\nkind = "declare"\ntype = "board"\ntitle = "表"\n'
                                 'table = [["日付", "媒体"], ["10/16", "紙"]]\nfooter = "出所"\n')
        self.assertEqual(0, code, said)
        self.assertEqual({"FFFFFF"}, self._colours_of("日付"))

    def _logo(self, name: str, colour: tuple[int, int, int] = (200, 40, 90)) -> Path:
        image = self.root / "assets" / name
        Image.new("RGB", (120, 24), colour).save(image)
        return image

    def test_a_template_picture_is_left_to_the_template(self) -> None:
        """複製した表紙のロゴは雛形の物 (= 案件は直せない)。案件が置いた同じ粗さの絵は言う。"""
        logo = self._logo("logo.png")
        specimen = Presentation(str(self.root / "specimen.pptx"))
        specimen.slides[0].shapes.add_picture(str(logo), Emu(500000), Emu(500000), Emu(3000000), Emu(600000))
        specimen.save(str(self.root / "specimen.pptx"))
        # 表紙の文言は残したまま複製する (= ここで見るのは絵だけ。雛形の語の残りの検査は外す)
        settings = self.root / "workspace.toml"
        settings.write_text(re.sub(r"(?m)^stale_words = .*$", "stale_words = []",
                                   settings.read_text(encoding="utf-8")), encoding="utf-8")
        code, said = self._build('[[pages]]\nkind = "copy"\npage = 1\n')
        self.assertEqual(0, code, said)
        self._logo("ours.png", (40, 90, 200))  # 雛形と違う絵 (= 同じ中身なら雛形の絵とみなす)
        code, said = self._build('[[pages]]\nkind = "declare"\ntype = "figure"\ntitle = "案件の絵"\n'
                                 'figure = "ours.png"\nfooter = "出所"\n')
        self.assertEqual(1, code, said)
        self.assertRegex(said, r"FAIL\s+low_resolution")


if __name__ == "__main__":
    unittest.main()
