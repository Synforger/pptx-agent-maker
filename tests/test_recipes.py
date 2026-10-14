"""Recipes: a page written the same way twice, kept once.

⚠ **上げる口が重いと、誰も上げない。**前の世代では昇格が 1 度も起きず、毎回複製して
少し直す形が溜まった。ここは上げる操作が 1 回で済み、上げた後に古いコピーが残らず、
**上げる前と後で焼いた頁が変わらない**ことを機械が見る。
"""

from __future__ import annotations

import io
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.project import Workspace, create  # noqa: E402
from pptx_agent_maker.project.files import recipes  # noqa: E402
from pptx_agent_maker.project.files.manifest import Manifest, ManifestError  # noqa: E402
from pptx_agent_maker.project.files.recipes import RecipeError, expand  # noqa: E402

DATA = Path(__file__).resolve().parent / "data"

RECIPES = {"result": {"type": "figure", "title": "{method} の結果",
                      "condition": "同じ入力、同じ条件"}}


class ExpandingARecipe(unittest.TestCase):
    """A page calling a recipe becomes an ordinary declaration."""

    def test_the_page_gives_only_what_changes(self) -> None:
        page = {"kind": "recipe", "recipe": "result", "fill": {"method": "A"},
                "figure": "a.png"}
        self.assertEqual(expand("p", page, RECIPES),
                         {"type": "figure", "title": "A の結果",
                          "condition": "同じ入力、同じ条件", "figure": "a.png"})

    def test_a_page_cannot_change_what_the_recipe_fixes(self) -> None:
        """⚠ 書き換えられると、同じ recipe の頁が週ごとに少しずつ違う形になる。"""
        page = {"recipe": "result", "fill": {"method": "A"}, "figure": "a.png",
                "condition": "今回だけ別の条件"}
        with self.assertRaises(RecipeError) as raised:
            expand("p", page, RECIPES)
        self.assertIn("already fixes condition", str(raised.exception))

    def test_a_hole_left_open_is_refused(self) -> None:
        with self.assertRaises(RecipeError) as raised:
            expand("p", {"recipe": "result", "figure": "a.png"}, RECIPES)
        self.assertIn("{method}", str(raised.exception))

    def test_a_fill_nobody_uses_is_refused(self) -> None:
        page = {"recipe": "result", "fill": {"method": "A", "methd": "typo"}, "figure": "a.png"}
        with self.assertRaises(RecipeError) as raised:
            expand("p", page, RECIPES)
        self.assertIn("methd", str(raised.exception))

    def test_an_unknown_recipe_names_the_ones_there_are(self) -> None:
        with self.assertRaises(RecipeError) as raised:
            expand("p", {"recipe": "reslt"}, RECIPES)
        self.assertIn("result", str(raised.exception))


W1 = '''specimen = "specimen.pptx"
out = "w1.pptx"

[[pages]]
kind = "copy"
page = 1
replace = [["案件名", "見本"], ["第 N 回 進捗報告", "第 1 回"]]

[[pages]]
kind = "declare"
type = "figure"
title = "手法 A の結果"
condition = "同じ入力、同じ条件"
figure = "dot.png"
conclusion = "A はこう読める"
footer = "採点表から"
why = "毎週の結果頁"

# 次の頁の前の注記 (= 書き換えで消えてはいけない)
[[pages]]
kind = "declare"
type = "cards"
title = "まとめ"
cards = [["一", "いち"], ["二", "に"]]
'''

W2 = W1.replace('out = "w1.pptx"', 'out = "w2.pptx"') \
       .replace("手法 A の結果", "手法 B の結果") \
       .replace("A はこう読める", "B はこう読める") \
       .replace('figure = "dot.png"', 'figure = "wide.png"')


class PromotingThroughTheCommandLine(unittest.TestCase):
    """`promote` once: the recipe is written, the pages call it, the decks do not change."""

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "project"
        create(self.root)
        for name in ("dot.png", "wide.png"):
            shutil.copy(DATA / name, self.root / "assets" / name)
        (self.root / "w1.toml").write_text(W1, encoding="utf-8")
        (self.root / "w2.toml").write_text(W2, encoding="utf-8")
        # 新しい案件は見本の部品を持つ recipe の file で始まる (= 拒んだ promote はそれを書き換えない)
        self.recipes_before = (self.root / "recipes.toml").read_text(encoding="utf-8")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _run(self, *argv: str) -> int:
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            return main(list(argv))

    def _texts(self, deck: str) -> list[list[str]]:
        slides = Presentation(str(self.root / deck)).slides
        return [sorted(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
                for slide in slides]

    def test_the_decks_read_the_same_before_and_after(self) -> None:
        for name in ("w1", "w2"):
            self.assertEqual(0, self._run("build", str(self.root), name, "--skip-checks"))
        before = {name: self._texts(f"{name}.pptx") for name in ("w1", "w2")}
        for name in ("w1", "w2"):
            (self.root / f"{name}.pptx").unlink()  # 焼き直す (= 手編集の検出を踏まない)

        self.assertEqual(0, self._run("promote", str(self.root), "result", "w1:2", "w2:2"))

        for name in ("w1", "w2"):
            self.assertEqual(0, self._run("build", str(self.root), name, "--skip-checks"))
            self.assertEqual(before[name], self._texts(f"{name}.pptx"), name)

    def test_what_is_shared_goes_up_and_what_differs_stays(self) -> None:
        self.assertEqual(0, self._run("promote", str(self.root), "result", "w1:2", "w2:2"))
        recipe = recipes.load(self.root)["result"]
        self.assertEqual(recipe, {"type": "figure", "condition": "同じ入力、同じ条件",
                                  "footer": "採点表から"})
        self.assertEqual(list(recipe)[0], "type", "the recipe should open with its type")
        text = (self.root / "w1.toml").read_text(encoding="utf-8")
        self.assertIn('kind = "recipe"', text)
        self.assertNotIn("同じ入力、同じ条件", text, "the old copy stayed in the manifest")
        self.assertIn('why = "毎週の結果頁"', text, "the page's own note went up with the shape")
        self.assertIn("# 次の頁の前の注記", text, "a comment before the next page was lost")

    def test_the_recipes_file_is_not_taken_for_a_manifest(self) -> None:
        self._run("promote", str(self.root), "result", "w1:2", "w2:2")
        names = [p.name for p in Workspace.load(self.root).manifests()]
        self.assertNotIn("recipes.toml", names)

    def test_pages_of_different_types_are_not_one_shape(self) -> None:
        self.assertEqual(1, self._run("promote", str(self.root), "mixed", "w1:2", "w1:3"))
        self.assertEqual((self.root / "recipes.toml").read_text(encoding="utf-8"), self.recipes_before)
        self.assertEqual(W1, (self.root / "w1.toml").read_text(encoding="utf-8"))

    def test_one_page_is_not_yet_a_shape(self) -> None:
        self.assertEqual(1, self._run("promote", str(self.root), "result", "w1:2"))
        self.assertEqual((self.root / "recipes.toml").read_text(encoding="utf-8"), self.recipes_before)

    def test_a_copied_page_has_nothing_to_share(self) -> None:
        self.assertEqual(1, self._run("promote", str(self.root), "cover", "w1:1", "w2:1"))
        self.assertEqual((self.root / "recipes.toml").read_text(encoding="utf-8"), self.recipes_before)

    def test_a_rewrite_that_would_change_a_page_writes_nothing(self) -> None:
        """⚠ 書き換えの検証が効いていることを、わざと壊した書き換えで確かめる。"""
        dropping = lambda page: recipes.page_block(  # noqa: E731
            {k: v for k, v in page.items() if k != "conclusion"})
        with mock.patch.object(recipes, "page_block", side_effect=dropping):
            self.assertEqual(1, self._run("promote", str(self.root), "result", "w1:2", "w2:2"))
        self.assertEqual((self.root / "recipes.toml").read_text(encoding="utf-8"), self.recipes_before)
        self.assertEqual(W1, (self.root / "w1.toml").read_text(encoding="utf-8"))
        self.assertEqual(W2, (self.root / "w2.toml").read_text(encoding="utf-8"))

    def test_a_manifest_calling_a_missing_recipe_says_so(self) -> None:
        (self.root / "w3.toml").write_text(
            'specimen = "specimen.pptx"\nout = "w3.pptx"\n\n[[pages]]\nkind = "recipe"\n'
            'recipe = "result"\n', encoding="utf-8")
        with self.assertRaises(ManifestError) as raised:
            Manifest.load(self.root / "w3.toml")
        self.assertIn("no recipe named 'result'", str(raised.exception))


if __name__ == "__main__":
    unittest.main()


# -- parts (= one compose cell kept once) --------------------------------------

from pptx_agent_maker.project.files.recipes import expand_parts, load_parts  # noqa: E402

PARTS = {"big_number": {"rows": [{"cells": [{"text": "{figure}", "role": "display"}]},
                                 {"cells": [{"text": "{meaning}"}]}]},
         "loop": {"rows": [{"cells": [{"part": "loop"}]}]}}


def _composed(*cells):
    return {"type": "compose", "title": "t", "rows": [{"cells": list(cells)}]}


class ExpandingAPart(unittest.TestCase):
    """A cell naming a part becomes the cell the part writes, its holes filled."""

    def test_the_cell_gives_only_what_changes_and_keeps_where_it_stands(self) -> None:
        page = _composed({"part": "big_number", "fill": {"figure": "98%", "meaning": "own a phone"},
                          "weight": 2, "id": "n"})
        (cell,) = expand_parts("p", page, PARTS)["rows"][0]["cells"]
        self.assertEqual(cell, {"rows": [{"cells": [{"text": "98%", "role": "display"}]},
                                         {"cells": [{"text": "own a phone"}]}], "weight": 2, "id": "n"})

    def test_cells_inside_rows_are_written_out_too(self) -> None:
        page = _composed({"rows": [{"cells": [{"part": "big_number", "fill": {"figure": "1", "meaning": "m"}}]}]})
        inner = expand_parts("p", page, PARTS)["rows"][0]["cells"][0]["rows"][0]["cells"][0]
        self.assertIn("rows", inner)
        self.assertNotIn("part", inner)

    def test_a_cell_cannot_change_what_its_part_fixes(self) -> None:
        with self.assertRaises(RecipeError) as raised:
            expand_parts("p", _composed({"part": "big_number", "fill": {"figure": "1", "meaning": "m"},
                                         "rows": []}), PARTS)
        self.assertIn("already fixes rows", str(raised.exception))

    def test_a_hole_left_open_a_fill_nobody_uses_and_an_unknown_part_are_refused(self) -> None:
        with self.assertRaises(RecipeError) as open_hole:
            expand_parts("p", _composed({"part": "big_number", "fill": {"figure": "1"}}), PARTS)
        self.assertIn("{meaning}", str(open_hole.exception))
        with self.assertRaises(RecipeError) as unused:
            expand_parts("p", _composed({"part": "big_number", "fill": {"figure": "1", "meaning": "m", "x": "y"}}),
                         PARTS)
        self.assertIn("names x", str(unused.exception))
        with self.assertRaises(RecipeError) as unknown:
            expand_parts("p", _composed({"part": "nope"}), PARTS)
        self.assertIn("big_number", str(unknown.exception))
        with self.assertRaises(RecipeError) as stray:
            expand_parts("p", _composed({"text": "x", "fill": {"a": "b"}}), PARTS)
        self.assertIn("`fill` belongs to a cell that names a `part`", str(stray.exception))

    def test_a_part_that_calls_itself_is_refused(self) -> None:
        with self.assertRaises(RecipeError) as raised:
            expand_parts("p", _composed({"part": "loop"}), PARTS)
        self.assertIn("calling itself", str(raised.exception))

    def test_a_part_carrying_where_its_cell_stands_is_refused_when_the_file_is_read(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "recipes.toml").write_text('[parts.wide]\ntext = "x"\nweight = 2\n', encoding="utf-8")
            with self.assertRaises(RecipeError) as raised:
                load_parts(tmp)
        self.assertIn("carries weight", str(raised.exception))


class TheSampleParts(unittest.TestCase):
    """Every part a new project is handed builds, with its holes filled, and the checks find nothing."""

    def test_each_sample_part_builds_on_a_page_of_its_own(self) -> None:
        from pptx_agent_maker import checks
        from pptx_agent_maker.layout import types
        from pptx_agent_maker.project.commands.scaffold import TEMPLATE
        from pptx_agent_maker.project.files.recipes import HOLE
        from pptx_agent_maker.write import add_page, aspect, new_deck, save

        parts = load_parts(TEMPLATE)
        self.assertEqual(len(parts), 9)
        deck = new_deck()
        for name, part in parts.items():
            with self.subTest(name):
                holes = sorted(set(HOLE.findall(repr(part))))
                fill = {hole: ("dot.png" if hole in ("picture", "map") else f"Words for {hole}") for hole in holes}
                cell = {"part": name, "fill": fill}
                if name in ("named_picture", "map"):
                    cell["marks"] = [{"at": [0.3, 0.4], "text": "Here"}]
                page = expand_parts(name, _composed(cell), parts)
                add_page(deck, types.build(page, lambda asset: DATA / asset, aspect).build())
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual([], checks.run_all(save(deck, Path(tmp) / "parts.pptx")))

    def test_a_project_built_from_a_template_with_recipes_keeps_the_samples_it_lacks(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            seed = create(Path(tmp) / "seed")
            template = Path(tmp) / "template"
            template.mkdir()
            for name in ("specimen.pptx", "workspace.toml"):
                shutil.copy(seed / name, template / name)
            (template / "recipes.toml").write_text(
                '[recipes.result]\ntype = "figure"\ntitle = "x"\n\n[parts.quote]\ntext = "the company\'s own"\n',
                encoding="utf-8")
            project = create(Path(tmp) / "project", specimen=template)
            parts = load_parts(project)
            self.assertEqual(parts["quote"], {"text": "the company's own"}, "the template's part wins")
            self.assertIn("big_number", parts, "a sample the template lacks is added")
            self.assertIn("result", recipes.load(project))

    def test_a_manifest_page_with_a_part_builds_through_the_command_line(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            project = create(Path(tmp) / "p")
            (project / "w1.toml").write_text(
                'specimen = "specimen.pptx"\nout = "w1.pptx"\n\n[[pages]]\nkind = "declare"\ntype = "compose"\n'
                'title = "One in four"\n\n[[pages.rows]]\n[[pages.rows.cells]]\npart = "big_number"\n'
                'fill = { figure = "25%", meaning = "of the work is waiting" }\n', encoding="utf-8")
            entries = Manifest.load(project / "w1.toml").entries
        (cell,) = entries[0].data["rows"][0]["cells"]
        self.assertEqual(cell["rows"][0]["cells"][0], {"text": "25%", "role": "display"})
