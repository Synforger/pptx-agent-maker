"""A template whose first layout is a page of content: a title placeholder in a colour, and a logo.

同梱の見本の 1 枚目のレイアウトは表紙 (= 真ん中の題)。案件のテンプレートは、1 枚目が本文の頁で、題の枠に
色が付き、隅にロゴが在ることが多い。その形を、同梱の見本のレイアウトを書き換えて作る (= 絵の file を
持たずに済むよう、ロゴは塗った矩形で置く。重なりの検査は絵と同じに数える)。
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))

from pptx_agent_maker.layout.base.geometry import Rect  # noqa: E402

SHIPPED = REPO / "src" / "pptx_agent_maker" / "templates" / "project" / "specimen.pptx"
LAYOUT = "ppt/slideLayouts/slideLayout1.xml"
PAGE = "ppt/slides/slide1.xml"
#: 同梱の見本のレイアウトが、頁番号の枠に付けている番号
NUMBER_IDX = 12
#: 頁番号を持つ頁が、その枠を持つ書き方 (= PowerPoint で「スライド番号」を入れた頁)
NUMBER_ON_A_PAGE = (
    '<p:sp><p:nvSpPr><p:cNvPr id="95" name="Slide Number Placeholder 1"/><p:cNvSpPr><a:spLocks noGrp="1"/>'
    f'</p:cNvSpPr><p:nvPr><p:ph type="sldNum" sz="quarter" idx="{NUMBER_IDX}"/></p:nvPr></p:nvSpPr><p:spPr/>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:fld id="{C1FF6DA9-008F-8B48-92A6-B652298478BF}" type="slidenum">'
    '<a:rPr lang="en-US"/><a:t>1</a:t></a:fld><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')
#: 題の枠の字の色 (= テーマの色のどれでもない青。継いだことが色で分かる)
TITLE_COLOUR = "0070C0"


def with_a_cover(template: Path, *, counts_from: int | None = None) -> Path:
    """Put a page with no number in front of the template's page (= a cover), in place.

    表紙は、番号を出さない頁として作られる (= 頁番号の枠を持たない)。python-pptx は、レイアウトから頁を
    作るとき日付・フッター・頁番号の枠を写さないので、足した頁はそのまま「番号を持たない頁」になる。

    `counts_from` は、テンプレートが数え始めを自分で書いている形 (= PowerPoint の「スライド開始番号」)。
    """
    from pptx import Presentation

    deck = Presentation(str(template))
    deck.slides.add_slide(deck.slide_layouts[6])
    listed = deck.slides._sldIdLst
    listed.insert(0, listed[-1])
    if counts_from is not None:
        deck.part._element.set("firstSlideNum", str(counts_from))
    deck.save(str(template))
    return template


def _line(across: Rect, how: str) -> str:
    """A line a layout draws, the two ways PowerPoint writes one: a connector, or a shape with no height."""
    place = (f'<a:xfrm><a:off x="{across.left}" y="{across.top}"/><a:ext cx="{across.width}" cy="0"/></a:xfrm>'
             '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>')
    stroke = {"connector": '<a:ln w="12700"><a:solidFill><a:srgbClr val="7F7F7F"/></a:solidFill></a:ln>',
              "flat": '<a:ln w="12700"><a:solidFill><a:srgbClr val="7F7F7F"/></a:solidFill></a:ln>',
              "unseen": "<a:ln><a:noFill/></a:ln>"}[how]
    if how == "connector":
        return ('<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="91" name="Line"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
                f"<p:spPr>{place}{stroke}</p:spPr></p:cxnSp>")
    return ('<p:sp><p:nvSpPr><p:cNvPr id="91" name="Line"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f"<p:spPr>{place}{stroke}</p:spPr></p:sp>")


def a_template(destination: Path, *, title: bool = True, logo: Rect | None = None,
               line: Rect | None = None, line_as: str = "connector",
               numbered: bool = False, number_frame: int | None = NUMBER_IDX) -> Path:
    """The shipped specimen, with its first layout made a page of content.

    `line` は、レイアウトが頁に描く線 (= 題の下の罫線など。高さは読まない)。`line_as` はその書かれ方 ―
    `connector` (= コネクタ) / `flat` (= 高さの無い図形) / `unseen` (= 線を引かない図形)。

    `numbered` は、見本の頁が頁番号を持つか (= 頁に番号を振る資料)。`number_frame` は、レイアウトの
    頁番号の枠の番号 (= `idx`。None なら、レイアウトから枠ごと抜く)。
    """
    with zipfile.ZipFile(SHIPPED) as source, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == PAGE and numbered:
                xml = data.decode("utf-8")
                assert xml.count("</p:spTree>") == 1
                data = xml.replace("</p:spTree>", NUMBER_ON_A_PAGE + "</p:spTree>").encode("utf-8")
            if item.filename == LAYOUT:
                xml = data.decode("utf-8")
                offered = f'<p:ph type="sldNum" sz="quarter" idx="{NUMBER_IDX}"/>'
                assert xml.count(offered) == 1
                if number_frame is None:
                    start = xml.rindex("<p:sp>", 0, xml.index(offered))
                    xml = xml[:start] + xml[xml.index("</p:sp>", start) + len("</p:sp>"):]
                else:
                    xml = xml.replace(offered, offered.replace(str(NUMBER_IDX), str(number_frame)))
                if title:
                    assert xml.count('<p:ph type="ctrTitle"/>') == 1
                    xml = xml.replace('<p:ph type="ctrTitle"/>', '<p:ph type="title"/>')
                    xml = xml.replace(
                        "<a:lstStyle/>",
                        '<a:lstStyle><a:lvl1pPr algn="ctr"><a:defRPr sz="2500" b="0"><a:solidFill>'
                        f'<a:srgbClr val="{TITLE_COLOUR}"/></a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle>', 1)
                if logo is not None:
                    xml = xml.replace("</p:spTree>", (
                        '<p:sp><p:nvSpPr><p:cNvPr id="90" name="Logo"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
                        f'<a:xfrm><a:off x="{logo.left}" y="{logo.top}"/><a:ext cx="{logo.width}" cy="{logo.height}"/>'
                        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill>'
                        '<a:srgbClr val="7F7F7F"/></a:solidFill></p:spPr></p:sp></p:spTree>'))
                if line is not None:
                    xml = xml.replace("</p:spTree>", _line(line, line_as) + "</p:spTree>")
                data = xml.encode("utf-8")
            out.writestr(item, data)
    return destination
