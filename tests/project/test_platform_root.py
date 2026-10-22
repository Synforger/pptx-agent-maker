"""Where the tool keeps its files: under `AGENT_PLATFORM_ROOT` when one is given, as before when not.

器の根 (= 環境変数 `AGENT_PLATFORM_ROOT`) を渡された時、道具は設定・見本集・焼いた頁を
`<根>/<役割>/pptx-agent-maker/` に置き、根の外には何も書かない。渡されない時は今までの置き場
(= `~/.config`・`~/.local/share`・`~/.cache`) のまま。後者は「今の挙動を変えない」の約束で、
どちらも置き場を引く経路を実際に走らせて確かめる。
"""

from __future__ import annotations

import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "preview" / "src"))

from pptx import Presentation  # noqa: E402
from pptx_agent_maker.project.files.places import PlacesError, default_places_file, load  # noqa: E402
from pptx_live_preview.cache import cache_root  # noqa: E402
from pptx_live_preview.platform_root import ROOT_VARIABLE, tool_folder  # noqa: E402
from pptx_live_preview.render import page_count, to_pdf, xdg_environment  # noqa: E402

SCENARIO = REPO / "tests" / "support" / "platform_root_scenario.py"
TOOL = "pptx-agent-maker"
#: 置き場を決める環境変数 (= 試験の外の機械の値が混ざらないよう、毎回消してから始める)
VARIABLES = (ROOT_VARIABLE, "XDG_CACHE_HOME")


def _examples_folder() -> Path:
    spec = importlib.util.spec_from_file_location(
        "fetch_examples_for_root", REPO / "scripts" / "examples" / "fetch-examples.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.default_examples_folder()


class Machine(unittest.TestCase):
    """A machine with an empty home, and none of the variables that decide where files go."""

    def setUp(self) -> None:
        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        self.dir = Path(scratch.name).resolve()
        self.home = self.dir / "home"
        self.home.mkdir()
        environ = {name: value for name, value in os.environ.items() if name not in VARIABLES}
        environ["HOME"] = str(self.home)
        patch = mock.patch.dict(os.environ, environ, clear=True)
        patch.start()
        self.addCleanup(patch.stop)


class WithoutARoot(Machine):
    def test_every_place_is_where_it_was_before_the_root_existed(self) -> None:
        self.assertIsNone(tool_folder("config"))
        self.assertEqual(self.home / ".config" / TOOL / "preview.toml", default_places_file())
        self.assertEqual(self.home / ".cache" / "pptx-live-preview", cache_root())
        self.assertEqual(self.home / ".local" / "share" / TOOL / "examples", _examples_folder())

    def test_an_empty_root_is_no_root(self) -> None:
        os.environ[ROOT_VARIABLE] = ""
        self.assertIsNone(tool_folder("cache"))
        self.assertEqual(self.home / ".cache" / "pptx-live-preview", cache_root())


class WithARoot(Machine):
    def setUp(self) -> None:
        super().setUp()
        self.root = self.dir / "platform"
        os.environ[ROOT_VARIABLE] = str(self.root)

    def test_each_place_is_in_the_folder_of_its_role(self) -> None:
        self.assertEqual(self.root / "config" / TOOL / "preview.toml", default_places_file())
        self.assertEqual(self.root / "cache" / TOOL / "pages", cache_root())
        self.assertEqual(self.root / "references" / TOOL / "examples", _examples_folder())
        for role in ("config", "secrets", "state", "references", "cache"):
            self.assertEqual(self.root / role / TOOL, tool_folder(role))

    def test_a_role_the_platform_does_not_have_is_refused(self) -> None:
        with self.assertRaises(ValueError):
            tool_folder("logs")

    def test_a_variable_that_already_named_a_place_is_stronger_than_the_root(self) -> None:
        os.environ["XDG_CACHE_HOME"] = str(self.dir / "xdg")
        self.assertEqual(self.dir / "xdg" / "pptx-live-preview", cache_root())

    def test_a_root_that_depends_on_the_folder_it_was_started_in_is_refused(self) -> None:
        os.environ[ROOT_VARIABLE] = "platform"
        with self.assertRaises(ValueError):
            tool_folder("cache")
        with self.assertRaises(PlacesError):
            load()

    def test_the_places_file_is_read_from_the_root(self) -> None:
        folder = self.root / "config" / TOOL
        folder.mkdir(parents=True)
        (folder / "preview.toml").write_text('[[folder]]\nname = "a"\npath = "~/decks"\n', encoding="utf-8")
        self.assertEqual([{"name": "a", "path": "~/decks"}], load()["folder"])


XDG_VARIABLES = ("XDG_CACHE_HOME", "XDG_CONFIG_HOME", "XDG_DATA_HOME")


class TheCommandsItStarts(Machine):
    """LibreOffice and poppler write to $HOME/.cache and $HOME/.config on Linux unless told where else."""

    def test_with_a_root_their_cache_and_settings_go_under_the_tools_cache(self) -> None:
        os.environ[ROOT_VARIABLE] = str(self.dir / "platform")
        home = cache_root() / "_poppler-xdg"
        environment = xdg_environment(home)
        for variable, leaf in zip(XDG_VARIABLES, ("cache", "config", "data")):
            self.assertEqual(str(home / leaf), environment[variable])
            self.assertTrue((home / leaf).is_dir())

    def test_without_a_root_they_inherit_the_environment_as_before(self) -> None:
        self.assertIsNone(xdg_environment(self.dir / "xdg"))
        self.assertFalse((self.dir / "xdg").exists())

    def test_every_command_the_preview_starts_is_given_that_environment(self) -> None:
        os.environ[ROOT_VARIABLE] = str(self.dir / "platform")
        deck, pdf = self.dir / "a.pptx", self.dir / "a.pdf"
        deck.write_bytes(b"not a zip: the conversion is stubbed, the deck is only passed along")
        pdf.write_bytes(b"%PDF")
        asked = []

        def run(command, **options):
            asked.append((Path(command[0]).name, options.get("env")))
            return subprocess.CompletedProcess(command, 0, stdout="Pages: 1\n", stderr="")

        profile = cache_root() / "_libreoffice" / "profile-0"
        with mock.patch("pptx_live_preview.render.subprocess.run", run), \
                mock.patch("pptx_live_preview.render.soffice_bin", return_value="soffice"), \
                mock.patch("pptx_live_preview.render.pdfinfo_bin", return_value="pdfinfo"):
            (self.dir / "work").mkdir()
            (self.dir / "work" / "a.pdf").write_bytes(b"%PDF")
            to_pdf(deck, self.dir / "work", profile)
            page_count(pdf)
        self.assertEqual(["soffice", "pdfinfo"], [name for name, _ in asked])
        for name, environment in asked:
            for variable in XDG_VARIABLES:
                self.assertTrue(Path(environment[variable]).is_relative_to(cache_root()), (name, variable))
        # 変換は 1 つのプロファイルに 1 本。設定と cache もそのプロファイルのものだけ
        self.assertEqual(str(cache_root() / "_libreoffice" / "profile-0-xdg" / "cache"),
                         asked[0][1]["XDG_CACHE_HOME"])


#: Linux の LibreOffice と poppler の真似 (= XDG の変数が無ければ HOME の `.cache` と `.config` に書く)
LINUX_LIKE = """#!/bin/sh
# Writes where Linux would when XDG_* is not set, then runs the real {name}.
mkdir -p "${{XDG_CACHE_HOME:-$HOME/.cache}}/fontconfig" "${{XDG_CONFIG_HOME:-$HOME/.config}}"
exec "{real}" "$@"
"""


@unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftoppm"),
                     "bakes a deck for real: needs LibreOffice (soffice) and poppler (pdftoppm)")
class NothingIsWrittenOutsideTheRoot(Machine):
    """Bake a deck, remember what was opened and gather the examples, and see where the files went."""

    def setUp(self) -> None:
        super().setUp()
        self.tmp = self.dir / "tmp"
        self.cwd = self.dir / "cwd"
        self.tmp.mkdir()
        self.cwd.mkdir()
        self.bin = self.dir / "bin"
        self.bin.mkdir()
        for name in ("soffice", "pdftoppm", "pdfinfo"):
            wrapper = self.bin / name
            wrapper.write_text(LINUX_LIKE.format(name=name, real=shutil.which(name)), encoding="utf-8")
            wrapper.chmod(0o755)
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "a page"
        self.deck = self.dir / "work" / "a.pptx"
        self.deck.parent.mkdir()
        deck.save(self.deck)

    def prepare(self, config: Path, examples: Path) -> None:
        """What the machine already has: a places file and the collection's list of decks."""
        config.mkdir(parents=True)
        (config / "preview.toml").write_text('[[folder]]\nname = "a"\npath = "~/decks"\n', encoding="utf-8")
        (examples / "slides-align").mkdir(parents=True)
        (examples / "slides-align" / "decks.json").write_text("[]", encoding="utf-8")

    def run_scenario(self, root: Path | None) -> None:
        environ = {"PATH": f"{self.bin}{os.pathsep}{os.environ.get('PATH', '')}", "HOME": str(self.home), "TMPDIR": str(self.tmp)}
        if root is not None:
            environ[ROOT_VARIABLE] = str(root)
        result = subprocess.run([sys.executable, str(SCENARIO), str(self.deck)], env=environ,
                                cwd=self.cwd, capture_output=True, text=True)
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)

    @staticmethod
    def entries(folder: Path) -> list[str]:
        return sorted(str(path.relative_to(folder)) for path in folder.rglob("*"))

    def test_with_a_root_every_file_is_under_it(self) -> None:
        root = self.dir / "platform"
        self.prepare(root / "config" / TOOL, root / "references" / TOOL / "examples")
        before = self.entries(self.home)

        self.run_scenario(root)

        self.assertEqual(before, self.entries(self.home), "the home folder was written to")
        self.assertEqual([], self.entries(self.tmp), "work files were left in the temporary folder")
        self.assertEqual([], self.entries(self.cwd), "the folder it was started in was written to")
        pages = root / "cache" / TOOL / "pages"
        self.assertTrue((pages / "_opened.json").is_file())
        self.assertTrue(list(pages.glob("*/render.json")), "the baked pages are not in the root's cache")
        self.assertTrue(list(pages.glob("_libreoffice/profile-*")), "LibreOffice's profile is not in the root's cache")
        self.assertTrue((root / "references" / TOOL / "examples" / "INDEX.md").is_file())

    def test_without_a_root_the_files_are_where_they_were(self) -> None:
        self.prepare(self.home / ".config" / TOOL, self.home / ".local" / "share" / TOOL / "examples")

        self.run_scenario(None)

        self.assertEqual([], self.entries(self.tmp), "work files were left in the temporary folder")
        self.assertEqual([], self.entries(self.cwd), "the folder it was started in was written to")
        pages = self.home / ".cache" / "pptx-live-preview"
        self.assertTrue((pages / "_opened.json").is_file())
        self.assertTrue(list(pages.glob("*/render.json")))
        self.assertTrue((self.home / ".local" / "share" / TOOL / "examples" / "INDEX.md").is_file())


if __name__ == "__main__":
    unittest.main()
