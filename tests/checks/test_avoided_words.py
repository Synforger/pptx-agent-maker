"""A word the project has said not to use is stopped wherever it stands on a page.

⚠ **表に書いた語は、見た目の run の境にも、グラフの中にも逃げ場が無い。**太字の run で 1 語が割れていても、
グラフの項目名・軸の名前・系列名の中でも当たる。テンプレートと案件の表は鍵ごとに重なり、案件の `false` で外れる。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from checks.test_checks import raw_deck, text_at  # noqa: E402
from pptx_agent_maker._common.slide import read  # noqa: E402
from pptx_agent_maker.__main__ import _checks  # noqa: E402
from pptx_agent_maker.checks import CHECKS, run_all  # noqa: E402
from pptx_agent_maker.checks.rules import avoided_words  # noqa: E402
from pptx_agent_maker.project.files.workspace import Workspace  # noqa: E402

AVOID = {"avoid": {"鏡の写し": "レプリカ", "二重の殻": "ラッパー"}}


def chart_at(slide, *, title=None, axis=None, category_axis=None, categories=("項目 A", "項目 B"),
             series="系列 X", top=Emu(500000)) -> None:
    """A column chart whose words are the ones given (= no title or axis name unless asked for)."""
    data = CategoryChartData()
    data.categories = list(categories)
    data.add_series(series, [1, 2])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(500000), top, Emu(4000000), Emu(2000000),
                                   data).chart
    chart.has_title = title is not None
    if title is not None:
        chart.chart_title.text_frame.text = title
    for axis_of, words in ((chart.value_axis, axis), (chart.category_axis, category_axis)):
        axis_of.has_title = words is not None
        if words is not None:
            axis_of.axis_title.text_frame.text = words


class Fixtures(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def deck(self, *pages, name: str = "deck.pptx") -> Path:
        """A deck of one page for each function in `pages` (= each places its page's shapes)."""
        from pptx import Presentation
        from pptx_agent_maker import DEFAULT

        deck = Presentation()
        deck.slide_width = Emu(DEFAULT.slide.width)
        deck.slide_height = Emu(DEFAULT.slide.height)
        for place in pages:
            place(deck.slides.add_slide(deck.slide_layouts[6]))
        deck.save(str(self.dir / name))
        return self.dir / name

    def said(self, *pages, config=AVOID) -> list[tuple[int, str]]:
        return [(f.page, f.why) for f in avoided_words.run(self.deck(*pages), config)]


class TheWordsOnAPage(Fixtures):
    def test_a_word_on_the_list_stops_the_page_it_stands_on(self) -> None:
        found = avoided_words.run(
            self.deck(lambda s: text_at(s, "ふつうの文"), lambda s: text_at(s, "この鏡の写しを見る")), AVOID)
        self.assertEqual([(2, "この鏡の写しを見る")], [(f.page, f.what) for f in found])
        self.assertEqual({"avoided_words"}, {f.check for f in found})

    def test_the_word_to_write_instead_is_in_the_sentence(self) -> None:
        self.assertEqual([(1, "still says '鏡の写し' — write 'レプリカ'")],
                         self.said(lambda s: text_at(s, "鏡の写しを見る")))

    def test_a_word_with_no_word_instead_says_it_is_on_the_list(self) -> None:
        self.assertEqual([(1, "'鏡の写し' is on the project's list of words to avoid")],
                         self.said(lambda s: text_at(s, "鏡の写しを見る"), config={"avoid": {"鏡の写し": ""}}))

    def test_a_word_split_by_a_run_boundary_is_caught(self) -> None:
        def split(slide):
            paragraph = text_at(slide, "鏡の").text_frame.paragraphs[0]
            paragraph.add_run().text = "写し"
            paragraph.runs[1].font.bold = True

        deck = self.deck(split)
        self.assertEqual([], [t for t in read(deck)[0].texts() if "鏡の写し" in t],
                         "the fixture no longer splits the word across runs")
        self.assertEqual(1, len(avoided_words.run(deck, AVOID)))

    def test_a_line_break_between_two_halves_is_not_one_word(self) -> None:
        def broken(slide):
            paragraph = text_at(slide, "鏡の").text_frame.paragraphs[0]
            paragraph.add_line_break()
            paragraph.add_run().text = "写し"

        self.assertEqual([], self.said(broken))

    def test_the_words_in_a_table_cell_are_read(self) -> None:
        def table(slide):
            cells = slide.shapes.add_table(2, 2, Emu(500000), Emu(500000), Emu(4000000), Emu(1000000)).table
            for row, words in enumerate((["手法", "値"], ["A", "二重の殻の厚さ"])):
                for column, word in enumerate(words):
                    cells.cell(row, column).text = word

        self.assertEqual([(1, "still says '二重の殻' — write 'ラッパー'")], self.said(table))

    def test_the_letters_are_matched_as_written(self) -> None:
        config = {"avoid": {"Mirror": "Replica"}}
        self.assertEqual([], self.said(lambda s: text_at(s, "a mirror of the page"), config=config))
        self.assertEqual(1, len(self.said(lambda s: text_at(s, "a Mirror of the page"), config=config)))

    def test_a_short_word_is_caught_inside_a_longer_one(self) -> None:
        self.assertEqual(1, len(self.said(lambda s: text_at(s, "書き写した"), config={"avoid": {"写し": "複製"}})))
        self.assertEqual(1, len(self.said(lambda s: text_at(s, "書き写した"), config={"avoid": {"写": "複製"}})))

    def test_a_word_is_said_once_for_a_page_however_often_it_stands_there(self) -> None:
        def twice(slide):
            text_at(slide, "鏡の写しと鏡の写し")
            text_at(slide, "もう一度、鏡の写し")

        self.assertEqual(1, len(self.said(twice)))

    def test_two_words_on_one_page_are_said_one_each_in_the_order_of_the_table(self) -> None:
        self.assertEqual(["鏡の写し", "二重の殻"],
                         [why.split("'")[1] for _, why in self.said(lambda s: text_at(s, "二重の殻と鏡の写し"))])

    def test_the_head_of_the_words_that_hit_is_shown_to_sixty_letters(self) -> None:
        text = "鏡の写し" + "あ" * 80
        found = avoided_words.run(self.deck(lambda s: text_at(s, text)), AVOID)
        self.assertEqual(text[:60], found[0].what)

    def test_a_word_in_no_page_is_not_said(self) -> None:
        self.assertEqual([], self.said(lambda s: text_at(s, "ふつうの文")))


class TheWordsInAChart(Fixtures):
    def test_each_place_a_chart_keeps_words_is_read(self) -> None:
        places = {"title": {"title": "鏡の写しの数"}, "value axis": {"axis": "鏡の写し (件)"},
                  "category axis": {"category_axis": "鏡の写し"}, "item": {"categories": ("鏡の写し", "B")},
                  "series": {"series": "鏡の写し"}}
        for place, words in places.items():
            with self.subTest(place):
                self.assertEqual([(1, "still says '鏡の写し' — write 'レプリカ'")],
                                 self.said(lambda s, w=words: chart_at(s, **w)))

    def test_a_chart_with_other_words_is_not_said(self) -> None:
        self.assertEqual([], self.said(lambda s: chart_at(s, title="件数", axis="件数 (件)")))

    def test_the_page_of_the_chart_is_the_page_said(self) -> None:
        found = avoided_words.run(self.deck(lambda s: text_at(s, "ふつう"), lambda s: chart_at(s, axis="鏡の写し")),
                                  AVOID)
        self.assertEqual([2], [f.page for f in found])

    def test_every_chart_of_a_page_is_read(self) -> None:
        def two(slide):
            chart_at(slide, title="一つ目")
            chart_at(slide, title="二つ目の鏡の写し", top=Emu(3000000))

        self.assertEqual(1, len(self.said(two)))

    def test_a_chart_the_file_does_not_hold_is_passed_over(self) -> None:
        deck = self.deck(lambda s: chart_at(s, title="件数"))
        broken = self.dir / "broken.pptx"
        with zipfile.ZipFile(deck) as source, zipfile.ZipFile(broken, "w") as target:
            for item in source.infolist():
                if not item.filename.startswith("ppt/charts/"):
                    target.writestr(item, source.read(item))
        page = read(broken)[0]
        self.assertEqual(((), []), (page.charts, page.chart_lines()))


class WhatTheProjectWrites(Fixtures):
    def test_without_a_table_nothing_is_said(self) -> None:
        deck = self.deck(lambda s: text_at(s, "鏡の写し"))
        for config in (None, {}, {"avoid": {}}, {"avoid": False}, {"stale_words": ["x"]}):
            with self.subTest(config):
                self.assertEqual([], avoided_words.run(deck, config))

    def test_a_word_set_to_false_is_let_through(self) -> None:
        config = {"avoid": {"鏡の写し": False, "二重の殻": "ラッパー"}}
        self.assertEqual(["二重の殻"], [why.split("'")[1] for _, why in
                                      self.said(lambda s: text_at(s, "鏡の写しと二重の殻"), config=config)])

    def test_an_empty_word_does_not_hit_every_page(self) -> None:
        self.assertEqual([], self.said(lambda s: text_at(s, "ふつう"), config={"avoid": {"": "x"}}))

    def test_a_table_that_is_not_a_table_is_refused(self) -> None:
        for table in (["鏡の写し"], "鏡の写し", True):
            with self.subTest(table), self.assertRaises(ValueError):
                avoided_words.avoided({"avoid": table})

    def test_a_word_to_write_instead_that_is_not_a_string_is_refused(self) -> None:
        for instead in (True, 3, ["レプリカ"]):
            with self.subTest(instead), self.assertRaises(ValueError) as caught:
                avoided_words.avoided({"avoid": {"鏡の写し": instead}})
            self.assertIn("鏡の写し", str(caught.exception))

    def test_it_runs_with_the_other_checks_of_a_deck(self) -> None:
        self.assertIn(avoided_words, CHECKS)
        found = run_all(self.deck(lambda s: text_at(s, "鏡の写し")), AVOID)
        self.assertEqual(["avoided_words"], [f.check for f in found if f.check == "avoided_words"])


class TheTemplateAndTheProject(Fixtures):
    """Both write `[checks.avoid]`; the two tables lie over each other key by key."""

    def project(self, template: str, own: str) -> Workspace:
        (self.dir / "template").mkdir()
        (self.dir / "template" / "workspace.toml").write_text(template, encoding="utf-8")
        (self.dir / "project").mkdir()
        (self.dir / "project" / "workspace.toml").write_text('template = "../template"\n' + own, encoding="utf-8")
        return Workspace.load(self.dir / "project")

    def test_the_words_of_both_tables_are_asked(self) -> None:
        workspace = self.project('[checks.avoid]\n"鏡の写し" = "レプリカ"\n',
                                 '[checks.avoid]\n"二重の殻" = "ラッパー"\n')
        self.assertEqual(AVOID["avoid"], avoided_words.avoided(_checks(workspace)))

    def test_the_project_writes_a_word_instead_of_its_own_for_a_word_of_the_template(self) -> None:
        workspace = self.project('[checks.avoid]\n"鏡の写し" = "レプリカ"\n', '[checks.avoid]\n"鏡の写し" = "複製"\n')
        self.assertEqual({"鏡の写し": "複製"}, avoided_words.avoided(_checks(workspace)))

    def test_the_project_lets_a_word_of_the_template_through_with_false(self) -> None:
        workspace = self.project('[checks.avoid]\n"鏡の写し" = "レプリカ"\n"二重の殻" = "ラッパー"\n',
                                 '[checks.avoid]\n"鏡の写し" = false\n')
        self.assertEqual({"二重の殻": "ラッパー"}, avoided_words.avoided(_checks(workspace)))
        found = avoided_words.run(self.deck(lambda s: text_at(s, "鏡の写しと二重の殻")), _checks(workspace))
        self.assertEqual(["二重の殻"], [f.why.split("'")[1] for f in found])

    def test_a_project_with_no_table_of_its_own_keeps_the_templates(self) -> None:
        workspace = self.project('[checks.avoid]\n"鏡の写し" = "レプリカ"\n', "")
        self.assertEqual({"鏡の写し": "レプリカ"}, avoided_words.avoided(_checks(workspace)))


if __name__ == "__main__":
    unittest.main()
