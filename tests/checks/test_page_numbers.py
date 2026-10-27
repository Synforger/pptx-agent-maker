"""A report names a page the way the person reading the deck counts it.

⚠ **人は刷られた番号で頁を言う。**表紙に番号を付けない資料は数え始めを 0 にするので、表紙の次 (= ファイルの
2 番目) の頁に刷られるのは 1 で、プレビューも PowerPoint もそう見せる。報告が位置だけを言うと、操作者の「2 ページ」
を別の頁と取り違える。数え始めが 1 の資料は位置と刷られる番号が同じなので、報告は今のまま。
"""

from __future__ import annotations

import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

from checks.test_checks import text_at  # noqa: E402
from support.templates import a_template, with_a_cover  # noqa: E402

from pptx_agent_maker import DEFAULT  # noqa: E402
from pptx_agent_maker._common.slide import first_slide_number  # noqa: E402
from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.checks import Finding, Findings, page_label, report, run_all  # noqa: E402
from pptx_agent_maker.project import create  # noqa: E402

WHITE_ON_BRAND = {"words": "FFFFFF", "ground": "2E9E5B"}


def a_deck(destination: Path, pages: int, *, counts_from: int | None = None, off_page: int | None = None) -> Path:
    """A deck of `pages` blank pages; `off_page` (= a place in the file) carries a shape that leaves the page."""
    deck = Presentation()
    deck.slide_width = Emu(DEFAULT.slide.width)
    deck.slide_height = Emu(DEFAULT.slide.height)
    for place in range(1, pages + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if place == off_page:
            text_at(slide, "はみ出す", left=Emu(DEFAULT.slide.width + 100000))
    if counts_from is not None:
        deck.part._element.set("firstSlideNum", str(counts_from))
    deck.save(str(destination))
    return destination


class FirstNumber(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_a_deck_that_does_not_say_counts_from_one(self) -> None:
        self.assertEqual(1, first_slide_number(a_deck(self.dir / "plain.pptx", 2)))

    def test_where_the_deck_says_the_count_starts_is_read(self) -> None:
        for said in (0, 1, 5):
            with self.subTest(said=said):
                self.assertEqual(said, first_slide_number(a_deck(self.dir / f"{said}.pptx", 2, counts_from=said)))


class PageLabel(unittest.TestCase):
    def test_a_deck_counted_from_one_is_said_as_it_always_was(self) -> None:
        self.assertEqual("page 3", page_label(3))
        self.assertEqual("page 3: 箱  <- 理由", Finding("c", 3, "箱", "理由").render())
        self.assertEqual("page 3: 箱  <- 理由", Finding("c", 3, "箱", "理由").render(1))

    def test_a_deck_counted_from_zero_adds_the_number_printed(self) -> None:
        self.assertEqual("page 3 (printed 2)", page_label(3, 0))
        self.assertEqual("page 1 (printed 0)", page_label(1, 0))
        self.assertEqual("page 3 (printed 7)", page_label(3, 5))
        self.assertEqual("page 3 (printed 2): 箱  <- 理由", Finding("c", 3, "箱", "理由").render(0))

    def test_the_whole_deck_is_not_a_page(self) -> None:
        self.assertEqual("page 0", page_label(0, 0))


class Report(unittest.TestCase):
    def test_a_deck_counted_from_one_has_no_note(self) -> None:
        said = report([Finding("off_page", 2, "箱", "理由")])
        self.assertIn("page 2: 箱", said)
        self.assertNotIn("printed", said)

    def test_a_deck_counted_from_zero_says_both_and_what_they_are(self) -> None:
        said = report([Finding("off_page", 2, "箱", "理由")], first_number=0)
        self.assertIn("page 2 (printed 1): 箱", said)
        self.assertIn("first page prints 0", said)

    def test_a_clean_report_has_nothing_to_explain(self) -> None:
        self.assertNotIn("printed", report([], first_number=0))
        self.assertNotIn("printed", report(Findings([Finding("untyped_parts", 0, "x", "y")]), first_number=1))

    def test_a_finding_on_the_whole_deck_is_not_given_a_printed_number(self) -> None:
        said = report([Finding("untyped_parts", 0, "x", "y")], first_number=0)
        self.assertIn("page 0: x", said)
        self.assertNotIn("printed", said)


class WhatTheCommandsPrint(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.root = self.dir / "project"

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def a_project(self) -> None:
        create(self.root)

    def run_command(self, *argv: str) -> tuple[int, str]:
        said = io.StringIO()
        with contextlib.redirect_stdout(said), contextlib.redirect_stderr(said):
            code = main([argv[0], str(self.root), *argv[1:]])
        return code, said.getvalue()

    def test_check_names_the_page_the_reader_sees(self) -> None:
        self.a_project()
        a_deck(self.root / "from-zero.pptx", 3, counts_from=0, off_page=3)
        code, said = self.run_command("check", "from-zero")
        self.assertEqual(1, code, said)
        self.assertIn("page 3 (printed 2): ", said)

    def test_check_leaves_a_deck_counted_from_one_as_it_was(self) -> None:
        self.a_project()
        a_deck(self.root / "from-one.pptx", 3, off_page=3)
        code, said = self.run_command("check", "from-one")
        self.assertEqual(1, code, said)
        self.assertIn("page 3: ", said)
        self.assertNotIn("printed", said)

    def test_a_pair_let_through_names_the_page_the_reader_sees(self) -> None:
        deck = a_deck(self.dir / "pair.pptx", 2, counts_from=0)
        presentation = Presentation(str(deck))
        box = text_at(presentation.slides[1], "字", size=Pt(14))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("2E9E5B")
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        presentation.save(str(deck))
        said = report(run_all(deck, {"contrast_allowed": [WHITE_ON_BRAND]}))
        self.assertRegex(said, r"pass\s+contrast\s+#FFFFFF on #2E9E5B \(contrast 3\.4\): 1 place on page 2 \(printed 1\)")

    def test_build_says_where_the_numbers_start_and_names_the_page_the_reader_sees(self) -> None:
        create(self.root, specimen=with_a_cover(a_template(self.dir / "cover.pptx", numbered=True)))
        (self.root / "assets" / "deck").mkdir(parents=True)
        (self.root / "assets" / "deck" / "example.png").write_bytes(
            (self.root / "assets" / "example" / "example.png").read_bytes())
        # 出所の無い図の頁は、宣言を読む検査 (`unsourced`) が止める
        body = '[[pages]]\nkind = "declare"\ntype = "figure"\ntitle = "出所の無い頁"\nfigure = "example.png"\n\n'
        (self.root / "deck.toml").write_text(
            'specimen = "specimen.pptx"\nout = "deck.pptx"\n\n[[pages]]\nkind = "copy"\npage = 1\n\n' + body * 2,
            encoding="utf-8")
        code, said = self.run_command("build", "deck")
        self.assertEqual(1, code, said)
        self.assertIn("(3 pages; the page numbers printed start at 0)", said)
        self.assertIn("page 3 (printed 2): ", said)
        self.assertNotIn("page 1 (printed", said, "the cover has nothing wrong with it")


if __name__ == "__main__":
    unittest.main()
