"""Taking a PowerPoint edit back into the manifest (= `review --apply`).

人が触るのは PowerPoint で、差分を manifest に戻すのはエージェント。ここは直しの種類ごとに
写し方が正しいこと、取り込んだ後に焼き直すと直した版と同じに読めること、確かめで止まった時は
何も書かないことを見る。直しは python-pptx で開いて保存して作る (= 保存のたびに全部の部品が
書き直されるので、直していない頁が「変わっていない」と読めるかも同時に見ている)。
"""

from __future__ import annotations

import io
import shutil
import sys
import tempfile
import tomllib
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx_agent_maker.__main__ import main  # noqa: E402
from pptx_agent_maker.checks.slide import read  # noqa: E402
from pptx_agent_maker.project import create  # noqa: E402
from pptx_agent_maker.review import apply as applying  # noqa: E402
from test_swap import arranged_by_hand  # noqa: E402

DATA = Path(__file__).resolve().parent / "data"

W1 = '''specimen = "specimen.pptx"
out = "w1.pptx"

[[pages]]
kind = "copy"
page = 1
replace = [["案件名", "見本"], ["第 N 回 進捗報告", "第 1 回"]]

[[pages]]
kind = "declare"
type = "figure"
title = "手法 A の結果"
figure = "dot.png"
conclusion = "A はこう読める"

# 持ち込んだ頁 (= 手で並べた絵 3 枚と表)
[[pages]]
kind = "import"
deck = "w0.pptx"
page = 1

[[pages]]
kind = "declare"
type = "cards"
title = "まとめ"
cards = [["一", "いち"], ["二", "に"]]
'''


class TakingAnEditIn(unittest.TestCase):

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = create(Path(self.tmp.name) / "project")
        for name in ("dot.png", "wide.png"):
            shutil.copy(DATA / name, self.root / "assets" / name)
        arranged_by_hand(self.root / "w0.pptx")
        (self.root / "w1.toml").write_text(W1, encoding="utf-8")
        self.assertEqual(self._run("build", str(self.root), "w1", "--skip-checks")[0], 0)
        self.deck = self.root / "w1.pptx"

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _run(self, *argv: str) -> tuple[int, str]:
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main(list(argv))
        return code, out.getvalue() + err.getvalue()

    def _apply(self) -> tuple[int, str]:
        return self._run("review", str(self.root), "w1.pptx", "--apply")

    def _pages(self) -> list[dict]:
        return tomllib.loads((self.root / "w1.toml").read_text(encoding="utf-8"))["pages"]

    def _edit(self, change) -> None:
        deck = Presentation(str(self.deck))
        change(deck)
        deck.save(str(self.deck))

    @staticmethod
    def _retext(deck, old: str, new: str) -> None:
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text == old:
                                run.text = new
                                return
        raise AssertionError(f"{old!r} is not on the deck")

    def _rebuilds_as_edited(self) -> None:
        """取り込んだ後の build は止まらず、直した版と同じ文言で焼ける。"""
        edited = [p.texts() for p in read(self.deck)]
        code, said = self._run("build", str(self.root), "w1", "--skip-checks")
        self.assertEqual(code, 0, said)
        self.assertEqual([p.texts() for p in read(self.deck)], edited)

    def test_words_on_a_declared_page_go_into_the_declaration(self) -> None:
        self._edit(lambda d: self._retext(d, "手法 A の結果", "手法 A の結果 (改)"))
        code, said = self._apply()
        self.assertEqual(code, 0, said)
        pages = self._pages()
        self.assertEqual(pages[1]["title"], "手法 A の結果 (改)")
        self.assertNotIn("replace", pages[1], "the declaration should say it, not a replacement")
        self.assertEqual(pages[2]["kind"], "import", "an untouched page was rewritten")
        self._rebuilds_as_edited()

    def test_words_that_a_replacement_put_there_move_the_replacement(self) -> None:
        self._edit(lambda d: self._retext(d, "見本", "新しい表紙"))
        self.assertEqual(self._apply()[0], 0)
        self.assertIn(["案件名", "新しい表紙"], self._pages()[0]["replace"])
        self._rebuilds_as_edited()

    def test_words_in_a_new_place_keep_the_page_as_edited(self) -> None:
        """人が新しい場所に文章を足した頁は、manifest では書けない ― 直した版を正にする。"""
        def add_box(deck):
            box = deck.slides[3].shapes.add_textbox(Emu(900000), Emu(5000000),
                                                    Emu(4000000), Emu(400000))
            box.text_frame.paragraphs[0].add_run().text = "新しく足した一文"
        self._edit(add_box)
        code, said = self._apply()
        self.assertEqual(code, 0, said)
        page = self._pages()[3]
        self.assertEqual(page["kind"], "import")
        self.assertTrue(page["deck"].startswith("_edits/"))
        self._rebuilds_as_edited()
        self.assertIn("新しく足した一文", [t for p in read(self.deck) for t in p.texts()])

    def test_a_picture_swapped_on_an_imported_page_goes_into_assets(self) -> None:
        with zipfile.ZipFile(self.deck) as archive:
            items = {i.filename: archive.read(i) for i in archive.infolist()}
        order = [n for n in items if n.startswith("ppt/slides/_rels/")]
        third = sorted(order, key=lambda n: int(n.split("slide")[-1].split(".")[0]))
        rels = next(items[n].decode() for n in third
                    if "image" in items[n].decode() and b"a:tbl" in items[n.replace("_rels/", "").replace(".rels", "")])
        media = rels.split('media/')[1].split('"')[0]
        items[f"ppt/media/{media}"] = (DATA / "dot.png").read_bytes()
        with zipfile.ZipFile(self.deck, "w") as archive:
            for name, data in items.items():
                archive.writestr(name, data)
        code, said = self._apply()
        self.assertEqual(code, 0, said)
        page = self._pages()[2]
        self.assertEqual(len(page["pictures"]), 3)
        for name in page["pictures"]:
            self.assertTrue((self.root / "assets" / "w1" / name).is_file())
        self._rebuilds_as_edited()

    def test_a_page_added_in_powerpoint_is_taken_in_where_it_stands(self) -> None:
        def add_page(deck):
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            box = slide.shapes.add_textbox(Emu(900000), Emu(900000), Emu(4000000), Emu(400000))
            box.text_frame.paragraphs[0].add_run().text = "足した頁"
        self._edit(add_page)
        self.assertEqual(self._apply()[0], 0)
        pages = self._pages()
        self.assertEqual(len(pages), 5)
        self.assertEqual((pages[4]["kind"], pages[4]["page"]), ("import", 5))
        self._rebuilds_as_edited()

    def test_a_page_removed_in_powerpoint_leaves_the_manifest(self) -> None:
        def remove_last(deck):
            ids = deck.slides._sldIdLst
            ids.remove(ids[-1])
        self._edit(remove_last)
        self.assertEqual(self._apply()[0], 0)
        self.assertEqual([p["kind"] for p in self._pages()], ["copy", "declare", "import"])
        self._rebuilds_as_edited()

    def test_a_take_that_would_not_read_as_edited_writes_nothing(self) -> None:
        """⚠ 確かめが効いていることを、わざと頁を落とす書き換えで見る。"""
        self._edit(lambda d: self._retext(d, "手法 A の結果", "手法 A の結果 (改)"))
        before = (self.root / "w1.toml").read_bytes()
        rewrite = applying._rewrite
        with mock.patch.object(applying, "_rewrite",
                               side_effect=lambda text, pages: rewrite(text, pages[:-1])):
            code, said = self._apply()
        self.assertEqual(code, 1)
        self.assertIn("nothing was written", said)
        self.assertEqual((self.root / "w1.toml").read_bytes(), before)
        self.assertEqual(self._run("build", str(self.root), "w1")[0], 1,
                         "the edited deck should still be protected after a refused take")

    def test_nothing_changed_is_said(self) -> None:
        code, said = self._apply()
        self.assertEqual(code, 1)
        self.assertIn("nothing to take in", said)


if __name__ == "__main__":
    unittest.main()
