"""A chart drawn from its numbers, which a person can still edit.

⚠ **絵で貼ったグラフは、数字が 1 つ変わるたびに描き直しになり、渡した先では誰も直せない。**ここで書く
グラフは数字を deck の中に持つ ― 見張るのは 3 つ:

* 書いた数字が、書いたとおりの形と色で描かれること (= 目立たせるのは 1 つ、滝グラフの台は道具が計算する)
* 数字が deck の中に在ること (= PowerPoint の「データの編集」が開く表)
* グラフの部品が頁と一緒に運ばれること ― 頁はグラフを関係で指すだけなので、頁の XML を写しただけでは
  グラフは付いて来ず、指す先の無い deck ができる
"""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO / "tests"))

from pptx import Presentation  # noqa: E402

from pptx_agent_maker import DEFAULT, PageFullError, PageTypeError, checks  # noqa: E402
from pptx_agent_maker.checks.base.slide import read as read_pages  # noqa: E402
from pptx_agent_maker.deck import Deck  # noqa: E402
from pptx_agent_maker.layout import types  # noqa: E402
from pptx_agent_maker.layout.base.tokens import theme_from  # noqa: E402
from pptx_agent_maker.layout.parts.elements import Chart  # noqa: E402
from pptx_agent_maker.write import add_page, aspect, new_deck, save  # noqa: E402
from support.pages import build  # noqa: E402
from test_deck import a_specimen, make_dot  # noqa: E402

P = DEFAULT.palette
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
NORTH = {"name": "North", "values": [12, 15, 21, 18]}
SOUTH = {"name": "South", "values": [8, 9, 14, 16]}
WEST = {"name": "West", "values": [5, 7, 6, 9]}


def a_chart(kind: str = "column", series=(NORTH,), categories=QUARTERS, theme=DEFAULT, **keys) -> Chart:
    page = build({"type": "chart", "title": "A chart",
                  "chart": {"kind": kind, "categories": list(categories), "series": list(series), **keys}}, theme)
    return next(e for e in page.build() if isinstance(e, Chart))


def colours(chart: Chart) -> list[str]:
    return [series.colour for series in chart.series]


class WhatIsDrawn(unittest.TestCase):
    def test_one_series_is_drawn_in_the_colour_that_marks_a_reading(self) -> None:
        chart = a_chart()
        self.assertEqual([P.accent], colours(chart))
        self.assertEqual((), chart.series[0].points)
        self.assertEqual(("column", tuple(QUARTERS), (12, 15, 21, 18)),
                         (chart.plot, chart.categories, chart.series[0].values))

    def test_a_category_that_stands_out_is_the_one_coloured_bar_among_grey(self) -> None:
        chart = a_chart(highlight="Q3")
        self.assertEqual([P.muted], colours(chart))
        self.assertEqual(((2, P.accent),), chart.series[0].points)

    def test_a_series_that_stands_out_is_the_one_coloured_series_among_grey(self) -> None:
        chart = a_chart(series=(NORTH, SOUTH, WEST), highlight="South")
        self.assertEqual([P.muted, P.accent, P.muted], colours(chart))
        self.assertTrue(all(series.points == () for series in chart.series))

    def test_series_nobody_coloured_run_from_the_accent_into_greys(self) -> None:
        self.assertEqual([P.accent, P.muted], colours(a_chart(series=(NORTH, SOUTH))))
        self.assertEqual([P.accent, P.muted, P.rule], colours(a_chart(series=(NORTH, SOUTH, WEST))))

    def test_a_fourth_series_needs_a_colour_of_its_own(self) -> None:
        east = {"name": "East", "values": [1, 2, 3, 4]}
        with self.assertRaises(PageTypeError) as refused:
            a_chart(series=(NORTH, SOUTH, WEST, east))
        self.assertIn("series 4 (East)", str(refused.exception))
        self.assertEqual(P.bad, colours(a_chart(series=(NORTH, SOUTH, WEST, {**east, "tone": "bad"})))[3])
        self.assertEqual(P.accent, colours(a_chart(series=(NORTH, SOUTH, WEST, east), highlight="East"))[3])

    def test_a_series_takes_the_ground_its_tone_names(self) -> None:
        named = theme_from({"grounds": {"Inside": "EAF0F8"}})
        chart = a_chart(series=({**NORTH, "tone": "good"}, {**SOUTH, "tone": "Inside"}, {**WEST, "tone": "tint"}),
                        theme=named)
        self.assertEqual([P.good, "EAF0F8", P.tint], colours(chart))
        with self.assertRaises(PageTypeError) as refused:
            a_chart(series=({**NORTH, "tone": "loud"},))
        self.assertIn("tone", str(refused.exception))

    def test_the_legend_comes_with_a_second_series(self) -> None:
        self.assertFalse(a_chart().legend)
        self.assertTrue(a_chart(series=(NORTH, SOUTH)).legend)

    def test_the_numbers_are_shown_once_on_the_bars_or_on_the_axis(self) -> None:
        on_the_bars, on_the_axis = a_chart(), a_chart(labels=False)
        self.assertEqual(("#,##0", False), (on_the_bars.series[0].number_format, on_the_bars.axis))
        self.assertEqual(("", True), (on_the_axis.series[0].number_format, on_the_axis.axis))

    def test_a_number_is_printed_to_the_places_written_with_its_unit_after_it(self) -> None:
        share = {"name": "Share", "values": [12.5, 15, 21.25, 1800]}
        self.assertEqual('#,##0.00"%"', a_chart(series=(share,), unit="%").series[0].number_format)
        self.assertEqual('#,##0" 件"', a_chart(unit=" 件").series[0].number_format)
        tenth = {"name": "Rate", "values": [0.1, 0.2, 0.3, 1]}
        self.assertEqual("#,##0.0", a_chart(series=(tenth,)).series[0].number_format)

    def test_the_words_of_a_chart_are_the_size_and_colour_of_the_body(self) -> None:
        self.assertEqual((DEFAULT.type.body, P.ink, P.rule), (a_chart().size, a_chart().colour, a_chart().line))
        self.assertEqual(18, a_chart(theme=theme_from({"use": "present"})).size)

    def test_the_numbers_inside_stacked_columns_read_on_the_ground_they_sit_on(self) -> None:
        chart = a_chart("stacked", series=({**NORTH, "tone": "accent"}, {**SOUTH, "tone": "box"}))
        self.assertEqual("stacked", chart.plot)
        self.assertEqual([P.paper, P.ink], [series.label_colour for series in chart.series])
        # 段の境は紙の色の細い線。薄い地の段は、自分の枠がその役をする
        self.assertEqual([P.paper, P.edge(P.box)], [series.outline for series in chart.series])
        self.assertEqual({""}, {series.label_colour + series.outline for series in a_chart().series})

    def test_a_series_on_a_light_ground_has_the_edge_a_box_on_that_ground_has(self) -> None:
        """⚠ 薄い地は紙と近い。枠が無いと、棒の端が背景に溶ける (= 箱と線表の棒で差し戻された事と同じ)。"""
        named = theme_from({"grounds": {"Inside": "EAF0F8", "Late": "8B1E3F"}})
        tones = ("box", "band", "tint", "Inside", "accent", "good", "bad", "Late")
        series = tuple({"name": tone, "values": [1, 2, 3, 4], "tone": tone} for tone in tones)
        chart = a_chart(series=series, theme=named)
        light = {"box": P.box, "band": P.band, "tint": P.tint, "Inside": "EAF0F8"}
        self.assertEqual([P.edge(light[tone]) if tone in light else "" for tone in tones],
                         [drawn.outline for drawn in chart.series])
        xml, _parts = written(a_chart(series=({**NORTH, "tone": "box"},)))
        self.assertIn(f'<a:ln w="{DEFAULT.spacing.hairline}"><a:solidFill><a:srgbClr val="{P.edge(P.box)}"/>', xml)

    def test_every_kind_is_drawn_as_itself(self) -> None:
        self.assertEqual({"bar": "bar", "column": "column", "line": "line", "stacked": "stacked"},
                         {kind: a_chart(kind).plot for kind in ("bar", "column", "line", "stacked")})

    def test_a_chart_is_what_the_page_shows(self) -> None:
        """グラフの頁は絵を持たなくても組める (= 数字を形で見せている)。部品としてもマスに入る。"""
        self.assertIsNotNone(a_chart())
        cell = {"chart": {"kind": "bar", "categories": ["a", "b"], "series": [{"name": "n", "values": [1, 2]}]}}
        page = build({"type": "compose", "title": "In a cell",
                      "rows": [{"cells": [cell, {"card": ["Beside", "its reading"]}]}]})
        chart = next(e for e in page.build() if isinstance(e, Chart))
        self.assertEqual("bar", chart.plot)
        self.assertTrue(page.theme.frame().contains(chart.rect))


class TheListingOfTypes(unittest.TestCase):
    def test_it_says_how_a_chart_is_written(self) -> None:
        said = types.describe()
        self.assertRegex(said, r"chart\s+needs: chart")
        self.assertIn("kind (bar | column | line | stacked | waterfall | scatter)", said)
        self.assertIn('a scatter:   kind = "scatter", x, y', said)
        self.assertIn("points = [[x, y], …]", said)
        self.assertIn('data = "x.csv"', said)
        self.assertIn("callouts = [{ at, series, text }]", said)


class AWaterfall(unittest.TestCase):
    STEPS = ["Start", "Up", "Down", "More", "End"]

    def drawn(self, values, totals=("Start", "End"), categories=None, **keys) -> dict:
        chart = a_chart("waterfall", series=({"name": "Profit", "values": list(values)},),
                        categories=categories or self.STEPS, totals=list(totals), **keys)
        self.assertEqual(("stacked", False), (chart.plot, chart.legend))
        return {series.name: series for series in chart.series}

    def test_each_step_floats_where_the_one_before_it_ended(self) -> None:
        series = self.drawn([100, 20, -15, 30, 135])
        self.assertEqual(["(base)", "Profit +", "Profit -", "Profit ="], list(series))
        self.assertEqual((None, 100, 105, 105, None), series["(base)"].values)
        self.assertEqual((None, 20, None, 30, None), series["Profit +"].values)
        self.assertEqual((None, None, 15, None, None), series["Profit -"].values)
        self.assertEqual((100, None, None, None, 135), series["Profit ="].values)

    def test_the_base_is_not_drawn_and_the_rest_say_what_they_are_by_colour(self) -> None:
        series = self.drawn([100, 20, -15, 30, 135])
        self.assertEqual(["", P.good, P.bad, P.accent], [s.colour for s in series.values()])
        self.assertEqual("", series["(base)"].number_format)
        self.assertEqual(["+#,##0;+#,##0", "-#,##0;-#,##0", "#,##0"],
                         [s.number_format for s in list(series.values())[1:]])
        self.assertEqual({P.words_on(P.good), P.words_on(P.bad), P.words_on(P.accent)},
                         {s.label_colour for s in list(series.values())[1:]})

    def test_a_step_is_printed_with_the_places_and_the_unit_of_the_rest(self) -> None:
        series = self.drawn([100, 20.5, -15, 30, 135.5], unit="M")
        self.assertEqual('+#,##0.0"M";+#,##0.0"M"', series["Profit +"].number_format)
        self.assertEqual("", self.drawn([100, 20, -15, 30, 135], labels=False)["Profit +"].number_format)

    def test_a_total_must_say_what_the_steps_come_to(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            self.drawn([100, 20, -15, 30, 140])
        self.assertIn("'End' is written as 140, and the steps before it add up to 135", str(refused.exception))

    def test_the_first_total_is_where_the_steps_start_from(self) -> None:
        self.assertEqual((250, None, None, None, 285), self.drawn([250, 20, -15, 30, 285])["Profit ="].values)

    def test_a_total_in_the_middle_stands_on_the_ground_too(self) -> None:
        series = self.drawn([100, 20, 120, -30, 90], totals=("Start", "Down", "End"))
        self.assertEqual((100, None, 120, None, 90), series["Profit ="].values)
        self.assertEqual((None, 100, None, 90, None), series["(base)"].values)

    def test_steps_with_no_total_start_from_nothing(self) -> None:
        series = self.drawn([10, 5, -3], totals=(), categories=["a", "b", "c"])
        self.assertEqual((None, 10, 12), series["(base)"].values)
        self.assertEqual((10, 5, None), series["Profit +"].values)
        self.assertEqual((None, None, 3), series["Profit -"].values)

    def test_fractions_add_up_as_they_are_written(self) -> None:
        """⚠ 0.1 を 3 回足して 0.3 にならない足し算では、合計の照合が狂う。"""
        series = self.drawn([0, 0.1, 0.2, -0.3, 0])
        self.assertEqual((None, None, 0.1, None, None), series["(base)"].values)
        self.assertEqual((None, None, None, 0.3, None), series["Profit -"].values)

    def test_steps_below_zero_hang_down_from_the_ground(self) -> None:
        series = self.drawn([-10, -20, 5], totals=(), categories=["a", "b", "c"])
        self.assertEqual((None, -10, -25), series["(base)"].values)
        self.assertEqual((-10, -20, None), series["Profit -"].values)
        self.assertEqual((None, None, -5), series["Profit +"].values)

    def test_a_step_across_zero_is_refused(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            self.drawn([10, -25], totals=(), categories=["a", "b"])
        self.assertIn("across zero", str(refused.exception))
        self.assertTrue(self.drawn([10, -10, -15], totals=(), categories=["a", "b", "c"]))

    def test_it_is_drawn_from_one_series_and_takes_no_colour_of_its_own(self) -> None:
        for keys, word in (({"series": [NORTH, SOUTH]}, "one series"), ({"highlight": "Q1"}, "highlight"),
                           ({"series": [{**NORTH, "tone": "bad"}]}, "tone")):
            with self.subTest(word), self.assertRaises(PageTypeError) as refused:
                build({"type": "chart", "title": "t", "chart": {"kind": "waterfall", "categories": QUARTERS,
                                                               "series": [NORTH], **keys}})
            self.assertIn(word, str(refused.exception))

    def test_totals_belong_to_a_waterfall_and_name_its_categories(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            a_chart(totals=["Q1"])
        self.assertIn("`totals` belongs to a waterfall", str(refused.exception))
        with self.assertRaises(PageTypeError) as refused:
            self.drawn([100, 20, -15, 30, 135], totals=("Start", "Finish"))
        self.assertIn("`totals` lists the categories", str(refused.exception))


class WhatIsRefused(unittest.TestCase):
    def refused(self, chart: dict, *words: str) -> None:
        with self.assertRaises(PageTypeError) as raised:
            build({"type": "chart", "title": "t", "chart": {"kind": "column", "categories": QUARTERS,
                                                           "series": [NORTH], **chart}})
        for word in words:
            self.assertIn(word, str(raised.exception))

    def test_a_kind_nobody_knows(self) -> None:
        self.refused({"kind": "pie"}, "`kind` is one of bar, column, line, stacked, waterfall")

    def test_a_key_nobody_reads(self) -> None:
        self.refused({"colour": "red"}, "chart does not take colour")
        self.refused({"series": [{**NORTH, "color": "red"}]}, "chart: series 1 does not take color")

    def test_no_categories_or_no_series(self) -> None:
        self.refused({"categories": []}, "`categories`")
        self.refused({"categories": ["a", " "]}, "`categories`")
        self.refused({"series": []}, "`series`")

    def test_a_series_without_a_name_or_with_too_few_numbers(self) -> None:
        self.refused({"series": [{"values": [1, 2, 3, 4]}]}, "series 1 has no `name`")
        self.refused({"series": [{"name": "n", "values": [1, 2, 3]}]}, "one number for each of the 4 categories")

    def test_something_that_is_not_a_number(self) -> None:
        for value in ("12", True, None, float("nan"), float("inf")):
            with self.subTest(value=value):
                self.refused({"series": [{"name": "n", "values": [1, 2, value, 4]}]},
                             "series 1 (n), value 3", "drawn from numbers")

    def test_a_name_written_twice(self) -> None:
        self.refused({"categories": ["a", "b", "a", "c"]}, "a category is written twice")
        self.refused({"series": [NORTH, NORTH]}, "a series is written twice")

    def test_a_highlight_that_names_nothing_or_two_things(self) -> None:
        self.refused({"highlight": "Q9"}, "neither a series (North) nor a category (Q1, Q2, Q3, Q4)")
        self.refused({"highlight": 3}, "`highlight` is the name of")
        self.refused({"series": [{"name": "Q1", "values": [1, 2, 3, 4]}], "highlight": "Q1"},
                     "names both a series and a category")

    def test_a_category_cannot_stand_out_among_several_series(self) -> None:
        self.refused({"series": [NORTH, SOUTH], "highlight": "Q3"}, "with 2 series, `highlight` names one of them")

    def test_a_highlight_and_a_tone_at_once(self) -> None:
        self.refused({"series": [{**NORTH, "tone": "good"}, SOUTH], "highlight": "South"},
                     "take the `tone` off North")

    def test_a_unit_or_a_switch_written_wrong(self) -> None:
        self.refused({"unit": 'in "M"'}, "`unit`")
        self.refused({"unit": 3}, "`unit`")
        self.refused({"labels": "no"}, "`labels` is true or false")

    def test_a_chart_laid_into_a_cell_names_its_cell(self) -> None:
        with self.assertRaises(PageTypeError) as raised:
            build({"type": "compose", "title": "t",
                   "rows": [{"cells": [{"card": ["a", "b"]}, {"chart": {"kind": "pie"}}]}]})
        self.assertIn("compose: row 1, cell 2, chart: `kind`", str(raised.exception))


class NumbersFromAFile(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def read(self, text: str, name: str = "numbers.csv", **keys) -> Chart:
        (self.dir / name).write_text(text, encoding="utf-8")
        page = types.build({"type": "chart", "title": "t", "chart": {"kind": "column", "data": name, **keys}},
                           lambda asset: self.dir / asset, aspect)
        return next(e for e in page.build() if isinstance(e, Chart))

    def test_the_first_column_names_the_categories_and_the_first_row_the_series(self) -> None:
        chart = self.read("﻿Quarter,North,South\nQ1,12,8\nQ2,15.5,-9\n\n")
        self.assertEqual(("Q1", "Q2"), chart.categories)
        self.assertEqual([("North", (12, 15.5)), ("South", (8, -9))], [(s.name, s.values) for s in chart.series])
        self.assertIsInstance(chart.series[0].values[0], int)
        self.assertEqual("#,##0.0", chart.series[0].number_format)

    def test_a_series_from_a_file_can_stand_out(self) -> None:
        chart = self.read("Quarter,North,South\nQ1,12,8\n", highlight="South")
        self.assertEqual([P.muted, P.accent], colours(chart))

    def test_a_cell_that_is_not_a_number_is_named_by_its_row_and_column(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            self.read("Quarter,North,South\nQ1,12,8\nQ2,n/a,9\n")
        self.assertIn("numbers.csv, row 3, column 2 is 'n/a'", str(refused.exception))

    def test_a_row_of_another_length_is_named(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            self.read("Quarter,North,South\nQ1,12\n")
        self.assertIn("numbers.csv, row 2 has 2 cells and the first row has 3", str(refused.exception))

    def test_a_file_with_no_numbers_or_no_names(self) -> None:
        for text in ("Quarter,North\n", "Quarter\nQ1\n", "Quarter,,South\nQ1,1,2\n"):
            with self.subTest(text=text), self.assertRaises(PageTypeError):
                self.read(text)

    def test_the_numbers_come_from_the_file_or_from_the_page_not_both(self) -> None:
        with self.assertRaises(PageTypeError) as refused:
            self.read("Quarter,North\nQ1,1\n", categories=["Q1"])
        self.assertIn("not from both", str(refused.exception))
        with self.assertRaises(PageTypeError) as refused:
            self.read("Quarter,North\nQ1,1\n", name="numbers.txt")
        self.assertIn("`data` is the name of a CSV file", str(refused.exception))


def written(chart: Chart, theme=DEFAULT) -> tuple[str, dict[str, bytes]]:
    """The chart's own XML once written out, and every part of the deck it was written into."""
    with tempfile.TemporaryDirectory() as tmp:
        deck = new_deck(theme)
        add_page(deck, [chart], theme)
        saved = save(deck, Path(tmp) / "chart.pptx")
        Presentation(str(saved))
        with zipfile.ZipFile(saved) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
    return parts["ppt/charts/chart1.xml"].decode("utf-8"), parts


class AsWrittenOut(unittest.TestCase):
    def test_each_kind_is_the_chart_it_says(self) -> None:
        marks = {"bar": ('<c:barDir val="bar"/>', '<c:grouping val="clustered"/>'),
                 "column": ('<c:barDir val="col"/>', '<c:grouping val="clustered"/>'),
                 "stacked": ('<c:barDir val="col"/>', '<c:grouping val="stacked"/>'),
                 "line": ("<c:lineChart>", '<c:grouping val="standard"/>')}
        for kind, wanted in marks.items():
            with self.subTest(kind):
                xml, _parts = written(a_chart(kind, series=(NORTH, SOUTH)))
                for mark in wanted:
                    self.assertIn(mark, xml)

    def test_it_has_no_title_of_its_own_and_no_grid(self) -> None:
        xml, _parts = written(a_chart())
        self.assertIn('<c:autoTitleDeleted val="1"/>', xml)
        self.assertNotIn("<c:title>", xml)
        self.assertNotIn("<c:majorGridlines", xml)
        self.assertEqual(['<c:majorTickMark val="none"/>'] * 2, re.findall(r"<c:majorTickMark[^>]*/>", xml))

    def test_the_value_axis_is_there_only_when_the_bars_carry_no_numbers(self) -> None:
        def value_axis(chart: Chart) -> str:
            return re.search(r"<c:valAx>.*?</c:valAx>", written(chart)[0], re.S).group(0)

        self.assertRegex(value_axis(a_chart()), r'<c:delete(/>| val="1"/>)')
        self.assertIn('<c:delete val="0"/>', value_axis(a_chart(labels=False)))
        self.assertIn("<c:showVal val=\"1\"/>", written(a_chart())[0])
        self.assertNotIn("<c:dLbls>", written(a_chart(labels=False))[0])

    def test_the_colours_and_the_one_that_stands_out_are_in_the_file(self) -> None:
        xml, _parts = written(a_chart(highlight="Q3"))
        series = re.search(r"<c:ser>.*?</c:ser>", xml, re.S).group(0)
        self.assertIn(f'<c:spPr><a:solidFill><a:srgbClr val="{P.muted}"/></a:solidFill>', series)
        point = re.search(r"<c:dPt>.*?</c:dPt>", series, re.S).group(0)
        self.assertIn('<c:idx val="2"/>', point)
        self.assertIn(f'<a:srgbClr val="{P.accent}"/>', point)
        self.assertIn('<c:invertIfNegative val="0"/>', series)

    def test_the_numbers_are_printed_the_way_the_page_asked(self) -> None:
        xml, _parts = written(a_chart(unit="%"))
        self.assertIn('<c:numFmt formatCode="#,##0&quot;%&quot;" sourceLinked="0"/>', xml)
        self.assertIn('<c:dLblPos val="outEnd"/>', xml)
        self.assertIn('<c:dLblPos val="t"/>', written(a_chart("line"))[0])
        self.assertIn('<c:dLblPos val="ctr"/>', written(a_chart("stacked", series=(NORTH, SOUTH)))[0])

    def test_the_legend_sits_under_the_plot_when_there_is_one(self) -> None:
        self.assertNotIn("<c:legend>", written(a_chart())[0])
        legend = re.search(r"<c:legend>.*?</c:legend>", written(a_chart(series=(NORTH, SOUTH)))[0], re.S).group(0)
        self.assertIn('<c:legendPos val="b"/>', legend)
        self.assertIn('<c:overlay val="0"/>', legend)

    def test_a_bar_chart_reads_from_the_top_like_a_table(self) -> None:
        self.assertIn('<c:orientation val="maxMin"/>', written(a_chart("bar"))[0])
        self.assertNotIn('<c:orientation val="maxMin"/>', written(a_chart("column"))[0])

    def test_the_value_axis_of_a_bar_chart_stays_under_the_plot(self) -> None:
        """⚠ 項目の並びを逆にすると、値の軸は上へ回る。いちばん下の項目の側で交わらせて下に戻す。"""
        def crossing(chart: Chart) -> str:
            axis = re.search(r"<c:valAx>.*?</c:valAx>", written(chart)[0], re.S).group(0)
            return re.search(r'<c:crosses val="(\w+)"/>', axis).group(1)

        self.assertEqual("max", crossing(a_chart("bar", labels=False)))
        self.assertEqual("autoZero", crossing(a_chart("column", labels=False)))

    def test_the_words_are_set_in_the_decks_face_at_the_size_of_the_body(self) -> None:
        theme = theme_from({"font": "Arial", "use": "present"})
        xml, _parts = written(a_chart(theme=theme), theme)
        words = re.search(r"<c:txPr>.*?</c:txPr>", xml[xml.rindex("</c:chart>"):], re.S).group(0)
        self.assertIn('sz="1800"', words)
        self.assertIn('<a:latin typeface="Arial"/>', words)
        self.assertIn(f'<a:srgbClr val="{P.ink}"/>', words)

    def test_a_line_is_drawn_in_its_colour_with_round_marks(self) -> None:
        xml, _parts = written(a_chart("line", highlight="Q2"))
        series = re.search(r"<c:ser>.*?</c:ser>", xml, re.S).group(0)
        self.assertIn(f'<a:ln w="{DEFAULT.spacing.strong_line}"><a:solidFill><a:srgbClr val="{P.muted}"/>', series)
        self.assertIn('<c:symbol val="circle"/>', series)
        self.assertIn('<c:smooth val="0"/>', series)
        point = re.search(r"<c:dPt>.*?</c:dPt>", series, re.S).group(0)
        self.assertIn(f'<a:srgbClr val="{P.accent}"/>', point)

    def test_the_base_of_a_waterfall_is_there_and_cannot_be_seen(self) -> None:
        chart = a_chart("waterfall", series=({"name": "Profit", "values": [100, 20, -15, 105]},),
                        categories=["Start", "Up", "Down", "End"], totals=["Start", "End"])
        xml, _parts = written(chart)
        base, up, down, total = re.findall(r"<c:ser>.*?</c:ser>", xml, re.S)
        self.assertIn("<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>", base)
        self.assertNotIn("<c:dLbls>", base)
        # 「この項目には無い」値は書かない (= 0 と書くと、0 の棒に「+0」の数字が付く)
        self.assertEqual(['idx="1"', 'idx="2"'], re.findall(r'<c:pt (idx="\d")><c:v>\d+</c:v>',
                                                             base[base.index("<c:val>"):]))
        self.assertIn('formatCode="+#,##0;+#,##0"', up)
        self.assertIn('formatCode="-#,##0;-#,##0"', down)
        self.assertIn(f'<a:srgbClr val="{P.accent}"/>', total)
        self.assertIn('<c:overlap val="100"/>', xml)

    def test_the_numbers_travel_inside_the_deck_for_a_person_to_edit(self) -> None:
        """PowerPoint の「データの編集」が開くのは、グラフが中に持つ表。数字がそこに無ければ、形は
        描けても人は直せない。"""
        xml, parts = written(a_chart(series=(NORTH, SOUTH)))
        self.assertIn("<c:externalData r:id=", xml)
        rels = parts["ppt/charts/_rels/chart1.xml.rels"].decode("utf-8")
        book = re.search(r'Target="\.\./embeddings/([^"]+\.xlsx)"', rels).group(1)
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / book).write_bytes(parts[f"ppt/embeddings/{book}"])
            with zipfile.ZipFile(Path(tmp) / book) as workbook:
                sheet = workbook.read("xl/worksheets/sheet1.xml").decode("utf-8")
                words = workbook.read("xl/sharedStrings.xml").decode("utf-8")
        for name in ("North", "South", *QUARTERS):
            self.assertIn(f">{name}<", words)
        numbers = [float(value) for value in re.findall(r"<v>([\d.]+)</v>", sheet)]
        for value in (*NORTH["values"], *SOUTH["values"]):
            self.assertIn(value, numbers)
        self.assertIn("Sheet1!$B$2:$B$5", xml)


class WordsPointingAtAPoint(unittest.TestCase):
    """⚠ **注記の位置を座標で書かせない。**書くのは「どの項目の、どの系列」で、点の位置は道具が計算する。
    計算できるのは、描く範囲と値の軸を道具が決めて書き出すから ― 言葉と線が指す先と、PowerPoint が
    棒を描く先が、同じ数から出る。"""

    S, LINE = DEFAULT.spacing, DEFAULT.line_height()

    def pointed(self, kind="column", series=(NORTH,), callouts=({"at": "Q3", "text": "the new product"},),
                **keys) -> Chart:
        return a_chart(kind, series=series, callouts=list(callouts), **keys)

    @staticmethod
    def inner(chart: Chart):
        """The plot's area in EMU, from the fractions the chart is pinned to."""
        x, y, w, h = chart.pinned.inner
        frame = chart.frame
        return (frame.left + x * frame.width, frame.top + y * frame.height, w * frame.width, h * frame.height)

    def height_of(self, chart: Chart, value: float) -> float:
        _left, top, _wide, tall = self.inner(chart)
        return top + tall - tall * (value - chart.pinned.low) / (chart.pinned.high - chart.pinned.low)

    def test_a_chart_with_no_callout_is_left_to_powerpoint(self) -> None:
        chart = a_chart()
        self.assertEqual((None, None, ()), (chart.frame, chart.pinned, chart.callouts))
        xml, parts = written(chart)
        for mark in ("manualLayout", "<c:max ", "<c:min ", "<c:majorUnit", "<c:gapWidth"):
            self.assertNotIn(mark, xml)
        self.assertNotIn("<p:grpSp>", parts["ppt/slides/slide1.xml"].decode("utf-8"))

    def test_the_words_take_a_band_above_the_chart_and_the_chart_takes_the_rest(self) -> None:
        chart = self.pointed()
        self.assertEqual((chart.rect.left, chart.rect.top + self.LINE + self.S.gap_s, chart.rect.width,
                          chart.rect.height - self.LINE - self.S.gap_s),
                         (chart.frame.left, chart.frame.top, chart.frame.width, chart.frame.height))
        words = chart.callouts[0].words
        self.assertEqual((chart.rect.top, self.LINE, DEFAULT.width("the new product")),
                         (words.top, words.height, words.width))
        self.assertLessEqual(words.bottom, chart.frame.top)

    def test_the_words_stand_over_their_bar_and_the_line_runs_down_to_its_number(self) -> None:
        chart = self.pointed()
        left, _top, wide, _tall = self.inner(chart)
        middle = left + wide * (2 + 0.5) / 4                       # Q3 は 4 つのうちの 3 つめ
        callout = chart.callouts[0]
        self.assertAlmostEqual(middle, callout.words.left + callout.words.width / 2, delta=1)
        self.assertAlmostEqual(middle, callout.line.left + callout.line.width / 2, delta=1)
        self.assertEqual(self.S.hairline, callout.line.width)
        self.assertEqual(callout.words.bottom, callout.line.top)
        # 棒の頭の上には数字が 1 行ぶん在る。線はその上で止まる
        self.assertAlmostEqual(self.height_of(chart, 21) - self.LINE - self.S.bar_pad_y, callout.line.bottom, delta=1)

    def test_among_several_series_the_line_lands_on_the_bar_of_the_series_named(self) -> None:
        chart = self.pointed(series=(NORTH, SOUTH), callouts=(
            {"at": "Q1", "series": "North", "text": "a"}, {"at": "Q1", "series": "South", "text": "b"},
            {"at": "Q4", "series": "South", "text": "c"}))
        left, _top, wide, _tall = self.inner(chart)
        slot = wide / 4
        bar = slot / (2 + 1.5)                                     # 棒 2 本と、項目どうしの間 (= 棒 1.5 本ぶん)
        wanted = [left + bar * (0.75 + 0.5), left + bar * (0.75 + 1.5), left + 3 * slot + bar * (0.75 + 1.5)]
        found = [callout.line.left + callout.line.width / 2 for callout in chart.callouts]
        for want, got in zip(wanted, found):
            self.assertAlmostEqual(want, got, delta=1)
        heights = [callout.line.bottom for callout in chart.callouts]
        for value, got in zip((12, 8, 16), heights):
            self.assertAlmostEqual(self.height_of(chart, value) - self.LINE - self.S.bar_pad_y, got, delta=1)

    def test_on_a_line_the_point_is_the_value_itself_at_the_middle_of_its_category(self) -> None:
        chart = self.pointed("line", series=(NORTH, SOUTH), labels=False,
                             callouts=({"at": "Q2", "series": "South", "text": "flat"},))
        left, _top, wide, _tall = self.inner(chart)
        callout = chart.callouts[0]
        self.assertAlmostEqual(left + wide * 1.5 / 4, callout.line.left + callout.line.width / 2, delta=1)
        self.assertAlmostEqual(self.height_of(chart, 9) - self.S.bar_pad_y, callout.line.bottom, delta=1)

    def test_a_stacked_column_and_a_waterfall_are_pointed_at_the_top_of_one_column(self) -> None:
        stacked = self.pointed("stacked", series=(NORTH, SOUTH), callouts=({"at": "Q3", "text": "the peak"},))
        self.assertAlmostEqual(self.height_of(stacked, 21 + 14) - self.S.bar_pad_y, stacked.callouts[0].line.bottom,
                               delta=1)
        fall = a_chart("waterfall", series=({"name": "Profit", "values": [100, 20, -15, 105]},),
                       categories=["Start", "Up", "Down", "End"], totals=["Start", "End"],
                       callouts=[{"at": "Down", "text": "a"}, {"at": "End", "text": "b"}, {"at": "Up", "text": "c"}])
        tops = [callout.line.bottom + self.S.bar_pad_y for callout in fall.callouts]
        for value, got in zip((120, 105, 120), tops):                # 減った段は、減る前の高さが頭
            self.assertAlmostEqual(self.height_of(fall, value), got, delta=1)
        for kind, chart in (("stacked", {"series": [NORTH, SOUTH]}), ("waterfall", {"series": [NORTH]})):
            with self.subTest(kind), self.assertRaises(PageTypeError) as refused:
                build({"type": "chart", "title": "t", "chart": {
                    "kind": kind, "categories": QUARTERS, **chart,
                    "callouts": [{"at": "Q1", "series": "North", "text": "a"}]}})
            self.assertIn("one column at a time", str(refused.exception))

    def test_on_a_bar_chart_the_words_stand_to_the_right_level_with_their_bar(self) -> None:
        chart = self.pointed("bar", series=(NORTH, SOUTH), unit=" pt", callouts=(
            {"at": "Q2", "series": "South", "text": "behind"}, {"at": "Q4", "series": "North", "text": "ahead of it"}))
        band = DEFAULT.width("ahead of it") + self.S.gap_s
        self.assertEqual((chart.rect.width - band, chart.rect.height, chart.rect.top),
                         (chart.frame.width, chart.frame.height, chart.frame.top))
        left, top, wide, tall = self.inner(chart)
        slot = tall / 4
        bar = slot / (2 + 1.5)
        for callout, (category, member, value) in zip(chart.callouts, ((1, 1, 9), (3, 0, 18))):
            level = top + slot * category + bar * (0.75 + member + 0.5)       # 最初の項目が上
            self.assertEqual(chart.frame.right + self.S.gap_s, callout.words.left)
            self.assertAlmostEqual(level, callout.words.top + callout.words.height / 2, delta=1)
            self.assertAlmostEqual(level, callout.line.top + callout.line.height / 2, delta=1)
            self.assertEqual(self.S.hairline, callout.line.height)
            tip = left + wide * (value - chart.pinned.low) / (chart.pinned.high - chart.pinned.low)
            number = DEFAULT.width(f"{value} pt") + self.S.gap_s            # 棒の先に付く数字のぶんを空ける
            self.assertAlmostEqual(tip + number + self.S.bar_pad_y, callout.line.left, delta=1)
            self.assertEqual(callout.words.left - self.S.bar_pad_y, callout.line.right)

    def test_the_value_axis_starts_at_nothing_and_leaves_room_for_the_number_on_the_tallest_bar(self) -> None:
        chart = self.pointed()
        _left, top, _wide, _tall = self.inner(chart)
        self.assertEqual(0, chart.pinned.low)
        self.assertGreaterEqual(self.height_of(chart, 21) - top, self.LINE + self.S.bar_pad_y - 1)
        self.assertEqual(0, chart.pinned.high % chart.pinned.step)
        self.assertEqual((0.0, 25.0, 5.0, 150), (chart.pinned.low, chart.pinned.high, chart.pinned.step,
                                                 chart.pinned.gap))

    def test_the_steps_of_the_axis_are_ones_twos_and_fives(self) -> None:
        for values, step in (([1, 2, 3, 4.2], 1), ([120, 135, 150, 190], 50), ([0.2, 0.5, 0.9, 0.7], 0.2),
                             ([1200, 3400, 8100, 7000], 2000)):
            with self.subTest(values=values):
                chart = self.pointed(series=({"name": "n", "values": values},), labels=False)
                self.assertAlmostEqual(step, chart.pinned.step)
                self.assertGreaterEqual(chart.pinned.high, max(values))

    def test_values_below_zero_get_their_room_under_the_ground(self) -> None:
        chart = self.pointed(series=({"name": "n", "values": [10, -20, 5, 8]},),
                             callouts=({"at": "Q2", "text": "a loss"}, {"at": "Q1", "text": "a gain"}))
        self.assertLess(chart.pinned.low, -20)
        self.assertGreater(chart.pinned.high, 10)
        loss, gain = chart.callouts
        # 負の棒は地面から下がる。指すのは地面の高さ (= 棒の根元)
        self.assertAlmostEqual(self.height_of(chart, 0) - self.LINE - self.S.bar_pad_y, loss.line.bottom, delta=1)
        self.assertAlmostEqual(self.height_of(chart, 10) - self.LINE - self.S.bar_pad_y, gain.line.bottom, delta=1)

    def test_with_the_numbers_on_the_axis_the_plot_makes_room_for_them(self) -> None:
        on_the_bars, on_the_axis = self.pointed(), self.pointed(labels=False)
        widest = DEFAULT.width("25")
        self.assertAlmostEqual(self.inner(on_the_bars)[0] + widest + self.S.gap_s, self.inner(on_the_axis)[0], delta=1)

    def test_the_legend_and_the_names_of_the_categories_keep_their_room_under_the_plot(self) -> None:
        one, two = self.pointed(), self.pointed(series=(NORTH, SOUTH),
                                                callouts=({"at": "Q3", "series": "North", "text": "a"},))
        bottom = [self.inner(chart)[1] + self.inner(chart)[3] for chart in (one, two)]
        self.assertAlmostEqual(one.frame.bottom - self.LINE - self.S.gap_s, bottom[0], delta=1)
        self.assertAlmostEqual(two.frame.bottom - 2 * (self.LINE + self.S.gap_s), bottom[1], delta=1)

    def test_words_that_would_print_over_each_other_stop_the_page(self) -> None:
        with self.assertRaises(PageFullError) as stopped:
            self.pointed(callouts=({"at": "Q2", "text": "a sentence long enough to reach its neighbour " * 2},
                                   {"at": "Q3", "text": "another one just as long as the first one " * 2}))
        self.assertIn("would print over each other", str(stopped.exception))
        self.assertEqual(2, len(self.pointed(callouts=({"at": "Q1", "text": "short"},
                                                       {"at": "Q4", "text": "short too"})).callouts))

    def test_words_wider_than_the_chart_stop_the_page(self) -> None:
        with self.assertRaises(PageFullError) as stopped:
            self.pointed(callouts=({"at": "Q2", "text": "word " * 80},))
        self.assertIn("does not fit beside the chart", str(stopped.exception))

    def test_words_near_an_edge_stay_inside_and_the_line_still_stands_over_the_bar(self) -> None:
        chart = self.pointed(callouts=({"at": "Q1", "text": "a sentence that is wider than the first bar by far"},))
        callout = chart.callouts[0]
        self.assertEqual(chart.rect.left, callout.words.left)
        left, _top, wide, _tall = self.inner(chart)
        self.assertAlmostEqual(left + wide * 0.5 / 4, callout.line.left + callout.line.width / 2, delta=1)

    def test_what_a_callout_may_not_be(self) -> None:
        base = {"kind": "column", "categories": QUARTERS, "series": [NORTH, SOUTH]}
        for callouts, word in (
                ("Q3", "`callouts` is a list"), ([], "`callouts` is a list"),
                ([{"at": "Q9", "series": "North", "text": "a"}], "`at` is 'Q9'"),
                ([{"at": "Q1", "text": "a"}], "say which one with `series`"),
                ([{"at": "Q1", "series": "East", "text": "a"}], "`series` is 'East'"),
                ([{"at": "Q1", "series": "North", "text": "  "}], "`text` is the few words"),
                ([{"at": "Q1", "series": "North", "text": "two\nlines"}], "`text` is the few words"),
                ([{"at": "Q1", "series": "North", "text": 3}], "`text` is the few words"),
                ([{"at": "Q1", "series": "North", "text": "a", "x": 3}], "callout 1 does not take x"),
                ([{"at": "Q1", "series": "North", "text": "a"}, {"at": "Q1", "series": "North", "text": "b"}],
                 "already has a callout")):
            with self.subTest(word), self.assertRaises(PageTypeError) as refused:
                build({"type": "chart", "title": "t", "chart": {**base, "callouts": callouts}})
            self.assertIn(word, str(refused.exception))

    def test_written_out_the_chart_and_its_words_are_one_group_and_the_plot_is_pinned(self) -> None:
        chart = self.pointed(series=(NORTH, SOUTH), callouts=({"at": "Q3", "series": "North", "text": "the peak"},
                                                              {"at": "Q1", "series": "South", "text": "slow"}))
        xml, parts = written(chart)
        slide = parts["ppt/slides/slide1.xml"].decode("utf-8")
        group = re.search(r"<p:grpSp>.*</p:grpSp>", slide, re.S).group(0)
        self.assertEqual(1, slide.count("<p:grpSp>"))
        self.assertEqual(1, group.count("<p:graphicFrame>"))
        self.assertEqual(["the peak", "slow"], re.findall(r"<a:t>([^<]+)</a:t>", group))
        self.assertEqual(4, group.count("<p:sp>"), "two callouts are two lines and two boxes of words")
        self.assertNotIn("<p:graphicFrame>", slide.replace(group, ""), "the chart was left outside the group")
        frame = re.search(r'<p:xfrm><a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', group).groups()
        self.assertEqual((chart.frame.left, chart.frame.top, chart.frame.width, chart.frame.height),
                         tuple(int(value) for value in frame))
        layout = re.search(r"<c:plotArea><c:layout><c:manualLayout>(.*?)</c:manualLayout>", xml, re.S).group(1)
        self.assertIn('<c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>', layout)
        self.assertEqual([f"{part:.6f}" for part in chart.pinned.inner],
                         re.findall(r'<c:[xywh] val="([\d.]+)"/>', layout))
        self.assertIn('<c:max val="25.0"/>', xml)
        self.assertIn('<c:min val="0.0"/>', xml)
        self.assertIn('<c:majorUnit val="5.0"/>', xml)
        self.assertIn('<c:gapWidth val="150"/>', xml)

    def test_a_group_puts_its_pieces_where_the_page_put_them(self) -> None:
        """グループの中の座標が頁の座標と同じでなければ、検査は別の場所を見ることになる。"""
        chart = self.pointed()
        _xml, parts = written(chart)
        group = re.search(r"<p:grpSpPr>(.*?)</p:grpSpPr>", parts["ppt/slides/slide1.xml"].decode("utf-8")).group(1)
        outer = re.search(r'<a:off x="(\d+)" y="(\d+)"/><a:ext cx="(\d+)" cy="(\d+)"/>', group).groups()
        inside = re.search(r'<a:chOff x="(\d+)" y="(\d+)"/><a:chExt cx="(\d+)" cy="(\d+)"/>', group).groups()
        self.assertEqual(outer, inside)

    def test_once_baked_the_checks_find_nothing_and_the_deck_opens(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            deck = new_deck()
            for kind, series, callouts in (
                    ("column", (NORTH, SOUTH), [{"at": "Q3", "series": "North", "text": "the peak"}]),
                    ("bar", (NORTH,), [{"at": "Q2", "text": "behind"}]),
                    ("line", (NORTH,), [{"at": "Q4", "text": "a dip"}]),
                    ("stacked", (NORTH, SOUTH), [{"at": "Q1", "text": "low"}])):
                page = build({"type": "chart", "title": f"A {kind}", "footer": "Source: made up",
                              "chart": {"kind": kind, "categories": QUARTERS, "series": list(series),
                                        "callouts": callouts}})
                frame = page.theme.frame()
                for element in page.build():
                    self.assertTrue(frame.contains(element.rect), element.kind)
                add_page(deck, page.build())
            saved = save(deck, Path(tmp) / "pointed.pptx")
            self.assertEqual([], checks.run_all(saved))
            self.assertEqual(4, len(Presentation(str(saved)).slides))

    def test_a_chart_in_a_cell_takes_callouts_too(self) -> None:
        cell = {"chart": {"kind": "column", "categories": ["a", "b"], "series": [{"name": "n", "values": [1, 2]}],
                          "callouts": [{"at": "b", "text": "up"}]}}
        page = build({"type": "compose", "title": "In a cell", "rows": [{"cells": [cell, {"card": ["Beside", "x"]}]}]})
        chart = next(e for e in page.build() if isinstance(e, Chart))
        self.assertTrue(chart.rect.contains(chart.callouts[0].words))
        self.assertTrue(chart.rect.contains(chart.frame))


CHART = {"kind": "column", "categories": QUARTERS, "series": [NORTH]}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CHART_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"


class TheChartTravelsWithItsPage(unittest.TestCase):
    """⚠ **頁はグラフを関係で指しているだけ。**頁の XML と関係を写しただけでは、指す先が行き先に無い。"""

    @classmethod
    def setUpClass(cls) -> None:
        make_dot()

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.specimen = a_specimen(self.dir / "specimen.pptx", ["表紙"])

    def charts(self, name: str, *numbers: list[int]) -> Path:
        """A deck of chart pages, one per list of numbers."""
        deck = new_deck()
        for values in numbers:
            page = build({"type": "chart", "title": f"Chart {values[0]}",
                          "chart": {**CHART, "series": [{"name": "North", "values": values}]}})
            add_page(deck, page.build())
        return save(deck, self.dir / name)

    @staticmethod
    def parts(deck: Path) -> dict[str, bytes]:
        with zipfile.ZipFile(deck) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def charts_of(self, deck: Path) -> list[tuple[str, list[str]]]:
        """Per page in reading order: the chart it points at, and the numbers that chart holds."""
        parts, found = self.parts(deck), []
        Presentation(str(deck))                       # 開けること (= 関係の指す先が全部在る)
        for slide in read_pages(deck):
            rels = parts[f"ppt/slides/_rels/{slide.name}.rels"].decode("utf-8")
            for target in re.findall(r'Target="\.\./charts/([^"]+)"', rels):
                xml = parts[f"ppt/charts/{target}"].decode("utf-8")
                values = re.findall(r"<c:v>(\d+)</c:v>", xml[xml.index("<c:val>"):])
                book = re.search(r'Target="\.\./embeddings/([^"]+)"',
                                 parts[f"ppt/charts/_rels/{target}.rels"].decode("utf-8")).group(1)
                self.assertIn(f"ppt/embeddings/{book}", parts)
                found.append((target, values, book))
        return found

    def test_a_page_brought_in_comes_with_its_chart_and_the_numbers_inside_it(self) -> None:
        source = self.charts("declared.pptx", [12, 15, 21, 18], [1, 2, 3, 4])
        out = self.dir / "built.pptx"
        with Deck.open(self.specimen, out) as deck:
            deck.copy(1)
            deck.bring(source, 2, relayout=True)
            deck.bring(source, 1, relayout=True)
        found = self.charts_of(out)
        self.assertEqual([["1", "2", "3", "4"], ["12", "15", "21", "18"]], [values for _c, values, _b in found])
        self.assertEqual(2, len({chart for chart, _v, _b in found}), "two pages point at one chart")
        self.assertEqual(2, len({book for _c, _v, book in found}), "two charts keep their numbers in one book")
        self.assertEqual([], checks.run_all(out))

    def test_the_package_says_what_the_chart_and_its_numbers_are(self) -> None:
        """⚠ 絵を 1 枚も持たないテンプレートに png の登録が無かったのと同じで、グラフを 1 つも持たない
        テンプレートには、グラフの登録も、中の表の登録も無い。"""
        out = self.dir / "built.pptx"
        with Deck.open(self.specimen, out) as deck:
            deck.bring(self.charts("declared.pptx", [12, 15, 21, 18]), 1, relayout=True)
        types_xml = self.parts(out)["[Content_Types].xml"].decode("utf-8")
        self.assertNotIn("xlsx", self.parts(self.specimen)["[Content_Types].xml"].decode("utf-8"))
        self.assertIn(f'<Default Extension="xlsx" ContentType="{XLSX}"/>', types_xml)
        self.assertIn(f'<Override PartName="/ppt/charts/chart1.xml" ContentType="{CHART_TYPE}"/>', types_xml)
        self.assertEqual([], checks.rules.untyped_parts.run(out))

    def test_a_chart_brought_into_a_deck_that_has_one_takes_a_name_of_its_own(self) -> None:
        """⚠ 元の名前のまま写すと、行き先に既に在る別のグラフを上書きするか、それを指す。"""
        template = self.charts("template.pptx", [7, 7, 7, 7])
        out = self.dir / "built.pptx"
        with Deck.open(template, out) as deck:
            deck.keep(1)
            deck.bring(self.charts("declared.pptx", [12, 15, 21, 18]), 1, relayout=True)
        found = self.charts_of(out)
        self.assertEqual([("chart1.xml", ["7", "7", "7", "7"]), ("chart2.xml", ["12", "15", "21", "18"])],
                         [(chart, values) for chart, values, _book in found])
        self.assertEqual(2, len({book for _c, _v, book in found}))

    def test_a_copy_within_the_deck_gets_a_chart_of_its_own(self) -> None:
        template = self.charts("template.pptx", [7, 7, 7, 7])
        out = self.dir / "built.pptx"
        with Deck.open(template, out) as deck:
            deck.copy(1)
            deck.copy(1)
        found = self.charts_of(out)
        self.assertEqual(2, len({chart for chart, _v, _b in found}))
        self.assertEqual(2, len({book for _c, _v, book in found}))
        self.assertEqual([["7", "7", "7", "7"]] * 2, [values for _c, values, _b in found])

    def test_the_chart_of_a_page_nobody_kept_is_gone_with_its_numbers(self) -> None:
        """⚠ 載せないと決めた頁の数字が、納品物の中まで付いて行く。"""
        template = self.charts("template.pptx", [7, 7, 7, 7], [9, 9, 9, 9])
        out = self.dir / "built.pptx"
        with Deck.open(template, out) as deck:
            deck.keep(2)
        parts = self.parts(out)
        kept = sorted(name for name in parts if name.startswith(("ppt/charts/", "ppt/embeddings/")))
        self.assertEqual(3, len(kept), kept)
        self.assertEqual([["9", "9", "9", "9"]], [values for _c, values, _b in self.charts_of(out)])
        self.assertEqual(1, parts["[Content_Types].xml"].decode("utf-8").count("/ppt/charts/"))
        self.assertNotIn(b"<c:v>7</c:v>", b"".join(parts[name] for name in kept if name.endswith(".xml")))
        self.assertEqual([], checks.run_all(out))

    def test_a_chart_that_keeps_something_this_toolkit_cannot_carry_stops_the_import(self) -> None:
        source = self.charts("declared.pptx", [12, 15, 21, 18])
        broken = self.dir / "styled.pptx"
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as after:
            for item in before.infolist():
                content = before.read(item.filename)
                if item.filename == "ppt/charts/_rels/chart1.xml.rels":
                    content = content.replace(
                        b"</Relationships>",
                        b'<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        b'relationships/image" Target="../media/fill.png"/></Relationships>')
                after.writestr(item, content)
        with self.assertRaises(ValueError) as stopped:
            with Deck.open(self.specimen, self.dir / "built.pptx") as deck:
                deck.bring(broken, 1)
        self.assertIn("does not carry", str(stopped.exception))

    def test_a_chart_inside_a_group_is_carried_like_any_other(self) -> None:
        """注記を持つグラフはグループの中に在る。頁がグラフを指す関係は同じなので、同じに運ばれる。"""
        deck = new_deck()
        page = build({"type": "chart", "title": "Pointed", "chart": {**CHART, "callouts": [{"at": "Q3", "text": "up"}]}})
        add_page(deck, page.build())
        source = save(deck, self.dir / "pointed.pptx")
        out = self.dir / "built.pptx"
        with Deck.open(self.specimen, out) as built:
            built.copy(1)
            built.bring(source, 1, relayout=True)
        self.assertEqual([["12", "15", "21", "18"]], [values for _c, values, _b in self.charts_of(out)])
        slide = self.parts(out)["ppt/slides/" + read_pages(out)[1].name].decode("utf-8")
        self.assertIn("<p:grpSp>", slide)
        self.assertIn("<a:t>up</a:t>", slide)
        self.assertEqual([], checks.run_all(out))

    def test_a_deck_with_charts_converts(self) -> None:
        """LibreOffice reading it end to end is the closest check to opening it."""
        out = self.dir / "built.pptx"
        with Deck.open(self.specimen, out) as deck:
            deck.copy(1)
            deck.bring(self.charts("declared.pptx", [12, 15, 21, 18], [1, 2, 3, 4]), 1, relayout=True)
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", str(out), "--outdir", str(self.dir)],
            capture_output=True, timeout=180)
        self.assertEqual(result.returncode, 0, result.stderr.decode()[:400])
        self.assertTrue((self.dir / "built.pdf").is_file(), "the deck did not convert")


if __name__ == "__main__":
    unittest.main()
