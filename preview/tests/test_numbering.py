"""Page numbers in the preview match PowerPoint when a deck starts counting elsewhere or hides pages.

⚠ LibreOffice は presentation.xml の開始番号 (`firstSlideNum`) を読まない。表紙に番号を付けない
資料 (= 開始 0、表紙の次が 1) が、プレビューでだけ 2・3 と描かれた。

⚠ 隠した頁 (`show="0"`) は、LibreOffice の PDF には既定で出ない。隠した頁より後ろの頁は位置が
1 つずつ詰まり、画面の札が PowerPoint の頁番号より小さくなった。
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

from pptx_live_preview import render
from pptx_live_preview.render import (RenderError, numbered_as_powerpoint, pdftoppm_bin, read_numbering,
                                      soffice_bin, to_pdf)
from pptx_live_preview.state import DeckState


def _has_libreoffice() -> bool:
    try:
        soffice_bin(), pdftoppm_bin()
    except RenderError:
        return False
    return shutil.which("pdftotext") is not None


needs_libreoffice = pytest.mark.skipif(not _has_libreoffice(), reason="LibreOffice and poppler are needed")

FIELD = ('<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
         'id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
         '<a:rPr lang="en-US" sz="1800"/><a:t>&lt;#&gt;</a:t></a:fld>')


def _deck(path, first: int | None, hidden: tuple[int, ...] = ()):
    prs = Presentation()
    for position in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if position in hidden:
            slide._element.set("show", "0")
        box = slide.shapes.add_textbox(Inches(8), Inches(6.5), Inches(1), Inches(0.5))
        box.text_frame.paragraphs[0]._p.append(etree.fromstring(FIELD))
    plain = path.with_suffix(".plain.pptx")
    prs.save(plain)
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(path, "w") as out:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "ppt/presentation.xml" and first is not None:
                data = data.replace(b"<p:presentation ", f'<p:presentation firstSlideNum="{first}" '.encode(), 1)
            out.writestr(item, data)
    return path


def _slide(deck, name: str) -> str:
    with zipfile.ZipFile(deck) as z:
        return z.read(name).decode()


def _numbers(deck) -> list[str]:
    with zipfile.ZipFile(deck) as z:
        order = sorted((n for n in z.namelist() if n.startswith("ppt/slides/slide")),
                       key=lambda n: int(n.rsplit("slide", 1)[1].split(".")[0]))
        found = []
        for name in order:
            xml = z.read(name).decode()
            assert 'type="slidenum"' not in xml, "the field is left for LibreOffice to count"
            found.append(xml.split("<a:t>")[-1].split("</a:t>")[0])
        return found


def test_a_deck_that_starts_at_zero_is_numbered_as_powerpoint_numbers_it(tmp_path):
    deck = _deck(tmp_path / "w1.pptx", first=0)
    copy = numbered_as_powerpoint(deck, tmp_path / "work")
    assert copy != deck and copy.name == deck.name, "the pdf is named after the deck"
    assert _numbers(copy) == ["0", "1", "2"]
    with zipfile.ZipFile(copy) as z:
        assert 'sz="1800"' in z.read("ppt/slides/slide2.xml").decode(), "the number keeps its look"


def test_a_deck_that_starts_at_one_is_handed_over_as_it_is(tmp_path):
    for first in (None, 1):
        deck = _deck(tmp_path / f"w{first}.pptx", first=first)
        assert numbered_as_powerpoint(deck, tmp_path / "work") == deck


def test_a_deck_reports_the_number_each_page_prints_and_which_pages_are_hidden(tmp_path):
    deck = _deck(tmp_path / "w.pptx", first=0, hidden=(1,))
    numbering = read_numbering(deck)
    assert numbering.numbers == (0, 1, 2), "a hidden page takes a number too"
    assert numbering.hidden == (False, True, False)
    assert read_numbering(_deck(tmp_path / "plain.pptx", first=None)).numbers == (1, 2, 3)


def test_a_file_that_is_not_a_deck_has_no_numbering(tmp_path):
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip")
    assert read_numbering(broken) is None


def _rewritten_while_read(monkeypatch, member: str) -> None:
    """The build writes the deck again while the preview is reading it: when `member` is read, the file has
    just been cut back to the end of that part's header (= the archive's list was read whole, the part runs off
    the end of the file)."""

    class Rewritten(zipfile.ZipFile):
        def read(self, name, pwd=None):
            if getattr(name, "filename", name) == member:  # 名前でも ZipInfo でも読まれる
                info = self.getinfo(member)
                path = Path(self.filename)
                path.write_bytes(path.read_bytes()[:info.header_offset + 30 + len(info.orig_filename.encode())
                                                   + len(info.extra)])
                self.fp.close()  # 開いた時の読み込みの溜めに、小さい deck は丸ごと載っている
                self.fp = open(path, "rb")
            return super().read(name, pwd)

    monkeypatch.setattr(render.zipfile, "ZipFile", Rewritten)


def test_a_deck_caught_half_written_has_no_numbering(tmp_path, monkeypatch):
    """⚠ 書きかけの deck は目録が読め、部品を読む所で中身が尽きる (= EOFError)。読めない deck として扱う。"""
    deck = _deck(tmp_path / "w.pptx", first=0)
    _rewritten_while_read(monkeypatch, "ppt/slides/slide3.xml")
    assert read_numbering(deck) is None


@needs_libreoffice
def test_a_deck_caught_half_written_goes_to_libreoffice_rather_than_raising(tmp_path, monkeypatch):
    """⚠ 写しを作る途中で中身が尽きた例外が、見張りの thread まで抜けて thread ごと止まった。"""
    deck = _deck(tmp_path / "w.pptx", first=0)
    _rewritten_while_read(monkeypatch, "docProps/app.xml")  # 番号は読め、写しを作る所で尽きる
    try:
        to_pdf(deck, tmp_path / "work", tmp_path / "profile")
    except RenderError:
        pass  # LibreOffice が読めないと言うのは正しい失敗 (= 画面に出て、次の書き込みで描き直す)


def test_a_hidden_page_is_handed_over_shown_and_the_rest_is_left_as_it_is(tmp_path):
    deck = _deck(tmp_path / "h.pptx", first=None, hidden=(1,))
    copy = numbered_as_powerpoint(deck, tmp_path / "work")
    assert copy != deck
    assert 'show="0"' not in _slide(copy, "ppt/slides/slide2.xml")
    assert 'type="slidenum"' in _slide(copy, "ppt/slides/slide3.xml"), "a deck that starts at 1 keeps its fields"
    assert 'show="0"' in _slide(deck, "ppt/slides/slide2.xml"), "the deck on disk is not touched"
    assert read_numbering(copy).hidden == (False, False, False)


def test_a_hidden_page_in_a_deck_that_starts_at_zero_is_shown_and_numbered(tmp_path):
    copy = numbered_as_powerpoint(_deck(tmp_path / "h0.pptx", first=0, hidden=(0,)), tmp_path / "work")
    assert 'show="0"' not in _slide(copy, "ppt/slides/slide1.xml")
    assert _numbers(copy) == ["0", "1", "2"]


def test_a_page_hidden_the_xml_schema_way_is_hidden_too(tmp_path):
    deck = _deck(tmp_path / "f.pptx", first=None)
    with zipfile.ZipFile(deck) as source, zipfile.ZipFile(tmp_path / "false.pptx", "w") as out:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "ppt/slides/slide3.xml":
                data = data.replace(b"<p:sld ", b'<p:sld show="false" ', 1)
            out.writestr(item, data)
    assert read_numbering(tmp_path / "false.pptx").hidden == (False, False, True)


def _detail(tmp_path, first: int | None, hidden: tuple[int, ...] = ()) -> dict:
    """The deck as the screen gets it, drawn by the real LibreOffice."""
    deck = DeckState(_deck(tmp_path / "real.pptx", first, hidden), cache_dir=tmp_path / "cache")
    deck.rerender_if_stale(force=True)
    assert deck.error is None, deck.error
    return deck.detail()


def _printed(tmp_path, first: int | None, hidden: tuple[int, ...] = ()) -> list[str]:
    """What each page of the drawn PDF prints as its number (= the text of the page, as LibreOffice drew it)."""
    pdf = to_pdf(_deck(tmp_path / "print.pptx", first, hidden), tmp_path / "work", tmp_path / "profile")
    text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, check=True).stdout
    return [page.strip() for page in text.split("\f") if page.strip()]


@needs_libreoffice
def test_the_screen_numbers_a_deck_that_starts_at_zero_as_the_pages_print(tmp_path):
    detail = _detail(tmp_path, first=0)
    assert len(detail["pages"]) == 3
    assert detail["numbers"] == [0, 1, 2]
    assert detail["hidden"] == [False, False, False]


@needs_libreoffice
def test_the_screen_counts_a_hidden_page_and_draws_it(tmp_path):
    detail = _detail(tmp_path, first=None, hidden=(1,))
    assert len(detail["pages"]) == 3, "the hidden page is drawn, so the pages after it keep their place"
    assert detail["numbers"] == [1, 2, 3]
    assert detail["hidden"] == [False, True, False]


@needs_libreoffice
def test_a_hidden_page_and_a_start_elsewhere_work_together(tmp_path):
    detail = _detail(tmp_path, first=0, hidden=(0,))
    assert len(detail["pages"]) == 3
    assert detail["numbers"] == [0, 1, 2]
    assert detail["hidden"] == [True, False, False]
    assert _printed(tmp_path, first=0, hidden=(0,)) == ["0", "1", "2"]
